#!/usr/bin/env python3
"""小慧 · 赛场完整初版 PPT（约 1 小时：开场 → 技能 → 联调 → 收尾）。"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
CREAM = RGBColor(0xF7, 0xF4, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xD2, 0xC5)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
TEAL = RGBColor(0x2F, 0x6F, 0x6B)
DASH = RGBColor(0x8A, 0x86, 0x7C)
FONT = "Microsoft YaHei"
OUT = Path("/opt/openclaw-hr/docs/小慧-赛场完整初版.pptx")
POLICY = Path("/opt/openclaw-hr/docs/ppt/assets/policy")
TOTAL = 29


def set_run(run, size, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    if line:
        sh.line.fill.solid()
        sh.line.color.rgb = line
    return sh


def box(slide, l, t, w, h, fill=WHITE, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.solid()
    sh.line.color.rgb = line if line else fill
    sh.adjustments[0] = 0.08
    return sh


def tb(slide, l, t, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    try:
        tf.anchor = anchor
    except Exception:
        pass
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color)
    return sh


def footer(slide, n, section=""):
    rect(slide, Inches(0), Inches(7.28), W, Inches(0.22), NAVY)
    left = f"小慧 · 完整初版{('  ·  ' + section) if section else ''}"
    tb(slide, Inches(0.4), Inches(7.28), Inches(9), Inches(0.22), left, 9, False, WHITE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    tb(slide, Inches(10.5), Inches(7.28), Inches(2.4), Inches(0.22), f"{n} / {TOTAL}", 9, False, WHITE, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def slide_bg(slide, color=CREAM):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def section_bar(slide, label):
    rect(slide, Inches(0), Inches(0), W, Inches(0.08), ACCENT)
    tb(slide, Inches(0.5), Inches(0.18), Inches(12), Inches(0.28), label, 11, True, ACCENT)


def title(slide, text, sub=None):
    tb(slide, Inches(0.5), Inches(0.42), Inches(12.3), Inches(0.42), text, 26, True, NAVY)
    if sub:
        tb(slide, Inches(0.5), Inches(0.88), Inches(12.3), Inches(0.36), sub, 14, False, MUTED)


def crop_caption(src: Path, bar_px: int = 54) -> Path:
    dest = src.with_name(src.stem + "-embed.png")
    im = PILImage.open(src).convert("RGB")
    if im.height > bar_px + 40:
        im = im.crop((0, 0, im.width, im.height - bar_px))
    im.save(dest, optimize=True)
    return dest


def pic(slide, path, l, t, w):
    return slide.shapes.add_picture(str(path), l, t, width=w)


def cover(prs):
    s = new_slide(prs)
    slide_bg(s, NAVY)
    rect(s, Inches(0), Inches(0), Inches(0.18), H, ACCENT)
    tb(s, Inches(0.7), Inches(1.2), Inches(12), Inches(0.4), "2026 云南省职业院校技能大赛  ·  人工智能赛道  ·  高职组", 14, False, RGBColor(0xC8, 0xD0, 0xE0))
    tb(s, Inches(0.7), Inches(1.75), Inches(12), Inches(1.0), "小慧", 54, True, WHITE)
    tb(s, Inches(0.7), Inches(2.9), Inches(11.5), Inches(1.2), "让一个人的人事岗接得住三件事\n拍板权仍在人", 28, False, RGBColor(0xF0, 0xE6, 0xD8))
    tb(s, Inches(0.7), Inches(4.4), Inches(11), Inches(1.0), "HR 数字员工  ·  招聘初筛 · 入职催办 · 制度问答\n脚本过关，人做决定", 16, False, RGBColor(0xA8, 0xB4, 0xC8))
    tb(s, Inches(0.7), Inches(6.4), Inches(11), Inches(0.4), "四人参赛  ·  姓名待填  ·  现场一小时：开场约 6–7 分钟 → 技能 → 联调", 13, False, RGBColor(0x8A, 0x96, 0xAA))
    notes(s, "5 秒。不要念赛道全称。开口进下一页。")


def team(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "开场")
    title(s, "本场怎么看", "四人各守一扇门。这一页是目录，不是自我介绍。")
    headers = ["同学（姓名待填）", "本场职责", "技能段（10–15 分钟）", "联调（一块屏）"]
    rows = [
        ["主控", "招聘（联调时串词）", "脱敏；评分只给建议；数字回核", "串词；员工1走完招聘主链"],
        ["女2", "人才发展（培训+绩效）", "培训点名/不编；绩效缺自评停、不定级", "不接进联调主链，待命"],
        ["女1", "入职", "合同/账号 halt，人批 yes/no", "主链走完后，顺序接「入职办到哪」"],
        ["男1", "行政", "非 HR 拒查工资；政策无命中不编", "飞书待命，不与招聘同屏"],
    ]
    col_w = [Inches(2.15), Inches(3.2), Inches(3.85), Inches(3.1)]
    x0, y0, rh = Inches(0.5), Inches(1.4), Inches(1.1)
    x = x0
    for i, h in enumerate(headers):
        rect(s, x, y0, col_w[i], Inches(0.4), NAVY)
        tb(s, x + Inches(0.08), y0, col_w[i] - Inches(0.1), Inches(0.4), h, 12, True, WHITE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
        x += col_w[i]
    for ri, row in enumerate(rows):
        y = y0 + Inches(0.4) + ri * rh
        fill = WHITE if ri % 2 == 0 else RGBColor(0xEE, 0xE8, 0xDF)
        x = x0
        for i, cell in enumerate(row):
            rect(s, x, y, col_w[i], rh, fill, LINE)
            tb(s, x + Inches(0.1), y + Inches(0.12), col_w[i] - Inches(0.18), rh - Inches(0.2), cell, 13, i == 0, INK)
            x += col_w[i]
    tb(s, Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.55), "技能段：一人一块屏，另外三人低头干活不抢镜头。联调：始终同一块投屏。", 14, False, TEAL)
    footer(s, 2, "开场")
    notes(s, "20 秒。指表即可。一人一块：招聘 / 入职 / 行政 / 人才发展。")


def pic_height(path, width):
    im = PILImage.open(path)
    return int(width * im.height / im.width)


def national(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "开场 · 项目背景")
    title(s, "国家到省：智能体，落到 OPC", "截图是书证。口播只念底下那一句。不要把智能体讲成多智能体协同。")
    n_agent = crop_caption(POLICY / "国家十五五-智能体段.png")
    n_plus = crop_caption(POLICY / "国家十五五-人工智能加-千行百业.png")
    yn_opc = crop_caption(POLICY / "云南方案-智能体OPC段.png")
    y0 = Inches(1.28)
    col_w = Inches(5.9)
    h1 = pic_height(n_agent, col_w)
    h2 = pic_height(n_plus, col_w)
    frame_h = max(h1, h2) + Inches(0.16)
    box(s, Inches(0.5), y0, Inches(6.05), frame_h, WHITE, LINE)
    box(s, Inches(6.75), y0, Inches(6.05), frame_h, WHITE, LINE)
    pic(s, n_agent, Inches(0.58), y0 + Inches(0.08), col_w)
    pic(s, n_plus, Inches(6.83), y0 + Inches(0.08), col_w)
    lab_y = y0 + frame_h + Inches(0.04)
    tb(s, Inches(0.5), lab_y, Inches(6.05), Inches(0.28), "国家十五五纲要  ·  印刷页第 35 页", 12, True, TEAL)
    tb(s, Inches(6.75), lab_y, Inches(6.05), Inches(0.28), "同上  ·  第十三章  人工智能+ 赋能千行百业", 12, True, TEAL)
    y1 = lab_y + Inches(0.32)
    tb(s, Inches(0.5), y1, Inches(12.3), Inches(0.28), "云南省人民政府办公厅  ·  云政办发〔2026〕27号  ·  2026-05-20", 12, True, ACCENT)
    yn_w = Inches(12.06)
    yn_h = pic_height(yn_opc, yn_w)
    box(s, Inches(0.5), y1 + Inches(0.28), Inches(12.3), yn_h + Inches(0.16), WHITE, LINE)
    pic(s, yn_opc, Inches(0.62), y1 + Inches(0.36), yn_w)
    box(s, Inches(0.5), Inches(6.12), Inches(12.3), Inches(1.02), RGBColor(0xEE, 0xE8, 0xDF), LINE)
    tb(
        s,
        Inches(0.7),
        Inches(6.22),
        Inches(11.9),
        Inches(0.82),
        "国家鼓励智能体和「人工智能+」；云南把智能体落到重点领域，并鼓励 OPC。\n这种极小团队往往没有人事部——小慧接的是人事三件事，拍板权仍在人。",
        15,
        True,
        NAVY,
    )
    footer(s, 3, "开场")
    notes(
        s,
        "40 秒。先指国家两张红框：智能体、人工智能+千行百业。再指云南红框：智能体+OPC。"
        "念总结句。不要念章节号，不要解释 OPC 工商含义，不要说 OPD，不要说多智能体协同。",
    )


def hr_pains(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "开场 · 项目背景")
    title(s, "人事现场：杂、高、差、少", "四字只为好记。内容是人事现场，不是降本百分比。")
    quads = [
        ("杂", "事务杂", "简历、催办、问制度同时砸来。\n制度散在手册、群消息、口头里。\n入口还有飞书、邮箱、当面。"),
        ("高", "风险高", "按年龄性别筛人，合规说不清。\n模型直接定录用，个保法过不了。\n年假天数答错，变成劳动争议。"),
        ("差", "协同差", "合同、IT、行政各填各的。\nHR 当催办员，进度对不齐。\n报表在多张表里对不拢。"),
        ("少", "人手少", "中小企业、OPC、院校人事处\n往往就一两个人。\n没有完整 eHR 班底。"),
    ]
    x = Inches(0.5)
    for word, tag, body in quads:
        box(s, x, Inches(1.4), Inches(3.0), Inches(4.35), WHITE, LINE)
        tb(s, x + Inches(0.2), Inches(1.55), Inches(2.6), Inches(0.7), word, 40, True, ACCENT)
        tb(s, x + Inches(0.2), Inches(2.3), Inches(2.6), Inches(0.4), tag, 16, True, NAVY)
        tb(s, x + Inches(0.2), Inches(2.8), Inches(2.6), Inches(2.6), body, 14, False, INK)
        x += Inches(3.2)
    tb(s, Inches(0.5), Inches(5.95), Inches(12.3), Inches(0.95), "急需的不是再买一套人事系统，是趴在现有入口上的数字员工：接得住杂事，把高风险的门留给人。", 15, True, NAVY)
    footer(s, 4, "开场")
    notes(s, "50 秒。四字各一句。禁止念「人力成本降低百分之多少」。落到：接得住、门留给人。")


def background(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "开场 · 项目背景")
    title(s, "落到谁身上，红线在哪", "OPC 是组织形态。个保法是不能越过的门。OPD 不是云南政策词。")
    box(s, Inches(0.5), Inches(1.35), Inches(6.0), Inches(5.5), WHITE, LINE)
    rect(s, Inches(0.5), Inches(1.35), Inches(0.12), Inches(5.5), TEAL)
    tb(s, Inches(0.85), Inches(1.55), Inches(5.4), Inches(0.35), "谁在用", 12, True, TEAL)
    tb(s, Inches(0.85), Inches(2.05), Inches(5.4), Inches(4.4), "上一页省政府方案鼓励人工智能 OPC。这种极小团队往往没有人事部。昆明西山、玉溪红塔已有公开报道的 OPC 创业社区。\n\n中小企业、职业院校人事处同样是一两个人扛初筛、催办、问制度。\n\n数据仍是教学夹具，不冒充已上线某企业。全国同期粤深也有 OPC 专项，本场主引云南文件。", 15, False, INK)
    box(s, Inches(6.8), Inches(1.35), Inches(6.0), Inches(5.5), WHITE, LINE)
    rect(s, Inches(6.8), Inches(1.35), Inches(0.12), Inches(5.5), ACCENT)
    tb(s, Inches(7.15), Inches(1.55), Inches(5.4), Inches(0.35), "不能做什么", 12, True, ACCENT)
    tb(s, Inches(7.15), Inches(2.05), Inches(5.4), Inches(4.4), "《个人信息保护法》第 24 条：对个人权益有重大影响的决定，个人有权拒绝仅由自动化作出的结果。\n\n《就业促进法》禁止就业歧视：年龄、性别、婚育不能当成筛选依据。\n\n所以：小慧让一个人接得住，这个人仍然拍板。一人 ≠ 无人。", 15, False, INK)
    footer(s, 5, "开场")
    notes(s, "45 秒。左边是谁，右边是红线。念完「一人不等于无人」。不要解释 OPD，不要把人社厅创业扶持说成项目甲方。")


def problems(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "开场 · 项目背景")
    title(s, "调研 · 落到三件可演的事", "痛点再多，本场只验这三件。不是人力资源全流程。")
    items = [
        ("01", "简历多，不敢用 AI 刷人", "年龄、性别一看就有。模型看得见，就无法证明没用过。"),
        ("02", "入职卡在合同和 IT 之间", "各部门各填各的，HR 当催办员。门必须人开。"),
        ("03", "员工把制度问题砸给 HR", "年假报销反复问。答错变成争议，不能用常识硬编。"),
    ]
    x = Inches(0.5)
    for num, t, b in items:
        box(s, x, Inches(1.4), Inches(3.95), Inches(3.15), WHITE, LINE)
        tb(s, x + Inches(0.25), Inches(1.55), Inches(3.4), Inches(0.4), num, 20, True, ACCENT)
        tb(s, x + Inches(0.25), Inches(2.05), Inches(3.45), Inches(0.85), t, 18, True, NAVY)
        tb(s, x + Inches(0.25), Inches(2.95), Inches(3.45), Inches(1.3), b, 14, False, MUTED)
        x += Inches(4.15)
    box(s, Inches(0.5), Inches(4.75), Inches(12.3), Inches(2.2), RGBColor(0xEE, 0xE8, 0xDF), LINE)
    tb(s, Inches(0.75), Inches(4.95), Inches(11.8), Inches(0.4), "实地调研照片（待补）", 16, True, NAVY)
    tb(s, Inches(0.75), Inches(5.4), Inches(11.8), Inches(1.3), "专家要求 1–2 张你们出镜的实拍。没有就空着，不要 P 图、不要写假公司名。\n数据是教学夹具。口述不要说成已上线某企业。有甲方后只换单位类型，不编花名册。", 14, False, MUTED)
    footer(s, 6, "开场")
    notes(s, "50 秒。三件事各一句。照片诚实说还没有。")


def scheme(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "开场")
    title(s, "整体方案", "人面对的是小慧。花名册在 HCM，制度在知识库，小慧只读不写。")
    cards = [
        ("门口", "WebChat / CLI → 总控只分发、自己不答\n飞书 → 直达行政"),
        ("原则", "按岗位步骤拆成可校验脚本\n模型写和抽，过关由脚本说"),
        ("人守的拍板", "录用 / 合同 / 账号 / 他人工资\n一人接住，人仍拍板"),
    ]
    x = Inches(0.5)
    for h, b in cards:
        box(s, x, Inches(1.35), Inches(3.95), Inches(1.55), WHITE, LINE)
        tb(s, x + Inches(0.2), Inches(1.45), Inches(3.55), Inches(0.32), h, 13, True, TEAL)
        tb(s, x + Inches(0.2), Inches(1.8), Inches(3.55), Inches(0.95), b, 13, False, INK)
        x += Inches(4.15)
    lines = [
        (True, "招聘", "脱敏评分，只给建议"),
        (True, "入职", "只读进度，两道人批的门"),
        (True, "培训", "点名出草稿，讲师审了才用"),
        (True, "绩效", "360 聚合，不定级"),
        (True, "行政", "制度引原文，月报脚本算"),
    ]
    x = Inches(0.5)
    for live, name, desc in lines:
        fill = NAVY if live else WHITE
        tag_c = ACCENT if live else DASH
        box(s, x, Inches(3.15), Inches(2.4), Inches(2.55), fill, NAVY if live else LINE)
        tb(s, x + Inches(0.12), Inches(3.25), Inches(2.15), Inches(0.32), "本场演示" if live else "本场不展开", 11, True, tag_c)
        tb(s, x + Inches(0.12), Inches(3.6), Inches(2.15), Inches(0.45), name, 22, True, WHITE if live else NAVY)
        tb(s, x + Inches(0.12), Inches(4.15), Inches(2.15), Inches(1.3), desc, 13, False, WHITE if live else MUTED)
        x += Inches(2.52)
    tb(s, Inches(0.5), Inches(5.9), Inches(12.3), Inches(1.1), "子 Agent 互不派活。筛完简历不会自动办入职——录用是人定的，入职是另一单。\n培训、绩效由女2在技能段打开，不接进联调主链。", 14, False, MUTED)
    footer(s, 7, "开场")
    notes(s, "70 秒。五条都点到。培训绩效说女2技能段演，不要在开场展开 ADDIE。禁止讲多智能体协同、模型名、端口。")


def doors(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "开场")
    title(s, "本场将打开的三扇门", "不是亮点汇报。每扇门后面都有能跑出退出码的动作。")
    rows = [
        ("门 1  ·  招聘", "敏感字段进评分", "物理删除。故意留「性别」，退出码 1，不能打分。", "规范性 · 难易度"),
        ("门 2  ·  入职", "合同 / 账号批准", "halt 等人。yes 才走，no 则 cancelled。", "规范性 · 熟练度"),
        ("门 3  ·  行政", "政策与他人工资", "无命中不编。非 HR 查工资，退出码 3。", "应用价值 · 规范性"),
    ]
    y = Inches(1.4)
    for a, b, c, d in rows:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.55), WHITE, LINE)
        rect(s, Inches(0.5), y, Inches(0.14), Inches(1.55), ACCENT)
        tb(s, Inches(0.9), y + Inches(0.15), Inches(4.3), Inches(0.4), a, 16, True, NAVY)
        tb(s, Inches(5.2), y + Inches(0.18), Inches(4.4), Inches(0.35), b, 16, True, ACCENT)
        tb(s, Inches(9.7), y + Inches(0.18), Inches(2.9), Inches(0.35), d, 12, False, TEAL, PP_ALIGN.RIGHT)
        tb(s, Inches(0.9), y + Inches(0.65), Inches(11.5), Inches(0.7), c, 16, False, INK)
        y += Inches(1.7)
    tb(s, Inches(0.5), Inches(6.55), Inches(12.3), Inches(0.5), "女2另开人才发展两刀（点名不编、缺自评停），不进联调。坏数据只在技能段用。", 14, True, NAVY)
    footer(s, 8, "开场")
    notes(s, "50 秒。三扇门各读动作。补一句女2两刀。然后进入技能段，不要翻回方案。")


def hour_map(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能展示")
    title(s, "这一小时怎么过", "评委主要在 45–50 分钟技能段结束时打分。开场加厚到约 7 分钟，不要再延长。")
    blocks = [
        ("0–7 分", "开场", "政策书证、四字痛点、红线、三件事、方案、三扇门"),
        ("7–45 分", "技能", "主控招聘三刀 → 女2培训绩效 → 女1入职 → 男1行政"),
        ("45–55 分", "联调", "一块屏：员工1走完招聘主链，再顺序接入职"),
        ("55–60 分", "收口", "二次开发、边界、展望、致谢"),
    ]
    x = Inches(0.5)
    for t, h, b in blocks:
        box(s, x, Inches(1.45), Inches(3.0), Inches(3.4), WHITE, LINE)
        tb(s, x + Inches(0.2), Inches(1.65), Inches(2.6), Inches(0.4), t, 14, True, ACCENT)
        tb(s, x + Inches(0.2), Inches(2.15), Inches(2.6), Inches(0.5), h, 22, True, NAVY)
        tb(s, x + Inches(0.2), Inches(2.8), Inches(2.6), Inches(1.7), b, 15, False, MUTED)
        x += Inches(3.2)
    tb(s, Inches(0.5), Inches(5.15), Inches(12.3), Inches(1.8), "技能段对着屏幕动手，PPT 只提示「问题 → 卡哪一刀 → 失败什么样」。\n不要在技能段念架构、模型、加班、市场潜力。\n现场写的是业务逻辑二三十行，不写初始化和 YAML。", 16, False, INK)
    footer(s, 9, "技能")
    notes(s, "15 秒过。这页给评委对表。立刻把投屏切到操作。")


def recruit_chain(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 招聘地图")
    title(s, "招聘主链八步", "模型只负责写和抽。过不过关由脚本说。本场精讲第 5–7 步。")
    steps = [
        ("1 收信", "inbox/\n队列不拦 / 必重跑"),
        ("2 转 MD", "不调模型\n不抽字段"),
        ("3 形状闸", "parsed/\n校 schema"),
        ("4 内容闸", "audited/\n对原文"),
        ("5 脱敏", "clean/\n本场手写"),
        ("6 评分", "八维+建议\n脚本定档"),
        ("7 数字回核", "只认原简历\n主控讲"),
        ("8 画像", "校验后再发\n不另算分"),
    ]
    x = Inches(0.4)
    for i, (h, b) in enumerate(steps):
        live = i in (4, 5, 6)
        fill = NAVY if live else WHITE
        box(s, x, Inches(1.45), Inches(1.55), Inches(3.55), fill, ACCENT if live else LINE)
        tb(s, x + Inches(0.08), Inches(1.55), Inches(1.4), Inches(0.85), h, 14, True, WHITE if live else NAVY)
        tb(s, x + Inches(0.08), Inches(2.5), Inches(1.4), Inches(2.2), b, 12, False, WHITE if live else MUTED)
        x += Inches(1.62)
    tb(s, Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.7), "人才库、面试是支链，本场不展开。被问：名次靠余弦；没有合格转写稿不编纪要。\n待确认岗位不停车；同一人新稿必须重跑转 MD→解析→脱敏→增量评分。", 15, False, MUTED)
    footer(s, 10, "招聘")
    notes(s, "30 秒指图。深色三步都是主控写/跑。然后开始脱敏。")


def desensitize(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 主控 · 门 1")
    title(s, "脱敏：看得见就无法证明没用过", "评分第一层。吃过内容闸的 parsed 副本，产出 scoring/clean。原件不动。")
    box(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(5.4), WHITE, LINE)
    tb(s, Inches(0.75), Inches(1.55), Inches(5.6), Inches(0.4), "删什么", 14, True, ACCENT)
    tb(s, Inches(0.75), Inches(2.05), Inches(5.6), Inches(4.4), "删：性别、年龄/出生、婚育、户籍、民族、健康/残疾、院校层次（985 / 第一学历）、政治面貌。\n\n不删：某某大学、计算机本科、学位——学历事实要留给评分。\n\n递归扫每一层。剥完再扫一遍。残留任一敏感键 → 退出码 1，直接停，不能打分。\n不是「加进黑名单再跑一遍算成功」。", 15, False, INK)
    box(s, Inches(6.85), Inches(1.4), Inches(5.95), Inches(5.4), WHITE, LINE)
    tb(s, Inches(7.1), Inches(1.55), Inches(5.5), Inches(0.4), "现场做什么", 14, True, TEAL)
    tb(s, Inches(7.1), Inches(2.05), Inches(5.5), Inches(4.4), "写内核：黑名单 + 递归 delete + 自检。\n\n喂一份仍带「性别」的 JSON。\n\n评委要看见退出码 1，没有分数。\n\nprompt 写「请忽略年龄」不算数。", 16, False, INK)
    footer(s, 11, "招聘")
    notes(s, "主控 10–15 分钟的前半。对着代码讲，PPT 不念。被问个保法：第 24 条可证明合规。")


def scoring(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 主控 · 建议权")
    title(s, "评分：模型打分，脚本定档，人不交权", "评分第二层。必须已经是脱敏 JSON。数字真不真是下一步。")
    items = [
        ("八维（脚本钉死）", "专业技能、经验匹配、逻辑能力、学习力、沟通能力、协作能力、抗压能力、价值观。每维 1–5。专业技能和经验匹配权重大。"),
        ("综合分 vs 评级", "综合分由模型打。A/B/C/D 由脚本切：≥85 A，≥75 B，≥60 C，否则 D。65 分是 C。模型写成 B 不算。"),
        ("只给建议", "标「建议进入面试 / 建议人工复核」。禁止出现淘汰、录用。硬伤只装真实性疑点，不能刷人。"),
        ("硬伤 ≠ 顾虑", "硬伤：经历时间重叠、学历存疑、简历自相矛盾。顾虑：薪资、城市、跳槽。顾虑进不了硬伤。"),
    ]
    y = Inches(1.35)
    for h, b in items:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.3), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.12), Inches(3.2), Inches(1.05), h, 16, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(4.1), y + Inches(0.15), Inches(8.4), Inches(1.05), b, 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.4)
    footer(s, 12, "招聘")
    notes(s, "主控后半。现场写 ratingOf + 禁止淘汰字样校验。不要说自研评分算法。")


def numbers(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 主控 · 回核")
    title(s, "数字回核：只认原简历", "评分第三层。防的是把 JD 数字安到候选人头上。主控讲完这刀再交女2。")
    box(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(5.4), WHITE, LINE)
    tb(s, Inches(0.75), Inches(1.55), Inches(5.6), Inches(0.4), "核什么", 14, True, ACCENT)
    tb(s, Inches(0.75), Inches(2.05), Inches(5.6), Inches(4.5), "核评分亮点/理由里的数字，能不能在原简历 md 里找到。\n\n脚本不接 JD 文件，基准不会被岗位描述污染。\n\n不管简历自己吹不吹——那是硬伤和面试的事。\n\n每一次评分都要跑，不是抽查。", 15, False, INK)
    box(s, Inches(6.85), Inches(1.4), Inches(5.95), Inches(5.4), WHITE, LINE)
    tb(s, Inches(7.1), Inches(1.55), Inches(5.5), Inches(0.4), "失败长什么样", 14, True, TEAL)
    tb(s, Inches(7.1), Inches(2.05), Inches(5.5), Inches(4.5), "找到串台并改回简历真值：退出码 0，综合分不改。\n\n判定程序没跑成：退出码 1。禁止告诉 HR「没有虚构」。\n\n现场喂：简历写 12.8 万 TPS，评分写成 JD 里的 58 万。", 15, False, INK)
    footer(s, 13, "招聘")
    notes(s, "主控招聘段收尾。写抽数字+只对原简历。讲完交女2，不要再筛一遍简历。")


def training_gate(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 女2 · 培训")
    title(s, "培训：点名才出稿，不编，讲师审了才用", "同一座位连续两刀的第一刀。不是办一场培训，是出能被拦住的草稿。")
    rows = [
        ("点名", "用户点哪几份就出哪几份。未点名不出全套 8 份。现场点讲师手册 + 测试题即可。"),
        ("不编", "source-docs 里没有的制度、数字、名称不写进去。没有素材就停，不靠常识补。"),
        ("标注", "每份标明「AI 辅助生成，讲师审核后交付」。讲师手册禁止逐字稿。"),
        ("现场", "评委要看见：只出点过的文件；header 有审核标注；没有把八份稿讲成 SOP。"),
    ]
    y = Inches(1.35)
    for h, b in rows:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.3), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.18), Inches(2.2), Inches(0.95), h, 18, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(3.2), y + Inches(0.18), Inches(9.3), Inches(0.95), b, 15, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.38)
    footer(s, 14, "人才发展")
    notes(s, "女2 前半。对着产出目录讲。不要组织签到、不要念 ADDIE 五阶段。点完立刻切绩效。")


def performance_gate(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 女2 · 绩效")
    title(s, "绩效：缺自评就停，聚合不定级", "第二刀仍在同一块屏。定级留给主管和校准会。")
    box(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(5.4), WHITE, LINE)
    tb(s, Inches(0.75), Inches(1.55), Inches(5.6), Inches(0.4), "闸在哪", 14, True, ACCENT)
    tb(s, Inches(0.75), Inches(2.05), Inches(5.6), Inches(4.5), "缺自评：拒绝聚合，不出总表。\n\n不给出 S / A / B / C。看板和报告里不能出现等级。\n\n评价人只用角色称呼（上级/平级/下属/自评），不用真名。\n\n分数由脚本从反馈表算，模型不改数。", 15, False, INK)
    box(s, Inches(6.85), Inches(1.4), Inches(5.95), Inches(5.4), WHITE, LINE)
    tb(s, Inches(7.1), Inches(1.55), Inches(5.5), Inches(0.4), "现场做什么", 14, True, TEAL)
    tb(s, Inches(7.1), Inches(2.05), Inches(5.5), Inches(4.5), "先抽掉自评再跑，评委要看见拒绝。\n\n再跑齐套 360，亮雷达看板：无定级字样。\n\n口述：定级仍走主管初评、隔级复核、校准会。\n\n这一刀不接进联调主链。", 15, False, INK)
    footer(s, 15, "人才发展")
    notes(s, "女2 后半。同一屏幕连做。不要讲绩效算法，不要把夹具看板说成校准会结果。")


def onboarding_gate(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 女1 · 门 2")
    title(s, "入职：只读进度，两道门必须人开", "不替代合同/IT/行政办事。不向 HCM 写入、不悄悄建人。")
    rows = [
        ("做什么", "对照 HCM 五节点催办；合同门、账号门 halt 等人；五节点全绿才挂试用期跟进。"),
        ("人做什么", "合同条款批不批、账号权限批不批。approvalId 不能顶替 resumeToken。"),
        ("停在哪", "halt ≥ 3 天只催办，不自动 resume yes。人说 no → cancelled。HCM 未齐时提示请勿批准。"),
        ("不讲什么", "不要把 Buddy、7/30/60/90 讲成国内标准 SOP。社保公积金仍由人事在原系统办。"),
    ]
    y = Inches(1.4)
    for h, b in rows:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.25), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.15), Inches(2.4), Inches(0.95), h, 16, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(3.3), y + Inches(0.15), Inches(9.2), Inches(0.95), b, 15, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.32)
    footer(s, 16, "入职")
    notes(s, "女1 前半。屏幕上亮 halt，不要念 stub。口述：跨部门催办是真痛点，节点清单不是全套国内 SOP。")


def onboarding_fail(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 女1 · 现场")
    title(s, "入职现场：流程必须停住", "评委要看见「系统不会替人点头」。")
    cols = [
        ("操作", "start 员工1 → 停在合同门\n把 PROMPT 原文和 resumeToken 亮给评委"),
        ("人批", "resume yes → 往下走\nresume no → cancelled，Buddy/complete 不执行"),
        ("坏路径", "不要代批。不要用模型常识编「已开通」。\nHCM 不可达就说不可达。"),
    ]
    x = Inches(0.5)
    for h, b in cols:
        box(s, x, Inches(1.45), Inches(3.95), Inches(4.0), WHITE, LINE)
        tb(s, x + Inches(0.25), Inches(1.65), Inches(3.45), Inches(0.5), h, 18, True, NAVY)
        tb(s, x + Inches(0.25), Inches(2.3), Inches(3.45), Inches(2.8), b, 16, False, INK)
        x += Inches(4.15)
    tb(s, Inches(0.5), Inches(5.7), Inches(12.3), Inches(1.2), "联调时入职不跟招聘抢屏：招聘主链走完，停一秒，总控再说「员工1入职办到哪了」——那是新的一单。", 15, False, TEAL)
    footer(s, 17, "入职")
    notes(s, "女1 后半。现场 halt。联调提醒写在页脚，技能段不必展开联调。")


def admin_access(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 男1 · 门 3")
    title(s, "行政：飞书直达，身份在脚本里挡", "员工问制度可以。他人工资和月报不是谁都能出。")
    box(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(5.4), WHITE, LINE)
    tb(s, Inches(0.75), Inches(1.55), Inches(5.6), Inches(0.4), "通道", 14, True, ACCENT)
    tb(s, Inches(0.75), Inches(2.1), Inches(5.6), Inches(4.4), "飞书私聊/群 @ 直达行政，不经过总控。\n\nWebChat / CLI 才走总控再转来。\n\n现场写/跑身份闸：飞书 HR 或无 sender（WebChat/CLI）才允许查他人工资、出月报。主管/员工 → 退出码 3。\n\n不要把「三级分权」讲成生产已落死。演示期有兜底，上场说：按发送者限制工具。", 15, False, INK)
    box(s, Inches(6.85), Inches(1.4), Inches(5.95), Inches(5.4), WHITE, LINE)
    tb(s, Inches(7.1), Inches(1.55), Inches(5.5), Inches(0.4), "现场", 14, True, TEAL)
    tb(s, Inches(7.1), Inches(2.1), Inches(5.5), Inches(4.4), "用非 HR 的身份去查「员工1工资」。\n\n评委要看见拒绝，不是看见工资数字。\n\n红线口述即可：个案争议、劳动仲裁不答、转人工。不要假装已经写成分类器。", 16, False, INK)
    footer(s, 18, "行政")
    notes(s, "男1 前半。跑 hr-data-access / er-admin-run。不要念 ou_。")


def admin_policy(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 男1 · 原文与数字")
    title(s, "政策只引原文，月报数字脚本算", "检索得到 ≠ 该回答。模型不改数。")
    cols = [
        ("政策", "先检索知识库。有命中：按原文答，关键数字逐字一致，结尾写出处。\n无命中 / 凭证缺失：不编，转人工。"),
        ("月报", "pandas 读表算出指标。缺 CSV 就停。模型不碰数字。\n现场最多跑一次补跑，不要手写整张报表。"),
        ("边界", "不替员工查「我的工资条」。劳动仲裁不下结论。\n本场不把某某公司制度说成真企业手册——知识库待换成调研材料。"),
    ]
    x = Inches(0.5)
    for h, b in cols:
        box(s, x, Inches(1.45), Inches(3.95), Inches(4.5), WHITE, LINE)
        tb(s, x + Inches(0.2), Inches(1.65), Inches(3.55), Inches(0.5), h, 18, True, NAVY)
        tb(s, x + Inches(0.2), Inches(2.3), Inches(3.55), Inches(3.3), b, 15, False, INK)
        x += Inches(4.15)
    footer(s, 19, "行政")
    notes(s, "男1 后半。问年假：有原文才答。问工资少了 800：转人工。")


def router(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "技能 · 总控")
    title(s, "总控只做三件事", "自己不答年假、分数、工资。判不准就问人，不猜。")
    items = [
        ("判断", "看业务动作：筛简历 → 招聘；入职办到哪 → 入职；点名出培训稿 → 培训；360 → 绩效；年假 → 行政。"),
        ("派发", "sessions_spawn 出独立会话。子 Agent 互不 spawn，避免改到别人工作区、一条请求串进另一条业务。"),
        ("转述", "把数字、结论、出处转回来。去掉内部结构。对用户来说门口就是小慧。"),
    ]
    y = Inches(1.4)
    for h, b in items:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.35), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.2), Inches(2.2), Inches(0.95), h, 20, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(3.2), y + Inches(0.2), Inches(9.3), Inches(0.95), b, 16, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.5)
    tb(s, Inches(0.5), Inches(5.95), Inches(12.3), Inches(1.0), "不要把总控讲成多智能体编排。深度在「为什么切断协同」，不在「它们怎么商量」。", 16, True, NAVY)
    footer(s, 20, "总控")
    notes(s, "主控在联调前 40 秒带过。不要写 yaml。")


def integration(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "联调")
    title(s, "联调：一块屏，一份简历走完主链", "不是四人同时演。不是坏数据穿透全场。")
    box(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(2.2), WHITE, LINE)
    tb(s, Inches(0.75), Inches(1.55), Inches(11.8), Inches(0.4), "第一步 · 员工1 招聘主链", 16, True, ACCENT)
    tb(s, Inches(0.75), Inches(2.05), Inches(11.8), Inches(1.3), "已有 inbox md 则从解析起，走到脱敏 → 评分 → 数字回核 → 画像 Canvas。\n总控开口派单。回核在主控招聘主链上跑完。评委始终看这一块屏。", 16, False, INK)
    box(s, Inches(0.5), Inches(3.8), Inches(12.3), Inches(2.2), WHITE, LINE)
    tb(s, Inches(0.75), Inches(3.95), Inches(11.8), Inches(0.4), "第二步 · 停一秒，再问入职（可选）", 16, True, ACCENT)
    tb(s, Inches(0.75), Inches(4.45), Inches(11.8), Inches(1.3), "「员工1入职办到哪了」是新请求、新会话。女1 亮合同门。\n口述：录用是人定的，入职是另一单。行政若要带，再换飞书问年假——仍是同一块屏，顺序接，不并行。", 16, False, INK)
    tb(s, Inches(0.5), Inches(6.2), Inches(12.3), Inches(0.7), "女2、男1待命递文件，不另开窗口。培训绩效不接进这条链。完整性靠招聘主链跑通，不靠四个人一起动。", 15, False, TEAL)
    footer(s, 21, "联调")
    notes(s, "联调开始把 PPT 翻到这页停住，去操作。不要再念。")


def integration_check(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "联调")
    title(s, "联调完成时评委应看见", "对照评分表：技能在操作里，价值在三件事里。")
    rows = [
        ("技能 60", "现场写过/跑过闸；坏数据在技能段被拦；干净数据在联调走通。"),
        ("应用价值 10", "一个人的人事岗接得住初筛、催办、问制度；录用权仍在人。OPC 作背景，不吹市场。"),
        ("规范 / 素养", "口述跟屏幕同步。失败就说失败。不编花名册、不把夹具当真实人事。"),
        ("团队 10", "四人都有别人替不了的一刀。联调一人主操，三人有岗位。"),
    ]
    y = Inches(1.4)
    for h, b in rows:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.25), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.2), Inches(2.6), Inches(0.85), h, 16, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(3.5), y + Inches(0.2), Inches(9.0), Inches(0.85), b, 15, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.32)
    footer(s, 22, "联调")
    notes(s, "联调跑通后扫一眼。不要逐条念。进入收口。")


def secondary(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "收口")
    title(s, "开源底座上我们改了什么", "不要讲多智能体。指刚才评委已经看见的闸。")
    rows = [
        ("招聘", "评分前物理脱敏；抽取/数字两道回核；只给建议不淘汰。"),
        ("入职", "HCM 只读；合同/账号人批；pending-approvals 留住待批。"),
        ("人才发展", "培训只出点名的稿并强制讲师审核标注；绩效缺自评拒绝、不定级。"),
        ("行政", "飞书直达行政；身份闸在脚本；月报改宿主机脚本，模型不改数。专线互不 spawn。"),
    ]
    y = Inches(1.4)
    for h, b in rows:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.25), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.2), Inches(2.2), Inches(0.85), h, 18, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(3.2), y + Inches(0.2), Inches(9.3), Inches(0.85), b, 16, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.32)
    footer(s, 23, "收口")
    notes(s, "60 秒。每条对应刚才的操作。禁止再列 Skill 清单。")


def honest(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "收口")
    title(s, "边界说清楚", "诚实比凑绿更像工业级。")
    cols = [
        ("现在就是夹具", "花名册、简历、制度、反馈表目前是教学数据。不把员工1说成真实雇员。"),
        ("演了但不接联调", "培训、绩效技能段打开。不把八份稿当 SOP。360 不定级。"),
        ("调研还要补", "缺你们出镜的实拍、缺可出示的单位制度。有了再换知识库，不改架构。"),
    ]
    x = Inches(0.5)
    for h, b in cols:
        box(s, x, Inches(1.45), Inches(3.95), Inches(4.3), WHITE, LINE)
        tb(s, x + Inches(0.2), Inches(1.65), Inches(3.55), Inches(0.7), h, 18, True, NAVY)
        tb(s, x + Inches(0.2), Inches(2.5), Inches(3.55), Inches(2.9), b, 16, False, INK)
        x += Inches(4.15)
    tb(s, Inches(0.5), Inches(6.0), Inches(12.3), Inches(0.9), "产教融合：课程骨架 + 本仓闸门二次开发。软著、浏览量如有，只当真实性佐证，不作为加分口号。", 14, False, MUTED)
    footer(s, 24, "收口")
    notes(s, "40 秒。专家要真，先承认夹具和未调研，比编甲方安全。")


def outlook(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "收口")
    title(s, "接下来补什么", "不加新 Agent。把闸接到能指认的现场。")
    rows = [
        ("1", "调研：校内人事处或实习点，拍 1–2 张出镜照片。"),
        ("2", "材料：把 IMA 制度换成可出示的手册（脱敏后）。"),
        ("3", "排练：每人 10–15 分钟坏数据技能段 + 一块屏联调。自备分屏器。"),
        ("4", "不做：完整 SOP 补开发、训基础模型、把培训绩效接进联调主链。"),
    ]
    y = Inches(1.4)
    for n, b in rows:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.2), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.25), Inches(1.0), Inches(0.7), n, 22, True, ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(2.0), y + Inches(0.25), Inches(10.4), Inches(0.7), b, 18, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.28)
    footer(s, 25, "收口")
    notes(s, "30 秒。展望要可执行。不要五年市场。")


def thanks(prs):
    s = new_slide(prs)
    slide_bg(s, NAVY)
    rect(s, Inches(0), Inches(0), Inches(0.18), H, ACCENT)
    tb(s, Inches(0.7), Inches(2.2), Inches(12), Inches(0.9), "谢谢评委", 44, True, WHITE)
    tb(s, Inches(0.7), Inches(3.3), Inches(12), Inches(1.2), "脚本过关，人做决定。\n请提问。", 22, False, RGBColor(0xF0, 0xE6, 0xD8))
    tb(s, Inches(0.7), Inches(5.5), Inches(12), Inches(0.8), "小慧  ·  2026 云南省职业院校技能大赛  ·  人工智能赛道", 14, False, RGBColor(0xA8, 0xB4, 0xC8))
    notes(s, "停。把投屏留在操作窗口，方便追问时回放退出码。")


def backup_chain(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "附录 · 一般不讲")
    title(s, "附录：招聘八步口播", "被追问主链时用。技能段不要翻到附录。")
    lines = [
        "1 收信：未读附件落 inbox。岗位待确认不停车；同一人新稿必须重跑下游。",
        "2 转 MD：不调模型，不抽字段。",
        "3 形状闸：原文没有留空；红线键不进 JSON；推断打标。",
        "4 内容闸：硬提取对原文；过不了不能打分。",
        "5 脱敏：物理删敏感字段；残留性别退出码 1。",
        "6 评分：模型打八维；脚本定档；只建议不淘汰。",
        "7 数字回核：只认原简历；核不成不能说无虚构。",
        "8 画像：只渲染已有数据；红线字段零渲染；过校验再发布。",
    ]
    y = Inches(1.35)
    for line in lines:
        tb(s, Inches(0.6), y, Inches(12.1), Inches(0.62), line, 16, False, INK)
        y += Inches(0.62)
    footer(s, 27, "附录")
    notes(s, "默认跳过。导师深挖招聘再翻。")


def backup_ask(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "附录 · 一般不讲")
    title(s, "附录：可能被问", "短答。讲错不如说「这是边界」。")
    qa = [
        ("为什么互不协同？", "一次请求一条专线。筛完不会自动办入职。协同切断是为了人守门。"),
        ("为什么不训模型？", "我们改的是闸和校验。基础模型不是学生能诚实声称自研的。"),
        ("ASR / 乐奇？", "转写本机补测过的才说。乐奇媒体不转写。本场主链不依赖它们。"),
        ("完整 SOP 呢？", "联调主链是初筛、催办、问制度。培训绩效技能段另演两刀，不接进联调。编制、背调、社保仍由人办。"),
    ]
    y = Inches(1.4)
    for q, a in qa:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.25), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.12), Inches(11.8), Inches(0.4), q, 15, True, ACCENT)
        tb(s, Inches(0.75), y + Inches(0.55), Inches(11.8), Inches(0.55), a, 14, False, INK)
        y += Inches(1.32)
    footer(s, 28, "附录")
    notes(s, "默认跳过。")


def backup_score(prs):
    s = new_slide(prs)
    slide_bg(s)
    section_bar(s, "附录 · 一般不讲")
    title(s, "附录：评分表怎么对上现场", "给自己看，不要当汇报页念。")
    rows = [
        ("技能 60", "规范性 / 熟练度 / 难易 / 先进性：招聘三刀、入职门、行政闸、女2两刀的退出码；讲解只占 5 分。"),
        ("职业素养 10", "失败就停；不编数据；口述跟屏幕一致。"),
        ("应用价值 10", "OPC/人少人事岗的三件事；人在回路。"),
        ("团队合作 10", "四人四刀：招聘 / 人才发展 / 入职 / 行政；联调一块屏。"),
        ("创新创意 10", "开源上的闸，不是新名词。"),
    ]
    y = Inches(1.4)
    for h, b in rows:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.0), WHITE, LINE)
        tb(s, Inches(0.75), y + Inches(0.2), Inches(2.4), Inches(0.6), h, 16, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(3.3), y + Inches(0.2), Inches(9.2), Inches(0.6), b, 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.05)
    footer(s, 29, "附录")
    notes(s, "默认跳过。备赛对表用。")


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cover(prs)
    team(prs)
    national(prs)
    hr_pains(prs)
    background(prs)
    problems(prs)
    scheme(prs)
    doors(prs)
    hour_map(prs)
    recruit_chain(prs)
    desensitize(prs)
    scoring(prs)
    numbers(prs)
    training_gate(prs)
    performance_gate(prs)
    onboarding_gate(prs)
    onboarding_fail(prs)
    admin_access(prs)
    admin_policy(prs)
    router(prs)
    integration(prs)
    integration_check(prs)
    secondary(prs)
    honest(prs)
    outlook(prs)
    thanks(prs)
    backup_chain(prs)
    backup_ask(prs)
    backup_score(prs)
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
