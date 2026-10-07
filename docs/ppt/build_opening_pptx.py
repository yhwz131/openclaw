#!/usr/bin/env python3
"""小慧 · 赛场开场 5 分钟初版 PPT。在仓库 docs/ 生成 pptx。"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# 16:9
W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
NAVY2 = RGBColor(0x2C, 0x3E, 0x6B)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
CREAM = RGBColor(0xF7, 0xF4, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xD2, 0xC5)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)  # 土橙，钩子/本场
TEAL = RGBColor(0x2F, 0x6F, 0x6B)
DASH = RGBColor(0x8A, 0x86, 0x7C)
FONT = "Microsoft YaHei"

OUT = Path("/opt/openclaw-hr/docs/小慧-开场5分钟-初版.pptx")


def set_run(run, size, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def add_text(tf, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, space_after=6):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
    if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1:
        p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return p


def box(slide, l, t, w, h, fill=WHITE, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.solid()
    sh.line.color.rgb = line if line else fill
    sh.adjustments[0] = 0.08
    return sh


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    if line:
        sh.line.fill.solid()
        sh.line.color.rgb = line
    return sh


def tb(slide, l, t, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(l, t, w, h)
    sh.text_frame.word_wrap = True
    sh.text_frame.auto_size = None
    try:
        sh.text_frame.anchor = anchor
    except Exception:
        pass
    tf = sh.text_frame
    p = tf.paragraphs[0]
    p.alignment = align
    p.clear()
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return sh


def footer(slide, n, total=6):
    rect(slide, Inches(0), Inches(7.28), W, Inches(0.22), NAVY)
    tb(slide, Inches(0.4), Inches(7.28), Inches(8), Inches(0.22), "小慧 · 开场（初版）  ·  技能展示前", 9, False, WHITE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    tb(slide, Inches(10.5), Inches(7.28), Inches(2.4), Inches(0.22), f"{n} / {total}", 9, False, WHITE, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def slide_bg(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def new_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def cover(prs):
    s = new_slide(prs)
    slide_bg(s, NAVY)
    rect(s, Inches(0), Inches(0), Inches(0.18), H, ACCENT)
    tb(s, Inches(0.7), Inches(1.35), Inches(12), Inches(0.4), "2026 云南省职业院校技能大赛  ·  人工智能赛道  ·  高职组", 14, False, RGBColor(0xC8, 0xD0, 0xE0))
    tb(s, Inches(0.7), Inches(1.9), Inches(12), Inches(1.1), "小慧", 54, True, WHITE)
    tb(
        s,
        Inches(0.7),
        Inches(3.15),
        Inches(11.5),
        Inches(1.4),
        "让一个人的人事岗接得住三件事\n拍板权仍在人",
        28,
        False,
        RGBColor(0xF0, 0xE6, 0xD8),
    )
    tb(s, Inches(0.7), Inches(5.15), Inches(11), Inches(0.8), "HR 数字员工  ·  招聘初筛 · 入职催办 · 制度问答\n开场约 5 分钟，随后进入技能实操", 16, False, RGBColor(0xA8, 0xB4, 0xC8))
    tb(s, Inches(0.7), Inches(6.55), Inches(11), Inches(0.4), "四人参赛  ·  姓名待填", 12, False, RGBColor(0x8A, 0x96, 0xAA))
    notes(
        s,
        "5 秒过封面，不要念副标题以外的字。开口第一句直接进下一页："
        "评委老师好，我们是小慧团队。接下来 5 分钟只讲一件事：一个人的人事岗要同时接住什么，以及人必须留下哪几扇门。",
    )


def team(prs):
    s = new_slide(prs)
    slide_bg(s, CREAM)
    tb(s, Inches(0.5), Inches(0.28), Inches(12), Inches(0.45), "本场怎么看", 26, True, NAVY)
    tb(s, Inches(0.5), Inches(0.78), Inches(12), Inches(0.35), "四人各守一扇门。这一页给评委当目录，不是自我介绍。", 14, False, MUTED)

    headers = ["同学（姓名待填）", "本场职责", "技能段看什么", "联调时干什么"]
    rows = [
        ["主控", "总控分发 + 招聘脱敏与建议", "敏感字段物理删除；只给建议不淘汰", "一块屏串词，员工1走完招聘主链"],
        ["女2", "招聘回核（培训本场不展开）", "评分数字对原简历；串台则改回或停", "招聘产物在盘上，待命"],
        ["女1", "入职", "合同/账号两道门 halt，人批 yes/no", "招聘走完后，顺序接「入职办到哪了」"],
        ["男1", "行政", "非 HR 查工资拒绝；政策无命中不编", "飞书侧待命；不与招聘同屏抢镜头"],
    ]
    col_w = [Inches(2.2), Inches(3.3), Inches(3.7), Inches(3.2)]
    x0, y0, row_h = Inches(0.5), Inches(1.3), Inches(1.15)
    # header
    x = x0
    for i, h in enumerate(headers):
        r = rect(s, x, y0, col_w[i], Inches(0.42), NAVY)
        tb(s, x + Inches(0.08), y0, col_w[i] - Inches(0.1), Inches(0.42), h, 12, True, WHITE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
        x += col_w[i]
    for ri, row in enumerate(rows):
        y = y0 + Inches(0.42) + ri * row_h
        fill = WHITE if ri % 2 == 0 else RGBColor(0xEE, 0xE8, 0xDF)
        x = x0
        for i, cell in enumerate(row):
            rect(s, x, y, col_w[i], row_h, fill, LINE)
            weight = i == 0
            tb(s, x + Inches(0.1), y + Inches(0.12), col_w[i] - Inches(0.18), row_h - Inches(0.18), cell, 13, weight, INK)
            x += col_w[i]

    tb(s, Inches(0.5), Inches(6.55), Inches(12.2), Inches(0.55), "技能段一人一块屏；另外三人低头做联调用的干净结果，不抢镜头。联调始终同一块投屏。", 13, False, TEAL)
    footer(s, 2)
    notes(
        s,
        "20 秒。不要念名字简历。只指两句：技能段四人轮流当主角；联调是员工1把招聘主链跑通，入职如要带上就顺序接，评委始终看一块屏。"
        "女2的培训、系统里的绩效，本场不演。",
    )


def background(prs):
    s = new_slide(prs)
    slide_bg(s, CREAM)
    tb(s, Inches(0.5), Inches(0.28), Inches(12), Inches(0.45), "项目背景", 26, True, NAVY)
    tb(s, Inches(0.5), Inches(0.78), Inches(12), Inches(0.4), "政策给了方向，现场给了约束：人少，三件事同时来，又不能把录用权交给模型。", 14, False, MUTED)

    # left card OPC
    box(s, Inches(0.5), Inches(1.35), Inches(6.0), Inches(5.5), WHITE, LINE)
    rect(s, Inches(0.5), Inches(1.35), Inches(0.12), Inches(5.5), TEAL)
    tb(s, Inches(0.85), Inches(1.55), Inches(5.4), Inches(0.4), "方向", 12, True, TEAL)
    tb(s, Inches(0.85), Inches(2.0), Inches(5.4), Inches(1.6), "2026 年广东、深圳等地出台人工智能 OPC（一人公司）专项政策：极小团队用智能体接全链路业务。", 16, False, INK)
    tb(s, Inches(0.85), Inches(3.7), Inches(5.4), Inches(2.6), "这种团队往往没有人事部。招聘初筛、入职催办、制度问答，全压在一两个人身上。\n\n云南本地专项未引用。政策出处写粤深，不写本省。", 15, False, MUTED)

    # right card constraint
    box(s, Inches(6.8), Inches(1.35), Inches(6.0), Inches(5.5), WHITE, LINE)
    rect(s, Inches(6.8), Inches(1.35), Inches(0.12), Inches(5.5), ACCENT)
    tb(s, Inches(7.15), Inches(1.55), Inches(5.4), Inches(0.4), "约束", 12, True, ACCENT)
    tb(s, Inches(7.15), Inches(2.0), Inches(5.4), Inches(4.4), "《个人信息保护法》第 24 条：对个人权益有重大影响的决定，个人有权拒绝仅由自动化作出的结果。\n\n所以小慧让一个人接得住，这个人仍然拍板。\n\n一人 ≠ 无人。", 16, False, INK)

    footer(s, 3)
    notes(
        s,
        "约 50 秒。念完「一人不等于无人」停半拍。不要展开 OPC 定义、不要念文件全称。被问政策再答广东方案、深圳行动计划。不要说 OPD。",
    )


def problems(prs):
    s = new_slide(prs)
    slide_bg(s, CREAM)
    tb(s, Inches(0.5), Inches(0.28), Inches(12), Inches(0.45), "调研 · 发现的三件事", 26, True, NAVY)
    tb(s, Inches(0.5), Inches(0.78), Inches(12), Inches(0.35), "本场只验这三件。不是人力资源全流程。", 14, False, MUTED)

    items = [
        ("01", "简历多，不敢用 AI 刷人", "年龄、性别、婚育一看就有，模型看得见就无法证明没用过。"),
        ("02", "入职卡在合同和 IT 之间", "各部门各填各的，HR 当催办员，进度看不清，门必须人开。"),
        ("03", "员工把制度问题砸给 HR", "年假、报销、社保反复问；答错变成劳动争议，不能用常识硬编。"),
    ]
    x = Inches(0.5)
    for num, title, body in items:
        box(s, x, Inches(1.3), Inches(3.95), Inches(3.35), WHITE, LINE)
        tb(s, x + Inches(0.25), Inches(1.5), Inches(3.4), Inches(0.45), num, 20, True, ACCENT)
        tb(s, x + Inches(0.25), Inches(2.05), Inches(3.45), Inches(0.9), title, 18, True, NAVY)
        tb(s, x + Inches(0.25), Inches(3.0), Inches(3.45), Inches(1.4), body, 14, False, MUTED)
        x += Inches(4.15)

    # photo placeholder
    box(s, Inches(0.5), Inches(4.85), Inches(12.3), Inches(2.15), RGBColor(0xEE, 0xE8, 0xDF), LINE)
    tb(s, Inches(0.75), Inches(5.05), Inches(11.8), Inches(0.4), "实地调研照片（待补）", 16, True, NAVY)
    tb(
        s,
        Inches(0.75),
        Inches(5.5),
        Inches(11.8),
        Inches(1.2),
        "专家要求：1–2 张你们自己出镜的实拍，证明做过调研。没有照片就不要 P 图、不要写假公司名。\n有照片后替换本框。数据仍是教学夹具，口述不要说成已上线某企业。",
        14,
        False,
        MUTED,
    )
    footer(s, 4)
    notes(
        s,
        "约 50 秒。三件事各一句。照片框诚实说还没有。不要把培训 360 加进来。若已有甲方，把「待补」换成单位类型（校内人事处 / 实习点），仍不编花名册。",
    )


def scheme(prs):
    s = new_slide(prs)
    slide_bg(s, CREAM)
    tb(s, Inches(0.5), Inches(0.22), Inches(12), Inches(0.4), "整体方案", 26, True, NAVY)
    tb(s, Inches(0.5), Inches(0.62), Inches(12), Inches(0.32), "人面对的是小慧。系统按业务派专线。花名册在 HCM，制度在知识库，小慧只读不写。", 14, False, MUTED)

    # user + entries
    box(s, Inches(0.5), Inches(1.05), Inches(3.7), Inches(1.55), WHITE, LINE)
    tb(s, Inches(0.7), Inches(1.15), Inches(3.3), Inches(0.35), "门口", 12, True, TEAL)
    tb(s, Inches(0.7), Inches(1.5), Inches(3.3), Inches(0.95), "WebChat / CLI → 总控只分发、自己不答\n飞书 → 直达行政（不经总控）", 13, False, INK)

    box(s, Inches(4.4), Inches(1.05), Inches(4.4), Inches(1.55), WHITE, LINE)
    tb(s, Inches(4.6), Inches(1.15), Inches(4.0), Inches(0.35), "原则", 12, True, ACCENT)
    tb(s, Inches(4.6), Inches(1.5), Inches(4.0), Inches(0.95), "按岗位步骤拆成可校验脚本\n模型只负责写和抽，过不过关由脚本说", 13, False, INK)

    box(s, Inches(9.0), Inches(1.05), Inches(3.8), Inches(1.55), WHITE, LINE)
    tb(s, Inches(9.2), Inches(1.15), Inches(3.4), Inches(0.35), "人守的拍板", 12, True, NAVY)
    tb(s, Inches(9.2), Inches(1.5), Inches(3.4), Inches(0.95), "录用 / 合同 / 账号 / 他人工资\n一人接住，人仍拍板", 13, False, INK)

    # five lines
    lines = [
        (True, "招聘", "脱敏评分，只给建议"),
        (True, "入职", "只读进度，两道人批的门"),
        (False, "培训", "点名出草稿，讲师审了才用"),
        (False, "绩效", "360 聚合，不定级"),
        (True, "行政", "制度引原文，月报脚本算"),
    ]
    x = Inches(0.5)
    for live, name, desc in lines:
        fill = NAVY if live else RGBColor(0xE7, 0xE2, 0xD8)
        fg = WHITE if live else DASH
        tag = "本场演示" if live else "同一套纪律，本场不展开"
        box(s, x, Inches(2.85), Inches(2.4), Inches(2.85), fill if live else WHITE, NAVY if live else LINE)
        tb(s, x + Inches(0.12), Inches(2.98), Inches(2.15), Inches(0.35), tag, 10, True, ACCENT if live else DASH)
        tb(s, x + Inches(0.12), Inches(3.4), Inches(2.15), Inches(0.5), name, 22, True, fg if live else NAVY)
        tb(s, x + Inches(0.12), Inches(4.05), Inches(2.15), Inches(1.35), desc, 13, False, fg if live else MUTED)
        x += Inches(2.52)

    tb(s, Inches(0.5), Inches(5.9), Inches(12.3), Inches(1.1), "子 Agent 互不派活。筛完简历不会自动去办入职——录用是人定的，入职是另一单。\n培训、绩效灰着，是地图图例，不是本章正文。", 14, False, MUTED)
    footer(s, 5)
    notes(
        s,
        "约 70 秒。先指三张小卡片，再指下面五条：加粗三条是今天要演的；灰的两条被问再各二十秒。"
        "禁止在这一页讲多智能体协同、模型名、端口。问另外两条：培训点名出稿讲师审；绩效聚合不定级。",
    )


def doors(prs):
    s = new_slide(prs)
    slide_bg(s, CREAM)
    tb(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.4), "本场将打开的三扇门", 26, True, NAVY)
    tb(s, Inches(0.5), Inches(0.7), Inches(12), Inches(0.35), "不是亮点汇报。每扇门后面都有能跑出退出码的动作，对着评分表看。", 14, False, MUTED)

    doors_data = [
        ("门 1  ·  招聘", "敏感字段进评分", "物理删除。故意留「性别」，脚本退出码 1，不能打分。", "规范性 · 任务难易"),
        ("门 2  ·  入职", "合同 / 账号批准", "流程 halt 等人。人说 yes 才走，说 no 就 cancelled。", "规范性 · 熟练度"),
        ("门 3  ·  行政", "政策与他人工资", "无命中不编年假天数。非 HR 查工资，退出码 3。", "应用价值 · 规范性"),
    ]
    y = Inches(1.2)
    for title, hook, action, score in doors_data:
        box(s, Inches(0.5), y, Inches(12.3), Inches(1.7), WHITE, LINE)
        rect(s, Inches(0.5), y, Inches(0.14), Inches(1.7), ACCENT)
        tb(s, Inches(0.9), y + Inches(0.15), Inches(4.2), Inches(0.4), title, 16, True, NAVY)
        tb(s, Inches(5.2), y + Inches(0.18), Inches(4.5), Inches(0.35), hook, 16, True, ACCENT)
        tb(s, Inches(9.8), y + Inches(0.18), Inches(2.8), Inches(0.35), score, 12, False, TEAL, PP_ALIGN.RIGHT)
        tb(s, Inches(0.9), y + Inches(0.7), Inches(11.5), Inches(0.8), action, 16, False, INK)
        y += Inches(1.85)

    tb(s, Inches(0.5), Inches(6.55), Inches(12.3), Inches(0.5), "下一页开始动手：门 1。坏数据只在技能段用。联调用干净的员工1把招聘主链跑到画像。", 14, True, NAVY)
    footer(s, 6)
    notes(
        s,
        "约 50 秒。三扇门各读动作那一句。读完立刻切屏去脱敏实操，不要再翻回方案。"
        "若时间紧，门 2、门 3 只报名字，说入职和行政同学会打开。",
    )


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cover(prs)
    team(prs)
    background(prs)
    problems(prs)
    scheme(prs)
    doors(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
