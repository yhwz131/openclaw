"""v2.4：在 v2.2 上换全貌版封面，四个模块背景光随页推进、亮度不减、最后一页铺满。v2.2 不改。"""
from lxml import etree
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

COVER_SCRIPT = "/opt/openclaw-hr/docs/ppt/build_covers_overview.py"
DST = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v2.4.pptx"
MODULES = [
    ("招聘", [14, 15, 16, 17, 18], "00D8FF", 24000),
    ("入职", [19, 20, 21, 22, 23], "FFA421", 17000),
    ("行政", [24, 25, 26, 27, 28], "FF6B8A", 24000),
    ("绩效", [29, 30, 31, 32], "26D098", 24000),
]
STEPS = {
    5: ([0, 20, 40, 65, 100], [40, 55, 75, 95, 100]),
    4: ([0, 30, 60, 100], [40, 65, 90, 100]),
}


def q(tag):
    return "{http://schemas.openxmlformats.org/drawingml/2006/main}%s" % tag


code = open(COVER_SCRIPT, encoding="utf-8").read()
code = code.replace("KEEP = [14, 19, 24, 29]", "KEEP = list(range(1, 39))")
code = code.replace("p.save(DST)", "pass").replace('print("saved", DST, len(p.slides))', "")
ns = {}
exec(compile(code, COVER_SCRIPT, "exec"), ns)
p, S, C, E = ns["p"], ns["S"], ns["C"], ns["E"]


def stops(k, n, alpha):
    full, reach = STEPS[n][0][k], STEPS[n][1][k]
    if full >= 100:
        return [(0, alpha), (100, alpha)]
    pts = [(0, alpha)] + ([(full, alpha)] if full > 0 else [])
    pts += [((full + reach) / 2, alpha / 3), (reach, 0)]
    return pts + ([(100, 0)] if reach < 100 else [])


def light(s, color, pts):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    r.line.fill.background()
    r.shadow.inherit = False
    spPr = r._element.spPr
    for child in list(spPr):
        if child.tag in (q("solidFill"), q("gradFill"), q("noFill")):
            spPr.remove(child)
    grad = etree.Element(q("gradFill"), rotWithShape="1")
    lst = etree.SubElement(grad, q("gsLst"))
    for pos, a in pts:
        gs = etree.SubElement(lst, q("gs"), pos=str(int(pos * 1000)))
        clr = etree.SubElement(gs, q("srgbClr"), val=color)
        etree.SubElement(clr, q("alpha"), val=str(int(a)))
    etree.SubElement(grad, q("lin"), ang="0", scaled="0")
    spPr.insert(list(spPr).index(spPr.find(q("prstGeom"))) + 1, grad)
    bg = next(sh for sh in s.shapes if sh.shape_id == 2)._element
    bg.addnext(r._element)


def tint_grid(s, color, alpha):
    for sh in s.shapes:
        if sh.shape_type == 9 and 3 <= sh.shape_id <= 9:
            ln = sh._element.spPr.find(q("ln"))
            if ln is None:
                continue
            for child in list(ln):
                if child.tag in (q("solidFill"), q("gradFill"), q("noFill")):
                    ln.remove(child)
            fill = etree.Element(q("solidFill"))
            clr = etree.SubElement(fill, q("srgbClr"), val=color)
            etree.SubElement(clr, q("alpha"), val=str(alpha))
            ln.insert(0, fill)


def strip(s, color, w):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.18), Inches(w), Inches(6.32))
    r.fill.solid()
    r.fill.fore_color.rgb = C(color)
    r.line.fill.background()
    r.shadow.inherit = False


def recolor_label(s, color):
    for sh in s.shapes:
        l, t = sh.left / E, sh.top / E
        if sh.has_text_frame and abs(l - 0.48) < 0.05 and 1.2 < t < 1.4:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    r.font.color.rgb = C(color)
        elif sh.shape_type == 1 and abs(l - 0.48) < 0.05 and 2.05 < t < 2.15 and sh.height / E < 0.08 and sh.width / E < 0.8:
            sh.fill.solid()
            sh.fill.fore_color.rgb = C(color)


for name, pages, color, alpha in MODULES:
    for k, n in enumerate(pages):
        s = S[n]
        cover = k == 0
        if not cover:
            recolor_label(s, color)
        light(s, color, stops(k, len(pages), alpha))
        tint_grid(s, color, 30000 if cover else 22000)
        strip(s, color, 0.12 if cover else 0.08)

p.save(DST)
print("saved", DST, len(p.slides))
