#!/usr/bin/env python3
"""培训模块四页完善版：点名生成、确定性编排、双重验收与现场守卫。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import build_opening_template_pptx as base


ROOT = Path("/opt/openclaw-hr/docs")
OUT = ROOT / "PPT准备/培训模块技能演示-4页完善版.pptx"
TOTAL = 4

GREEN = RGBColor(0x26, 0xD0, 0x98)
GREEN_DARK = RGBColor(0x08, 0x50, 0x57)
PURPLE = RGBColor(0x8C, 0x7C, 0xFF)
PURPLE_DARK = RGBColor(0x20, 0x1B, 0x55)
RED = RGBColor(0xF1, 0x5A, 0x66)
RED_DARK = RGBColor(0x56, 0x17, 0x26)
CODE_BG = RGBColor(0x00, 0x11, 0x2C)


def connector(slide, x1, y1, x2, y2, color=base.CYAN, width=1.1, arrow=True):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def secondary_nav(slide) -> None:
    labels = ["总控", "招聘", "入职", "培训", "绩效", "行政"]
    base.rect(slide, Inches(0), Inches(0.72), base.W, Inches(0.45), base.NAVY, line=base.LINE)
    base.text(
        slide,
        Inches(0.38),
        Inches(0.79),
        Inches(0.9),
        Inches(0.27),
        "技能地图",
        10,
        bold=True,
        color=base.MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    for i, label in enumerate(labels):
        x = Inches(1.33 + i * 1.77)
        active = i == 3
        base.rect(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            base.CYAN if active else base.NAVY_2,
            line=base.CYAN if active else base.LINE,
            radius=True,
        )
        base.text(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            label,
            10.5,
            bold=active,
            color=base.NAVY if active else base.MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def header(slide, page: int, eyebrow: str, title: str, subtitle: str) -> None:
    base.add_navigation(slide, 2, page)
    secondary_nav(slide)
    base.text(slide, Inches(0.48), Inches(1.28), Inches(5.1), Inches(0.25), eyebrow, 10.5, bold=True, color=base.CYAN)
    base.text(slide, Inches(0.46), Inches(1.53), Inches(12.2), Inches(0.49), title, 22, bold=True, font=base.TITLE_FONT)
    base.rect(slide, Inches(0.47), Inches(2.08), Inches(0.75), Inches(0.05), base.CYAN)
    base.text(slide, Inches(1.4), Inches(1.99), Inches(11.1), Inches(0.29), subtitle, 11.2, color=base.MUTED)


def card(slide, x, y, w, h, title: str, body: str, *, accent=base.CYAN, fill=base.NAVY_2, body_size=11.5):
    base.rect(slide, x, y, w, h, fill, line=accent, radius=True, transparency=3)
    base.rect(slide, x, y, Inches(0.08), h, accent)
    base.text(slide, x + Inches(0.2), y + Inches(0.14), w - Inches(0.4), Inches(0.3), title, 13, bold=True, color=accent)
    base.text(slide, x + Inches(0.2), y + Inches(0.58), w - Inches(0.4), h - Inches(0.7), body, body_size, color=base.PALE, line_spacing=0.95)


def pill(slide, x, y, w, text_value: str, color=base.CYAN):
    base.rect(slide, x, y, w, Inches(0.34), base.NAVY_3, line=color, radius=True)
    base.text(
        slide,
        x,
        y,
        w,
        Inches(0.34),
        text_value,
        10,
        bold=True,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def slide_positioning(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        1,
        "04  TRAINING · NAMED DELIVERY",
        "培训 Agent：用户点名几份，系统就生成几份",
        "八类是ADDIE能力上限，不是默认产物；当前使用现有脱敏业务素材验证流程",
    )

    chain = [
        ("点名交付物", base.CYAN),
        ("读取业务素材", base.CYAN),
        ("ADDIE编排", base.ORANGE),
        ("指定物料", base.ORANGE),
        ("结构自检", GREEN),
        ("讲师审核", PURPLE),
    ]
    for i, (label, accent) in enumerate(chain):
        x = Inches(0.48 + i * 2.04)
        base.rect(slide, x, Inches(2.38), Inches(1.8), Inches(0.58), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x, Inches(2.38), Inches(1.8), Inches(0.58), label, 11.2, bold=True, color=accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(chain) - 1:
            connector(slide, x + Inches(1.8), Inches(2.67), x + Inches(2.0), Inches(2.67), accent)

    labels = [
        "01 需求分析",
        "02 课程设计",
        "03 PPT大纲",
        "04 讲师手册",
        "05 学员手册",
        "06 课后测试",
        "07 交付指南",
        "08 评估反馈",
    ]
    for i, label in enumerate(labels):
        row, col = divmod(i, 4)
        pill(slide, Inches(0.48 + col * 2.34), Inches(3.32 + row * 0.55), Inches(2.12), label, base.CYAN if row == 0 else GREEN)

    card(
        slide,
        Inches(9.98),
        Inches(3.28),
        Inches(2.86),
        Inches(1.42),
        "责任边界",
        "Agent｜生成草稿\n脚本｜范围与结构\n讲师｜内容与交付",
        accent=PURPLE,
        fill=PURPLE_DARK,
        body_size=11.2,
    )

    base.rect(slide, Inches(0.48), Inches(5.16), Inches(12.36), Inches(1.24), base.NAVY_3, line=base.ORANGE, radius=True)
    base.text(slide, Inches(0.72), Inches(5.35), Inches(3.5), Inches(0.3), "不是“自动生成八份材料”", 14, bold=True, color=base.ORANGE)
    base.text(
        slide,
        Inches(4.0),
        Inches(5.28),
        Inches(8.55),
        Inches(0.7),
        "用户决定生成范围｜缺依据不补写｜生成文件仍是AI辅助草稿\n未经HR与讲师审核，不进入正式交付",
        12.2,
        color=base.PALE,
        align=PP_ALIGN.CENTER,
    )
    base.add_notes(
        slide,
        "约50秒。培训Agent不是再让模型写一篇长文，而是按用户指定范围生成前后承接、能够验收的物料。"
        "八类只是能力上限，点名几份就生成几份。Agent生成草稿，脚本负责范围、顺序和结构，"
        "HR与讲师负责最终审核。当前使用现有脱敏业务素材，不表述为真实企业项目。",
    )


def slide_orchestration(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        2,
        "04  ADDIE · DETERMINISTIC ORCHESTRATION",
        "确定性编排：前阶段输出，成为后阶段指定输入",
        "确定的是调用顺序、文件依赖和失败条件，不是宣称LLM每次生成的正文完全相同",
    )

    stages = [
        ("A + D", "01需求分析\n02课程设计", base.CYAN),
        ("D", "03 PPT大纲", base.CYAN),
        ("D", "04讲师手册\n05学员手册", base.ORANGE),
        ("D + E", "06课后测试", base.ORANGE),
        ("I", "07交付指南", GREEN),
        ("E", "08评估反馈", GREEN),
    ]
    for i, (stage, body, accent) in enumerate(stages):
        x = Inches(0.48 + i * 2.04)
        base.rect(slide, x, Inches(2.4), Inches(1.8), Inches(1.06), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x + Inches(0.12), Inches(2.51), Inches(1.56), Inches(0.24), stage, 10.5, bold=True, color=accent, align=PP_ALIGN.CENTER)
        base.text(slide, x + Inches(0.12), Inches(2.84), Inches(1.56), Inches(0.42), body, 10.2, color=base.PALE, align=PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            connector(slide, x + Inches(1.8), Inches(2.93), x + Inches(2.0), Inches(2.93), accent)

    card(
        slide,
        Inches(0.48),
        Inches(3.84),
        Inches(3.92),
        Inches(1.52),
        "点名生成｜--only",
        "用户只要05、06，就只选择对应阶段\n缺必要上游立即停止，不扩成全套",
        accent=base.CYAN,
        body_size=11.5,
    )
    card(
        slide,
        Inches(4.7),
        Inches(3.84),
        Inches(3.92),
        Inches(1.52),
        "单一内容源｜防漂移",
        "04、05共同读取03 PPT大纲\n06读取01中的T1—T5与05学员手册",
        accent=base.ORANGE,
        body_size=11.5,
    )
    card(
        slide,
        Inches(8.92),
        Inches(3.84),
        Inches(3.92),
        Inches(1.52),
        "断点与重生成",
        "已有且长度达标的文件默认skip\nskip≠合格｜重生成必须显式--force",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=11.5,
    )

    base.rect(slide, Inches(0.48), Inches(5.68), Inches(12.36), Inches(0.82), base.NAVY_3, line=PURPLE, radius=True)
    base.text(
        slide,
        Inches(0.72),
        Inches(5.68),
        Inches(11.9),
        Inches(0.82),
        "7次阶段调用 → 8类物料｜阶段一一次产生01+02｜后续阶段只读取规定上游，不依赖会话记忆承接",
        11.8,
        bold=True,
        color=PURPLE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约1分钟。ADDIE过程由脚本编排：第一阶段一次生成需求分析和课程设计，后续按规定读取上游，"
        "七次调用形成八类物料。--only只选择点名阶段，缺上游就停止。04与05共同读取03，"
        "06读取01中的目标和05。确定的是顺序、输入和失败条件，不是正文完全确定。"
        "已有文件只按存在与长度skip，最终是否合格仍看质量门。",
    )


def slide_quality_gate(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        3,
        "04  QUALITY GATE · MACHINE + HUMAN",
        "生成不等于交付：结构由机器检查，内容由讲师终审",
        "把机器能稳定判断的规则写进脚本；语义准确性、业务可用性和最终发布仍由人确认",
    )

    card(
        slide,
        Inches(0.48),
        Inches(2.4),
        Inches(3.85),
        Inches(2.54),
        "落盘强制｜generate_pack.py",
        "【T1】→【对应培训目标T1】\n每份强制追加AI辅助标识\n缺必要上游输出halted_at\n只按点名范围写文件",
        accent=base.CYAN,
        body_size=11.5,
    )
    card(
        slide,
        Inches(4.74),
        Inches(2.4),
        Inches(3.85),
        Inches(2.54),
        "机器自检｜self_check.py",
        "文件存在且有效\nT1—T5是否齐全\n目标矩阵与题目覆盖\nAI辅助标识是否存在",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=11.5,
    )
    card(
        slide,
        Inches(9.0),
        Inches(2.4),
        Inches(3.84),
        Inches(2.54),
        "人工终审｜Agent + 讲师",
        "阶段语义是否真正承接\n制度数字是否忠于源材料\n讲师手册是否只有要点\n是否适合真实授课与发布",
        accent=PURPLE,
        fill=PURPLE_DARK,
        body_size=11.5,
    )

    y = Inches(5.28)
    labels = [
        ("源材料", "《劳动合同法》第九条", base.CYAN),
        ("培训目标", "T1｜理解入职合规原则", base.ORANGE),
        ("课后考核", "题目标注对应培训目标T1", GREEN),
        ("交付", "讲师复核后发布", PURPLE),
    ]
    for i, (top, body, accent) in enumerate(labels):
        x = Inches(0.48 + i * 3.15)
        base.rect(slide, x, y, Inches(2.9), Inches(0.92), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x + Inches(0.13), y + Inches(0.1), Inches(0.78), Inches(0.2), top, 9.5, bold=True, color=accent)
        base.text(slide, x + Inches(0.18), y + Inches(0.43), Inches(2.54), Inches(0.26), body, 9.8, color=base.PALE, align=PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            connector(slide, x + Inches(2.9), y + Inches(0.46), x + Inches(3.1), y + Inches(0.46), accent)

    base.add_notes(
        slide,
        "约1分10秒。generate_pack在落盘时展开题目标识、强制AI辅助标注并处理缺上游停止。"
        "self_check检查文件、T1到T5、目标矩阵、题目覆盖和AI标识。机器通过不等于内容正确；"
        "阶段是否真正承接、制度数字是否忠于源材料、材料是否可用于授课仍由Agent和讲师复核。"
        "以劳动合同法第九条为例，源依据进入T1，再进入对应T1的测试题，形成可追溯链。",
    )


def slide_live_demo(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        4,
        "04  NAMED GENERATION · LIVE DEMO",
        "只生成05、06：范围不扩张，缺上游就停止",
        "现场需求：围绕入职流程，只重新生成学员手册和课后测试题；本页保留全场唯一一段代码讲解",
    )

    base.rect(slide, Inches(0.48), Inches(2.4), Inches(7.2), Inches(0.82), CODE_BG, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.7), Inches(2.54), Inches(6.76), Inches(0.2), "主命令｜已有上游与产物，显式 --force 重生成点名文件", 10.4, bold=True, color=base.CYAN)
    code = "bash scripts/training-run.sh generate --theme 入职流程 --only 05,06 --force"
    base.text(slide, Inches(0.7), Inches(2.86), Inches(6.76), Inches(0.22), code, 8.8, color=base.PALE, font="Courier New")

    base.rect(slide, Inches(7.96), Inches(2.4), Inches(4.88), Inches(0.82), base.NAVY_3, line=base.ORANGE, radius=True)
    base.text(slide, Inches(8.17), Inches(2.55), Inches(4.45), Inches(0.2), "只运行 t4（05）+ t5（06）｜未点名文件不重写", 10.5, bold=True, color=base.ORANGE, align=PP_ALIGN.CENTER)

    cards = [
        ("01", "范围验证", "FILES只统计05、06\n未点名01—04、07、08不重写", base.CYAN),
        ("02", "质量验证", "检查T1—T5与AI标识\nself_check_exit决定RESULT", GREEN),
        ("03", "失败守卫", "缺03或01等必要上游\nhalted_at + RESULT: failure", RED),
    ]
    for i, (num, title, body, accent) in enumerate(cards):
        x = Inches(0.48 + i * 4.18)
        fill = RED_DARK if accent == RED else (GREEN_DARK if accent == GREEN else base.NAVY_2)
        base.rect(slide, x, Inches(3.55), Inches(3.9), Inches(1.45), fill, line=accent, radius=True)
        base.text(slide, x + Inches(0.18), Inches(3.72), Inches(0.45), Inches(0.26), num, 10, bold=True, color=accent)
        base.text(slide, x + Inches(0.65), Inches(3.68), Inches(2.95), Inches(0.3), title, 12.4, bold=True, color=base.PALE)
        base.text(slide, x + Inches(0.22), Inches(4.2), Inches(3.46), Inches(0.52), body, 10.6, color=base.PALE, align=PP_ALIGN.CENTER)

    base.rect(slide, Inches(0.48), Inches(5.3), Inches(8.3), Inches(1.2), CODE_BG, line=PURPLE, radius=True)
    base.text(slide, Inches(0.7), Inches(5.43), Inches(2.2), Inches(0.2), "唯一代码｜目标覆盖差集", 10.4, bold=True, color=PURPLE)
    tiny = (
        "found = set(T_GOAL.findall(text))\n"
        "expected = {str(i) for i in range(1, 6)}\n"
        "missing = sorted(expected - found)\n"
        "if missing: problems.append('测试题未覆盖培训目标T' + ','.join(missing))"
    )
    base.text(slide, Inches(0.7), Inches(5.74), Inches(7.84), Inches(0.58), tiny, 8.1, color=base.PALE, font="Courier New", line_spacing=0.75)

    base.rect(slide, Inches(9.08), Inches(5.3), Inches(3.76), Inches(1.2), base.NAVY_3, line=base.ORANGE, radius=True)
    base.text(slide, Inches(9.28), Inches(5.44), Inches(3.35), Inches(0.22), "现场备用", 10.5, bold=True, color=base.ORANGE, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(9.28), Inches(5.77), Inches(3.35), Inches(0.48), "模型响应慢：去掉 --force\n验证断点skip + partial自检\n明确说明“未重新生成”", 9.8, color=base.PALE, align=PP_ALIGN.CENTER)

    base.add_notes(
        slide,
        "约2分30秒。业务方只要05学员手册和06测试题，因此指定--only 05,06；"
        "现有目录已有产物，现场要重生成必须显式--force。脚本只运行t4和t5，未点名文件不重写，"
        "最后输出FILES、MISSING、self_check_exit和RESULT。打开测试题只展示T1到T5、AI标识，"
        "避开企业微信和假期余额等旧冲突内容。唯一代码讲解是从期望T1到T5减去已出现目标，"
        "有缺失就加入问题并使自检失败。缺必要上游时脚本halt，不擅自补全。"
        "若外部模型响应慢，去掉--force展示断点skip与自检，并明确未重新生成。",
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL
    slide_positioning(prs)
    slide_orchestration(prs)
    slide_quality_gate(prs)
    slide_live_demo(prs)
    prs.core_properties.title = "小慧 · 培训模块技能演示完善版"
    prs.core_properties.subject = "点名生成、确定性编排、素材边界、结构自检与讲师终审"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
