#!/usr/bin/env python3
"""生成“小慧”开头 + 主控/招聘技能演示 PPT 初版。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import build_opening_template_pptx as base


OUT = Path("/opt/openclaw-hr/docs/小慧-开头及总控招聘-模板初版.pptx")
TOTAL = 12
RED_DARK = RGBColor(0x56, 0x17, 0x26)
GREEN = RGBColor(0x28, 0xD3, 0x9A)
GREEN_DARK = RGBColor(0x08, 0x54, 0x58)
CODE_BG = RGBColor(0x00, 0x11, 0x2C)
GRAY_BLUE = RGBColor(0x5E, 0x86, 0xA3)


def connector(slide, x1, y1, x2, y2, color=base.CYAN, width=1.2, arrow=True):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    sh.line.color.rgb = color
    sh.line.width = Pt(width)
    if arrow:
        sh.line.end_arrowhead = True
    return sh


def secondary_nav(slide, active: int) -> None:
    labels = ["总控", "招聘", "入职", "培训", "绩效", "行政"]
    base.rect(slide, Inches(0), Inches(0.72), base.W, Inches(0.45), base.NAVY, line=base.LINE)
    base.text(
        slide,
        Inches(0.38),
        Inches(0.79),
        Inches(0.9),
        Inches(0.27),
        "技能地图",
        10,
        bold=True,
        color=base.MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    for i, label in enumerate(labels):
        x = Inches(1.33 + i * 1.77)
        base.rect(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            base.CYAN if i == active else base.NAVY_2,
            line=base.CYAN if i == active else base.LINE,
            radius=True,
        )
        base.text(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            label,
            10.5,
            bold=i == active,
            color=base.NAVY if i == active else base.MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def skill_header(slide, page: int, active: int, eyebrow: str, title: str, subtitle: str) -> None:
    base.add_navigation(slide, 2, page)
    secondary_nav(slide, active)
    base.text(slide, Inches(0.48), Inches(1.28), Inches(3.0), Inches(0.25), eyebrow, 10.5, bold=True, color=base.CYAN)
    base.text(slide, Inches(0.46), Inches(1.53), Inches(12.2), Inches(0.49), title, 23, bold=True, font=base.TITLE_FONT)
    base.rect(slide, Inches(0.47), Inches(2.08), Inches(0.75), Inches(0.05), base.CYAN)
    base.text(slide, Inches(1.4), Inches(1.99), Inches(11.1), Inches(0.29), subtitle, 11.5, color=base.MUTED)


def layer_label(slide, y, index, label):
    base.rect(slide, Inches(0.48), y, Inches(1.05), Inches(0.48), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.48), y, Inches(1.05), Inches(0.48), index, 10, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    base.text(slide, Inches(1.7), y, Inches(1.4), Inches(0.48), label, 13, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def small_card(slide, x, y, w, h, heading, body, *, accent=base.CYAN, fill=base.NAVY_2, heading_size=13, body_size=10.5):
    base.rect(slide, x, y, w, h, fill, line=base.LINE, radius=True, transparency=3)
    base.rect(slide, x, y, Inches(0.07), h, accent)
    base.text(slide, x + Inches(0.16), y + Inches(0.11), w - Inches(0.28), Inches(0.28), heading, heading_size, bold=True, color=accent)
    base.text(slide, x + Inches(0.16), y + Inches(0.48), w - Inches(0.28), h - Inches(0.58), body, body_size, color=base.PALE)


def architecture(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    base.add_navigation(slide, 1, 5)
    base.add_title(
        slide,
        "02  TECHNICAL ARCHITECTURE",
        "完整架构：入口、Router、Agent与业务数据",
        "上一页回答“怎样分工”，这一页回答“系统怎样连接”",
    )

    # 四层结构；避免与上一页“1+5模块”重复。
    ys = [Inches(2.15), Inches(3.08), Inches(4.03), Inches(5.15)]
    layer_label(slide, ys[0], "L1", "任务入口")
    layer_label(slide, ys[1], "L2", "调度中枢")
    layer_label(slide, ys[2], "L3", "业务专线")
    layer_label(slide, ys[3], "L4", "工具与数据")

    inputs = [("WebChat / CLI", 3.05), ("企业邮箱", 5.75), ("飞书入口", 8.45)]
    for label, x in inputs:
        base.rect(slide, Inches(x), ys[0], Inches(2.15), Inches(0.48), base.NAVY_2, line=base.LINE, radius=True)
        base.text(slide, Inches(x), ys[0], Inches(2.15), Inches(0.48), label, 12.5, bold=True, color=base.PALE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    base.rect(slide, Inches(4.0), ys[1] - Inches(0.05), Inches(5.45), Inches(0.59), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(4.15), ys[1] + Inches(0.02), Inches(1.6), Inches(0.28), "OpenClaw Router", 14.5, bold=True, color=base.CYAN)
    base.text(slide, Inches(5.75), ys[1] + Inches(0.02), Inches(3.45), Inches(0.28), "识别意图｜隔离会话｜派发任务｜回传结果", 11.5, color=base.PALE, align=PP_ALIGN.CENTER)

    names = ["招聘", "入职", "培训", "绩效", "行政"]
    for i, name in enumerate(names):
        x = Inches(2.42 + i * 2.05)
        base.rect(slide, x, ys[2] - Inches(0.02), Inches(1.7), Inches(0.53), base.NAVY_2, line=base.LINE, radius=True)
        base.text(slide, x, ys[2] - Inches(0.02), Inches(1.7), Inches(0.53), f"{name} Agent", 12.5, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        connector(slide, Inches(6.72), ys[1] + Inches(0.54), x + Inches(0.85), ys[2] - Inches(0.02), base.LINE, 0.9, False)

    resources = [
        ("机制脚本", "校验·约束"),
        ("人才库", "向量检索"),
        ("HCM", "进度只读"),
        ("制度库", "原文问答"),
        ("报表产物", "审计留痕"),
    ]
    for i, (name, sub) in enumerate(resources):
        x = Inches(2.42 + i * 2.05)
        base.rect(slide, x, ys[3] - Inches(0.03), Inches(1.7), Inches(0.72), base.NAVY_2, line=base.LINE, radius=True)
        base.text(slide, x, ys[3] + Inches(0.04), Inches(1.7), Inches(0.25), name, 11.5, bold=True, align=PP_ALIGN.CENTER)
        base.text(slide, x, ys[3] + Inches(0.34), Inches(1.7), Inches(0.22), sub, 9.5, color=base.MUTED, align=PP_ALIGN.CENTER)

    base.rect(slide, Inches(0.48), Inches(6.2), Inches(12.36), Inches(0.63), base.NAVY_2, line=base.ORANGE, radius=True)
    base.text(slide, Inches(0.72), Inches(6.34), Inches(2.0), Inches(0.25), "贯穿全链路的人机边界", 10.5, bold=True, color=base.ORANGE)
    base.text(slide, Inches(2.82), Inches(6.28), Inches(9.55), Inches(0.35), "Agent负责理解与生成｜脚本负责校验与约束｜HR负责授权、复核与最终决策", 13, bold=True, color=base.PALE)
    base.add_notes(
        slide,
        "本页约45秒。上一页是业务分工，本页是技术连接。WebChat和CLI进入Router；企业邮箱进入招聘收件链；"
        "飞书可直达行政入口。Router集中调度，五条业务专线隔离运行，再调用各自脚本、人才库、HCM和制度库。"
        "贯穿全链路的原则是：Agent负责理解与生成，脚本负责校验与约束，HR负责最终决策。",
    )


def router_slide(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    skill_header(slide, 6, 0, "03  CONTROL AGENT", "总控只做三件事：判断、派发、转述", "简单带过，现场只验证任务确实进入正确业务模块")

    steps = [
        ("01", "判断意图", "识别HR要做的是招聘、入职、培训、绩效还是行政"),
        ("02", "派发专线", "创建隔离会话，把任务交给对应子Agent处理"),
        ("03", "返回结果", "接收业务产物，保留数字、结论和出处后转给HR"),
    ]
    for i, (num, heading, body) in enumerate(steps):
        x = Inches(0.55 + i * 4.23)
        shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, Inches(2.45), Inches(3.88), Inches(1.5))
        base.set_fill(shape, base.NAVY_3 if i == 1 else base.NAVY_2)
        shape.line.color.rgb = base.CYAN if i == 1 else base.LINE
        shape.line.width = Pt(1.2)
        base.text(slide, x + Inches(0.18), Inches(2.61), Inches(0.54), Inches(0.3), num, 12, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(0.68), Inches(2.57), Inches(2.55), Inches(0.34), heading, 17, bold=True)
        base.text(slide, x + Inches(0.28), Inches(3.08), Inches(2.92), Inches(0.6), body, 11.5, color=base.PALE)

    small_card(
        slide,
        Inches(0.55),
        Inches(4.28),
        Inches(5.85),
        Inches(1.72),
        "总控的边界",
        "不直接评分、不回答制度、不替子Agent做决定。\n意图不明确时追问HR，不凭猜测跨模块执行。",
        accent=base.ORANGE,
        body_size=13,
    )
    small_card(
        slide,
        Inches(6.65),
        Inches(4.28),
        Inches(6.13),
        Inches(1.72),
        "现场演示（2—3分钟）",
        "输入：“请筛选后端开发岗位简历并生成候选人画像”\n观察：总控识别为招聘任务 → 派发招聘Agent → 返回会话结果",
        accent=base.CYAN,
        body_size=13,
    )
    base.rect(slide, Inches(0.55), Inches(6.23), Inches(12.23), Inches(0.54), base.NAVY_2, line=base.LINE, radius=True)
    base.text(slide, Inches(0.75), Inches(6.35), Inches(11.85), Inches(0.25), "判断不准就问人｜子Agent互不派活｜筛完简历不会自动办理入职", 12, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)
    base.add_notes(
        slide,
        "总控控制在2—3分钟。口播：总控不负责具体业务，只判断、派发和转述。"
        "现场输入一条招聘任务，看到任务进入招聘Agent即可，不在这里展开底层配置。"
        "最后一句：筛完简历不会自动办理入职，因为录用由HR决定，入职是另一张任务单。",
    )


def chain_step(slide, x, y, w, label, *, hot=False, small=False):
    base.rect(slide, x, y, w, Inches(0.52 if not small else 0.43), base.NAVY_3 if hot else base.NAVY_2, line=base.CYAN if hot else base.LINE, radius=True)
    base.text(
        slide,
        x,
        y,
        w,
        Inches(0.52 if not small else 0.43),
        label,
        10.5 if not small else 9.5,
        bold=hot,
        color=base.CYAN if hot else base.PALE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def recruitment_map(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    skill_header(slide, 7, 1, "03  RECRUITMENT AGENT", "招聘Agent：一条主链，两条支链", "三条链共用简历解析能力，本场重点展示主链的三个机制")

    # 主链
    base.text(slide, Inches(0.52), Inches(2.42), Inches(1.58), Inches(0.3), "主链｜简历筛选与画像", 12.5, bold=True, color=base.CYAN)
    main = ["简历收取", "文档转换", "结构化解析", "敏感字段脱敏", "八维评分", "数字回核", "候选人画像"]
    widths = [1.35, 1.35, 1.48, 1.62, 1.36, 1.42, 1.45]
    x = 2.18
    for i, (label, width) in enumerate(zip(main, widths)):
        chain_step(slide, Inches(x), Inches(2.33), Inches(width), label, hot=i in (3, 4, 5))
        if i < len(main) - 1:
            connector(slide, Inches(x + width), Inches(2.59), Inches(x + width + 0.16), Inches(2.59), base.CYAN if i in (2, 3, 4, 5) else base.LINE, 1.0)
        x += width + 0.19

    # 人才库链
    base.text(slide, Inches(0.52), Inches(3.38), Inches(1.58), Inches(0.3), "支链A｜人才库检索", 12.5, bold=True, color=base.MUTED)
    talent = ["历史简历", "真实性核对", "向量化入库", "余弦检索", "推荐理由"]
    x = 2.18
    for i, label in enumerate(talent):
        chain_step(slide, Inches(x), Inches(3.29), Inches(1.62), label, small=True)
        if i < len(talent) - 1:
            connector(slide, Inches(x + 1.62), Inches(3.5), Inches(x + 1.83), Inches(3.5), base.LINE, 0.9)
        x += 1.86

    # 面试链
    base.text(slide, Inches(0.52), Inches(4.28), Inches(1.58), Inches(0.3), "支链B｜面试纪要", 12.5, bold=True, color=base.MUTED)
    interview = ["面试录音", "语音转写", "说话人分离", "转写校验", "结构化纪要"]
    x = 2.18
    for i, label in enumerate(interview):
        chain_step(slide, Inches(x), Inches(4.19), Inches(1.62), label, small=True)
        if i < len(interview) - 1:
            connector(slide, Inches(x + 1.62), Inches(4.4), Inches(x + 1.83), Inches(4.4), base.LINE, 0.9)
        x += 1.86

    small_card(
        slide,
        Inches(0.52),
        Inches(5.12),
        Inches(3.88),
        Inches(1.42),
        "机制1｜评分前物理脱敏",
        "敏感字段不进入模型上下文",
        accent=base.CYAN,
        body_size=12,
    )
    small_card(
        slide,
        Inches(4.72),
        Inches(5.12),
        Inches(3.88),
        Inches(1.42),
        "机制2｜八维评分边界",
        "Agent打分，脚本定档，只给建议",
        accent=base.CYAN,
        body_size=12,
    )
    small_card(
        slide,
        Inches(8.92),
        Inches(5.12),
        Inches(3.88),
        Inches(1.42),
        "机制3｜评分数字回核",
        "只以候选人原简历为真实性基准",
        accent=base.CYAN,
        body_size=12,
    )
    base.add_notes(
        slide,
        "约60秒。主链用于简历初筛和候选人画像；人才库链用于历史人才相似度检索；"
        "面试链用于转写、说话人分离和纪要。今天重点展示主链中三个深色节点：物理脱敏、八维评分边界、数字回核。",
    )


def mechanism_problem(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    skill_header(slide, 8, 1, "03  MECHANISM DESIGN", "从概率性问题，沉淀为可重复验证的机制", "不是复现模型偶发错误，而是现场验证机制能否稳定发现并处理")

    small_card(
        slide,
        Inches(0.52),
        Inches(2.43),
        Inches(5.9),
        Inches(1.55),
        "问题A｜敏感字段进入评分上下文",
        "年龄、性别、婚育、户籍、民族、院校层次等字段一旦被模型看见，便无法排除其对评分的影响。",
        accent=base.RED,
        fill=RED_DARK,
        body_size=12.5,
    )
    small_card(
        slide,
        Inches(6.68),
        Inches(2.43),
        Inches(6.1),
        Inches(1.55),
        "问题B｜简历与JD发生上下文串台",
        "开发中曾出现：候选人原简历写12.8万TPS，评分却引用了JD里的58万笔/秒。",
        accent=base.RED,
        fill=RED_DARK,
        body_size=12.5,
    )

    base.rect(slide, Inches(0.52), Inches(4.27), Inches(12.26), Inches(1.02), base.NAVY_3, line=base.CYAN, radius=True)
    principle = [("Agent", "理解与生成"), ("脚本", "校验与约束"), ("HR", "授权与最终决策")]
    for i, (role, task) in enumerate(principle):
        x = Inches(0.82 + i * 4.05)
        base.text(slide, x, Inches(4.43), Inches(1.15), Inches(0.3), role, 15, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(1.12), Inches(4.42), Inches(2.35), Inches(0.3), task, 14, bold=True)
        if i < 2:
            base.text(slide, x + Inches(3.42), Inches(4.4), Inches(0.42), Inches(0.3), "→", 18, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)

    base.rect(slide, Inches(0.52), Inches(5.57), Inches(12.26), Inches(1.18), base.NAVY_2, line=base.LINE, radius=True)
    base.text(slide, Inches(0.76), Inches(5.73), Inches(2.22), Inches(0.31), "沉淀为回归测试用例", 12, bold=True, color=base.ORANGE)
    base.text(
        slide,
        Inches(2.95),
        Inches(5.69),
        Inches(9.5),
        Inches(0.7),
        "开发过程中捕捉问题样本 → 固化为回归测试用例 → 现场重新运行机制 → 观察删除清单、退出码与修正快照",
        13,
        color=base.PALE,
    )
    base.add_notes(
        slide,
        "本页用于优化措辞。口播：这两类问题具有概率性，现场重新调用模型不一定重复出现。"
        "因此我们将开发过程中捕捉到的问题样本沉淀为回归测试用例，现场运行的是机制本身，"
        "看它能否稳定给出删除清单、退出码和修正快照。不要说“提前准备的错误文件”。",
    )


def desensitization(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    skill_header(slide, 9, 1, "03  GATE 1 · DESENSITIZE", "机制一：评分前物理删除敏感字段", "不是提示模型“不要使用”，而是让评分模型根本拿不到")

    small_card(
        slide,
        Inches(0.52),
        Inches(2.43),
        Inches(3.85),
        Inches(2.66),
        "输入｜带风险字段的JSON",
        "gender: 女\nage: 29\nmarital_status: 未婚\neducation.school_tier: 985\npolitical_status: 中共党员",
        accent=base.RED,
        fill=RED_DARK,
        body_size=12.5,
    )
    small_card(
        slide,
        Inches(4.63),
        Inches(2.43),
        Inches(4.0),
        Inches(2.66),
        "脚本｜递归物理删除",
        "嵌套字段也要删掉。\n只扫顶层会漏 school_tier。\n原始 parsed 不动，只给评分一份 clean 副本。",
        accent=base.CYAN,
        body_size=13,
    )
    small_card(
        slide,
        Inches(8.89),
        Inches(2.43),
        Inches(3.89),
        Inches(2.66),
        "输出｜clean JSON + 清单",
        "已物理删除 5 个字段\n评分输入剩余敏感字段 = 0\n✓ 自检通过\nexit code = 0",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=13,
    )

    base.rect(slide, Inches(0.52), Inches(5.36), Inches(12.26), Inches(1.35), base.NAVY_2, line=base.CYAN, radius=True)
    actions = ["展示敏感字段", "运行脱敏脚本", "查看删除清单", "确认自检与退出码"]
    for i, action in enumerate(actions):
        x = Inches(0.9 + i * 3.0)
        base.text(slide, x, Inches(5.53), Inches(0.34), Inches(0.3), f"{i+1:02}", 11, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(0.38), Inches(5.48), Inches(2.2), Inches(0.42), action, 13, bold=True, color=base.PALE)
        if i < 3:
            connector(slide, x + Inches(2.55), Inches(5.7), x + Inches(2.8), Inches(5.7), base.LINE, 0.9)
    base.text(slide, Inches(0.75), Inches(6.2), Inches(11.75), Inches(0.27), "残留任一疑似敏感键 → exit code 1 → 不允许进入评分", 11.5, bold=True, color=base.ORANGE, align=PP_ALIGN.CENTER)
    base.add_notes(
        slide,
        "本页不手写。约2分钟。展示嵌套在education里的school_tier，运行脚本。"
        "评委看见删除清单、剩余敏感字段为0和退出码即可。"
        "口述：原始parsed数据不动，只给评分模型clean副本。手写技能在下一道回核。",
    )


def scoring(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    skill_header(slide, 10, 1, "03  GATE 2 · SCORING", "机制二：八维评分，但只输出辅助建议", "模型负责匹配判断；维度、评级切分和输出边界由脚本固定")

    dims = ["专业技能", "经验匹配", "逻辑能力", "学习力", "沟通能力", "协作能力", "抗压能力", "价值观"]
    base.text(slide, Inches(0.52), Inches(2.42), Inches(3.55), Inches(0.3), "八个固定维度｜每维1—5分", 12.5, bold=True, color=base.CYAN)
    for i, dim in enumerate(dims):
        col, row = i % 2, i // 2
        x = Inches(0.52 + col * 1.83)
        y = Inches(2.84 + row * 0.64)
        base.rect(slide, x, y, Inches(1.64), Inches(0.47), base.NAVY_2, line=base.LINE, radius=True)
        base.text(slide, x, y, Inches(1.64), Inches(0.47), dim, 11.5, bold=i < 2, color=base.CYAN if i < 2 else base.PALE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    small_card(
        slide,
        Inches(4.4),
        Inches(2.43),
        Inches(4.12),
        Inches(2.95),
        "脚本固定评级口径",
        "综合分 ≥ 85  →  A 强烈推荐\n综合分 ≥ 75  →  B 推荐\n综合分 ≥ 60  →  C 待定\n其他             →  D 不匹配\n\n例：65分固定为C，模型不能临场改成B。",
        accent=base.CYAN,
        body_size=13,
    )
    small_card(
        slide,
        Inches(8.78),
        Inches(2.43),
        Inches(4.0),
        Inches(2.95),
        "决策边界",
        "允许：建议进入面试、建议人工复核\n禁止：自动录用、自动淘汰\n\n硬伤只记录真实性疑点；薪资、城市、跳槽频率只能作为顾虑。",
        accent=base.ORANGE,
        body_size=13,
    )
    base.rect(slide, Inches(0.52), Inches(5.72), Inches(12.26), Inches(0.98), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.75), Inches(5.89), Inches(2.2), Inches(0.28), "现场动作", 12, bold=True, color=base.CYAN)
    base.text(slide, Inches(2.25), Inches(5.84), Inches(10.05), Inches(0.42), "使用clean JSON正常评分 → 展示八维分、综合分和脚本定档 → 强调“建议”而非“决定”", 13, bold=True, color=base.PALE)
    base.add_notes(
        slide,
        "约2—3分钟。使用上一页生成的clean JSON跑评分。强调八个维度和评级切分由脚本固定，"
        "模型只完成匹配判断。评分产出只能写建议进入面试或人工复核，不能自动录用和淘汰。"
        "不要说这是自研评分算法，可以说这是对评分机制和决策边界的二次开发。",
    )


def number_audit(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    skill_header(slide, 11, 1, "03  LIVE CODING · NUMBER AUDIT", "机制三：数字只在原简历里核对", "防止模型把JD里的场景数字安到候选人头上")

    small_card(
        slide,
        Inches(0.52),
        Inches(2.43),
        Inches(3.25),
        Inches(2.66),
        "输入｜错误评分文本",
        "“具备58万笔/秒保障经验”\n这是JD里的场景数字\n候选人原简历没有这个数",
        accent=base.RED,
        fill=RED_DARK,
        body_size=13,
    )

    base.rect(slide, Inches(4.0), Inches(2.43), Inches(5.2), Inches(2.66), CODE_BG, line=base.LINE, radius=True)
    base.text(slide, Inches(4.22), Inches(2.58), Inches(4.76), Inches(0.3), "现场手写 verifyInResume", 11.5, bold=True, color=base.CYAN)
    code = (
        "const t = norm(token);\n"
        "if (resumeNorm.includes(t))\n"
        "  return true;\n"
        "const n = t.match(/\\d+(?:\\.\\d+)?/);\n"
        "if (n && resumeNorm.includes(n[0]))\n"
        "  return null;\n"
        "return false;"
    )
    base.text(slide, Inches(4.22), Inches(2.92), Inches(4.76), Inches(1.95), code, 13, color=base.PALE, font="Courier New", line_spacing=0.92)

    small_card(
        slide,
        Inches(9.44),
        Inches(2.43),
        Inches(3.34),
        Inches(2.66),
        "基准｜候选人原简历",
        "大促峰值 12.8万 TPS\n比对未命中 → 进入修正\n58万笔/秒 → 12.8万 TPS\n综合分不改，前后快照留痕",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=13,
    )

    base.rect(slide, Inches(0.52), Inches(5.36), Inches(12.26), Inches(1.35), base.NAVY_2, line=base.CYAN, radius=True)
    actions = ["展示58万评分", "手写比对函数", "运行回核脚本", "查看修正快照", "确认退出码"]
    for i, action in enumerate(actions):
        x = Inches(0.75 + i * 2.42)
        base.text(slide, x, Inches(5.53), Inches(0.34), Inches(0.3), f"{i+1:02}", 11, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(0.38), Inches(5.48), Inches(1.72), Inches(0.42), action, 11.5, bold=True, color=base.PALE)
        if i < 4:
            connector(slide, x + Inches(2.08), Inches(5.7), x + Inches(2.28), Inches(5.7), base.LINE, 0.9)
    base.text(slide, Inches(0.75), Inches(6.2), Inches(11.75), Inches(0.27), "脚本不接收JD。判定失败 → exit code 1 → 不能当成“无虚构”", 11.5, bold=True, color=base.ORANGE, align=PP_ALIGN.CENTER)
    base.add_notes(
        slide,
        "现场约4—6分钟。手写只写比对：数字在不在原简历里。"
        "先打开开发中保存的58万错误评分，对照原简历12.8万TPS，写完verifyInResume立刻跑脚本。"
        "评委要看见未命中、修正快照和退出码。"
        "抽取是脚本里的正则；判定和替换仍调用模型，失败则退出码1，不能当成无虚构。"
        "脚本接口不接JD。不要说整条回核都是手写的。",
    )


def close_loop(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    skill_header(slide, 12, 1, "03  RECRUITMENT · CLOSED LOOP", "招聘实操闭环：从问题样本到清洁画像", "坏样本验证机制，清洁数据跑通业务主链")

    steps = [
        ("01", "查看风险字段"),
        ("02", "运行脱敏脚本"),
        ("03", "clean JSON评分"),
        ("04", "展示58万评分"),
        ("05", "手写回核比对"),
        ("06", "运行数字回核"),
        ("07", "查看修正快照"),
    ]
    for i, (num, label) in enumerate(steps):
        x = Inches(0.47 + i * 1.82)
        base.rect(slide, x, Inches(2.46), Inches(1.55), Inches(0.92), base.NAVY_3 if i in (4, 5, 6) else base.NAVY_2, line=base.CYAN if i in (4, 5, 6) else base.LINE, radius=True)
        base.text(slide, x + Inches(0.1), Inches(2.56), Inches(0.35), Inches(0.25), num, 10, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(0.12), Inches(2.88), Inches(1.31), Inches(0.3), label, 10.5, bold=True, color=base.PALE, align=PP_ALIGN.CENTER)
        if i < 6:
            connector(slide, x + Inches(1.55), Inches(2.92), x + Inches(1.78), Inches(2.92), base.LINE, 0.9)

    base.text(slide, Inches(0.52), Inches(3.74), Inches(5.7), Inches(0.3), "机制加入前｜风险不可见", 13, bold=True, color=base.RED)
    base.text(slide, Inches(6.68), Inches(3.74), Inches(5.7), Inches(0.3), "机制加入后｜结果可校验、可追溯", 13, bold=True, color=GREEN)
    before = [
        "敏感字段进入模型上下文",
        "JD数字可能串入候选人画像",
        "只能提醒模型“请不要出错”",
        "HR难以核对结果产生过程",
    ]
    after = [
        "评分输入剩余敏感字段 = 0",
        "数字只回溯候选人原简历",
        "失败有退出码，不静默放行",
        "删除清单与修正快照可留痕",
    ]
    for i, value in enumerate(before):
        y = Inches(4.17 + i * 0.52)
        base.rect(slide, Inches(0.52), y, Inches(5.7), Inches(0.39), RED_DARK, line=base.RED, radius=True)
        base.text(slide, Inches(0.72), y, Inches(5.3), Inches(0.39), f"×  {value}", 11.5, color=base.PALE, anchor=MSO_ANCHOR.MIDDLE)
    for i, value in enumerate(after):
        y = Inches(4.17 + i * 0.52)
        base.rect(slide, Inches(6.68), y, Inches(6.1), Inches(0.39), GREEN_DARK, line=GREEN, radius=True)
        base.text(slide, Inches(6.88), y, Inches(5.7), Inches(0.39), f"✓  {value}", 11.5, color=base.PALE, anchor=MSO_ANCHOR.MIDDLE)

    base.rect(slide, Inches(0.52), Inches(6.43), Inches(12.26), Inches(0.39), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.72), Inches(6.43), Inches(11.86), Inches(0.39), "招聘模块演示完成  →  交接入职模块：候选人是否录用仍由HR决定，入职是下一张任务单", 11.5, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    base.add_notes(
        slide,
        "招聘段收口约30秒。概括前后差异，不重复解释代码。最后交接：招聘Agent已经形成可用于HR复核的候选人画像，"
        "但是否录用仍由HR决定；接下来交给入职模块展示录用后的流程追踪。"
        "如果队友顺序不是入职，可把最后一句改成实际下一位模块。",
    )


def main() -> None:
    source = Presentation(str(base.TEMPLATE))
    prs = Presentation(str(base.TEMPLATE))
    base.clear_slides(prs)
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL

    # 已确认的开头四页
    base.cover(prs, source)
    base.policy(prs)
    base.research(prs)
    base.scheme(prs)

    # 总控 + 招聘技能演示
    architecture(prs)
    router_slide(prs)
    recruitment_map(prs)
    mechanism_problem(prs)
    desensitization(prs)
    scoring(prs)
    number_audit(prs)
    close_loop(prs)

    prs.core_properties.title = "小慧 · 开头及总控招聘技能演示初版"
    prs.core_properties.subject = "职业院校技能大赛人工智能赛道展示"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
