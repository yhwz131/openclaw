#!/usr/bin/env python3
"""一期成果页加入样本测算条：不新开效率专页。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

import build_closing_pptx as closing
import build_opening_template_pptx as base


OUT = Path("/opt/openclaw-hr/docs/PPT准备/量化数据-成果页样本测算稿.pptx")
GREEN = closing.GREEN
PURPLE = closing.PURPLE


def compact_tile(slide, path: Path, x, y, title: str, sub: str, accent) -> None:
    base.rect(slide, x, y, Inches(3.75), Inches(1.72), base.NAVY_2, line=base.LINE, radius=True)
    slide.shapes.add_picture(str(path), x + Inches(0.08), y + Inches(0.08), width=Inches(3.59), height=Inches(1.12))
    base.rect(slide, x + Inches(0.08), y + Inches(1.22), Inches(3.59), Inches(0.4), base.NAVY_3)
    base.text(slide, x + Inches(0.18), y + Inches(1.26), Inches(1.7), Inches(0.3), title, 11, bold=True, color=accent, anchor=MSO_ANCHOR.MIDDLE)
    base.text(slide, x + Inches(1.85), y + Inches(1.26), Inches(1.7), Inches(0.3), sub, 9.5, color=base.MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def result_with_sample(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    if not getattr(base, "COMPOSE_TOTAL", None):
        base.TOTAL = 4
    closing.closing_header(
        slide,
        1,
        0,
        "05  PHASE I DELIVERABLES",
        "从设计到落地：一期建设成果",
        "先交付可运行产物，再用样本测算说明初筛整理效率",
    )

    base.rect(slide, Inches(0.48), Inches(2.38), Inches(4.18), Inches(3.42), base.NAVY_2, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.74), Inches(2.52), Inches(3.64), Inches(0.3), "小慧 · HR数字员工", 16, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(0.74), Inches(2.86), Inches(3.64), Inches(0.3), "1个协调Agent + 5个业务Agent", 13, bold=True, align=PP_ALIGN.CENTER)
    modules = [
        ("总控", "识别·分发·返回"),
        ("招聘", "画像·人才库·纪要"),
        ("入职", "追踪·审批·跟进"),
        ("培训", "按需生成资料"),
        ("绩效", "360聚合·报告看板"),
        ("行政", "问答·权限·月报"),
    ]
    for i, (name, body) in enumerate(modules):
        col, row = i % 2, i // 2
        x = Inches(0.73 + col * 1.9)
        y = Inches(3.28 + row * 0.58)
        base.rect(slide, x, y, Inches(1.72), Inches(0.48), base.NAVY_3, line=base.LINE, radius=True)
        base.text(slide, x + Inches(0.08), y + Inches(0.08), Inches(0.5), Inches(0.3), name, 10.5, bold=True, color=base.CYAN, anchor=MSO_ANCHOR.MIDDLE)
        base.text(slide, x + Inches(0.56), y + Inches(0.08), Inches(1.08), Inches(0.3), body, 8.5, color=base.PALE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    base.rect(slide, Inches(0.73), Inches(5.08), Inches(3.66), Inches(0.48), base.NAVY_3, line=GREEN, radius=True)
    base.text(
        slide,
        Inches(0.73),
        Inches(5.08),
        Inches(3.66),
        Inches(0.28),
        "可运行  ·  可检查  ·  可交付",
        11,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.text(
        slide,
        Inches(0.73),
        Inches(5.34),
        Inches(3.66),
        Inches(0.18),
        "WebChat / CLI / 飞书 / HCM / 本地人才库",
        8.5,
        color=base.MUTED,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    images = [
        (closing.crop_image(closing.ASSETS / "recruitment-profile.png", "recruitment-profile-wide.png", focus_y=0.0), "招聘画像", "八维评分与画像", base.CYAN),
        (closing.crop_image(closing.ASSETS / "onboarding-dashboard.png", "onboarding-dashboard-wide.png", focus_y=0.1), "入职看板", "HCM节点与跟进", GREEN),
        (closing.crop_image(closing.ASSETS / "performance-dashboard.png", "performance-dashboard-wide.png", focus_y=0.0), "绩效看板", "360反馈聚合", PURPLE),
        (closing.crop_image(closing.ASSETS / "monthly-report.png", "monthly-report-wide.png", focus_y=0.0), "行政月报", "报告期 2026-09", base.ORANGE),
    ]
    positions = [
        (Inches(4.94), Inches(2.38)),
        (Inches(8.91), Inches(2.38)),
        (Inches(4.94), Inches(4.22)),
        (Inches(8.91), Inches(4.22)),
    ]
    for (path, title, sub, accent), (x, y) in zip(images, positions):
        compact_tile(slide, path, x, y, title, sub, accent)

    base.rect(slide, Inches(0.48), Inches(6.12), Inches(12.36), Inches(0.7), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.68), Inches(6.16), Inches(12.0), Inches(0.2), "样本测算｜招聘初筛整理 · 脱敏演示样本 · 同一岗位说明 · 非正式上线数据", 10.5, bold=True, color=base.CYAN)
    metrics = [
        ("人工", "约 8 分钟/份"),
        ("系统", "约 1.2 分钟/份"),
        ("该环节", "耗时约降 85%"),
        ("一期事项", "约 18 / 30"),
    ]
    for i, (label, value) in enumerate(metrics):
        x = Inches(0.68 + i * 3.02)
        base.text(slide, x, Inches(6.38), Inches(0.95), Inches(0.28), label, 10, color=base.MUTED, anchor=MSO_ANCHOR.MIDDLE)
        base.text(slide, x + Inches(0.92), Inches(6.38), Inches(2.0), Inches(0.28), value, 12.5, bold=True, color=base.PALE, anchor=MSO_ANCHOR.MIDDLE)

    base.add_notes(
        slide,
        "本页先展示一期产物。底部是样本测算，不是企业上线KPI。"
        "口播只说：脱敏演示样本、同一份岗位说明，计时从打开附件到形成可呈报材料；"
        "人工大约八分钟一份，系统辅助链大约一分多钟，只比这一环节的整理耗时。"
        "一期大约覆盖十八项事务，不把事项覆盖说成工时百分比，也不比录用对错。"
        "机制有效性仍用前面技能演示证明。",
    )


def main() -> None:
    prs = Presentation(str(base.TEMPLATE))
    base.clear_slides(prs)
    prs.slide_width, prs.slide_height = base.W, base.H
    result_with_sample(prs)
    prs.core_properties.title = "小慧 · 一期成果（含样本测算）"
    prs.core_properties.subject = "成果栏中的效率部分，不新开专页"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
