"""从 v2.2 导出四张模块封面（招聘 / 入职 / 行政 / 绩效），三栏预告改为模块业务全貌。v2.2 不改。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v2.2.pptx"
DST = "/opt/openclaw-hr/docs/PPT准备/封面页-模块全貌版.pptx"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
E = 914400

NAVY, CARD, MID, DEEP = "02163B", "032B5C", "04457D", "00112C"
CYAN, GREEN, ORANGE = "00D8FF", "26D098", "FFA421"
WHITE, LIGHT, GREY = "FFFFFF", "E9F8FF", "9EC2D8"
C = RGBColor.from_string
WHO = {"模型": CYAN, "脚本": GREEN, "Lobster": GREEN, "Cron": GREEN, "网关": GREEN, "人": ORANGE, "数据": MID}

p = Presentation(SRC)
KEEP = [14, 19, 24, 29]
lst = p.slides._sldIdLst
for k, el in enumerate(list(lst)):
    if k + 1 not in KEEP:
        p.part.drop_rel(el.get(R + "id"))
        lst.remove(el)
S = dict(zip(KEEP, p.slides))


def rect(s, l, t, w, h, fill, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = C(fill)
    if line:
        r.line.color.rgb = C(line)
        r.line.width = Pt(lw)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def tb(s, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, font="微软雅黑"):
    b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, (text, size, color, bold) in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run()
        r.text = text
        f = r.font
        f.size, f.color.rgb, f.bold, f.name = Pt(size), C(color), bold, font
        rpr = r._r.get_or_add_rPr()
        rpr.append(rpr.makeelement(A + "ea", {"typeface": font}))
    return b


def chip(s, l, t, w, h, text, fill, color=NAVY, size=10.5):
    rect(s, l, t, w, h, fill)
    tb(s, l, t, w, h, [(text, size, color, True)], align=PP_ALIGN.CENTER)


def prepare(s, head):
    for sh in list(s.shapes):
        if 44 <= sh.shape_id <= 57:
            sh._element.getparent().remove(sh._element)
    for sh in s.shapes:
        if 40 <= sh.shape_id <= 43 or (sh.shape_id == 36 and sh.has_text_frame and "道坎" in sh.text_frame.text):
            sh.top -= Inches(0.8)
        if sh.shape_id == 43:
            for para in sh.text_frame.paragraphs:
                if para.runs and para.runs[-1].text.endswith("；"):
                    para.runs[-1].text = para.runs[-1].text[:-1]
    tb(s, 0.9, 3.6, 7.0, 0.36, [(head, 13, CYAN, True)])
    for k, who in enumerate(["模型", "脚本", "人"]):
        chip(s, 10.15 + k * 0.78, 3.66, 0.68, 0.26, who, WHO[who], size=10)


def arrow(s, l, t, h, size=16):
    tb(s, l, t, 0.16, h, [("›", size, GREY, True)], align=PP_ALIGN.CENTER)


# ---------- 招聘：三条链 ----------
s = S[14]
prepare(s, "业务全貌 · 三条链")
MAIN = [("收简历", "脚本", ""), ("转文档", "脚本", ""), ("解析字段", "模型", ""), ("原文回溯", "脚本", ""),
        ("删敏感字段", "脚本", "闸门 ①"), ("八维评分", "模型", ""), ("数字回核", "脚本", "闸门 ②"),
        ("生成画像", "模型", ""), ("HR 约面", "人", "")]
SIDE = [("人才库", [("收简历 · 解析（共用）", "数据"), ("向量入库", "模型"), ("余弦检索 Top N", "脚本"),
                  ("讲清“为什么像”", "模型"), ("HR 挑人", "人")]),
        ("面试纪要", [("面试录音", "数据"), ("转写 · 分说话人", "模型"), ("整理纪要", "模型"),
                    ("口述数字回核", "脚本"), ("交 HR", "人")])]
X0, W_ALL = 2.62, 9.83

rect(s, 0.9, 4.08, 1.6, 1.0, CARD, line=CYAN, lw=1.5)
tb(s, 0.9, 4.12, 1.6, 0.56, [("筛选画像", 14, WHITE, True)], align=PP_ALIGN.CENTER)
tb(s, 0.9, 4.66, 1.6, 0.34, [("★ 今天演示", 10.5, ORANGE, True)], align=PP_ALIGN.CENTER)
w = (W_ALL - 0.13 * 8) / 9
for k, (text, who, tag) in enumerate(MAIN):
    x = X0 + k * (w + 0.13)
    col = WHO[who]
    rect(s, x, 4.08, w, 1.0, CARD, line=col, lw=1.5)
    rect(s, x, 4.08, w, 0.07, col)
    tb(s, x, 4.18, w, 0.55, [(text, 12, WHITE, True)], align=PP_ALIGN.CENTER)
    tb(s, x, 4.72, w, 0.3, [(tag or who, 9.5, col, bool(tag))], align=PP_ALIGN.CENTER)
    if k < 8:
        arrow(s, x + w - 0.015, 4.3, 0.5, 13)

for r_, (name, steps) in enumerate(SIDE):
    y = 5.3 + r_ * 0.8
    rect(s, 0.9, y, 1.6, 0.62, DEEP, line=MID)
    tb(s, 0.9, y, 1.6, 0.62, [(name, 13, LIGHT, True)], align=PP_ALIGN.CENTER)
    w2 = (W_ALL - 0.16 * 4) / 5
    for k, (text, who) in enumerate(steps):
        x = X0 + k * (w2 + 0.16)
        col = WHO[who]
        rect(s, x, y, w2, 0.62, DEEP, line=col)
        tb(s, x, y, w2, 0.62, [(text, 12, LIGHT, False)], align=PP_ALIGN.CENTER)
        if k < 4:
            arrow(s, x + w2, y + 0.08, 0.46, 13)

s.notes_slide.notes_text_frame.text = """【男1】招聘封面（约 25 秒）

招聘 Agent 一共三条链。【指颜色】颜色就是分工：青色是模型，绿色是脚本，橙色是人。

【指主链】第一条，筛简历出画像：收简历、转文档、解析、评分、出画像。模型每写一步，后面都有脚本核：原文回溯、删敏感字段、数字回核。约不约面，由 HR 定。今天演示的就是这一条。

【指下面两条】另外两条：人才库按向量检索相似候选人，名次由脚本算；面试录音转写分人，口述的数字也要回核。

招聘是第一道坎：说得不准。"""

# ---------- 入职：一条跨部门时间轴 ----------
s = S[19]
prepare(s, "业务全貌 · 一条跨部门时间轴（今天现场走完整条）")
PHASES = [
    ("① 报到前 · 五个部门准备", MID, WHITE,
     [("人", "五个部门在 HCM 各自办结：合同 · IT · 行政 · 培训 · Buddy"),
      ("脚本", "抽取 HCM 节点进度，只读不写"),
      ("模型", "汇报卡点，建议催哪个部门")]),
    ("② 报到时 · 两道审批门", ORANGE, NAVY,
     [("Lobster", "合同门、账号门写死在流程里，到门即停、进程退出"),
      ("人", "HR 主管审合同，IT 核账号权限；批一次只走一步"),
      ("模型", "原样转交审批问题，不代批")]),
    ("③ 报到后 · 试用期跟进", MID, WHITE,
     [("Cron", "五个部门节点全绿，才挂 7/30/60/90 四个提醒"),
      ("模型", "到点在飞书提醒 HR，不改表"),
      ("人", "导师带教、岗位培训、关怀沟通、胜任度初评；HR 回填才算完成")]),
]
for k, (head, hfill, hcolor, rows) in enumerate(PHASES):
    x = 0.9 + k * 3.9
    rect(s, x, 4.08, 3.75, 2.78, CARD, line=ORANGE if k == 1 else MID)
    rect(s, x, 4.08, 3.75, 0.46, hfill)
    tb(s, x, 4.08, 3.75, 0.46, [(head, 13.5, hcolor, True)], align=PP_ALIGN.CENTER)
    for j, (who, text) in enumerate(rows):
        y = 4.66 + j * 0.72
        chip(s, x + 0.14, y + 0.17, 0.82, 0.3, who, WHO[who], size=10)
        tb(s, x + 1.04, y, 2.65, 0.64, [(text, 11.5, LIGHT, False)])
    if k < 2:
        arrow(s, x + 3.75, 5.2, 0.5, 18)

s.notes_slide.notes_text_frame.text = """【女1】入职封面（约 25 秒）

【接过男1，走到台中】

谢谢。录用确认之后，进入入职。入职是一条跨部门的时间轴。【指 ①】报到前，五个部门在 HCM 里各自办结，Agent 只读进度、报卡点，不替部门办。【指 ②】报到时有两道审批门，写死在 Lobster 流程里：合同由 HR 主管审，账号权限由 IT 核，到门就停。【指 ③】报到后，五个节点全绿，才挂上第 7、30、60、90 天的提醒，HR 谈完回填才算完成。

这是第二道坎：做得越界。今天现场把这条时间轴走一遍。"""

# ---------- 行政：三项业务 ----------
s = S[24]
prepare(s, "业务全貌 · 三项业务")
BIZ = [
    ("制度答疑", "★ 今天演示", ORANGE,
     [("数据", "员工在飞书提问"), ("脚本", "检索 14 份制度，无命中即退出"),
      ("模型", "照原文作答，标出处"), ("人", "个案、争议转 HR")]),
    ("员工数据查询", "★ 今天演示", ORANGE,
     [("网关", "飞书配对：未登记只收到配对码"), ("网关", "按发送人分配工具"),
      ("脚本", "闸门先判身份：非 HR 退出码 3"), ("人", "HR 过闸，读 HCM")]),
    ("HR 月报", "不展开", GREY,
     [("数据", "6 张台账"), ("脚本", "算 18 个指标，逐项回核来源"),
      ("模型", "转述结果、解释口径，不改数"), ("人", "HR 审阅后使用")]),
]
for k, (name, tag, tcol, steps) in enumerate(BIZ):
    x = 0.9 + k * 3.9
    rect(s, x, 4.08, 3.75, 2.78, CARD, line=MID)
    tb(s, x + 0.15, 4.12, 2.3, 0.44, [(name, 14.5, WHITE, True)])
    tb(s, x + 2.3, 4.12, 1.35, 0.44, [(tag, 10.5, tcol, True)], align=PP_ALIGN.RIGHT)
    for j, (who, text) in enumerate(steps):
        y = 4.64 + j * 0.54
        col = WHO[who]
        rect(s, x + 0.15, y, 3.45, 0.42, DEEP, line=col)
        rect(s, x + 0.15, y, 0.07, 0.42, col)
        tb(s, x + 0.3, y, 0.62, 0.42, [(who, 9.5, col if who != "数据" else GREY, True)])
        tb(s, x + 0.9, y, 2.68, 0.42, [(text, 11.5, LIGHT, False)])

s.notes_slide.notes_text_frame.text = """【男2】行政封面（约 25 秒）

【接过女1，走到台中】

谢谢。行政 Agent 有三项业务。【指左】制度答疑：先检索 14 份制度原文，照原文答、标出处，个案和争议交给 HR。【指中】员工数据查询：飞书先认人，网关按人发工具，闸门脚本在读数据前判身份，只有 HR 能读 HCM。【指右】HR 月报：6 张台账，脚本算 18 个指标，每个数字回核到来源，模型只转述、不改数。

这是第三道坎：查得不清。今天现场演示前两项。"""

# ---------- 绩效：一条链 + 视角与维度 ----------
s = S[29]
prepare(s, "业务全貌 · 360 反馈一条链（今天现场分三轮走完）")
CHAIN = [("收齐 360 反馈", "数据"), ("校验齐全：缺自评即停", "脚本"), ("按身份分组求均值", "脚本"),
         ("极差 ≥ 1 标疑似分歧", "脚本"), ("读评语，判共识或分歧", "模型"), ("写面谈提纲，出看板", "脚本"),
         ("主管面谈后定级", "人")]
w = (11.55 - 0.14 * 6) / 7
for k, (text, who) in enumerate(CHAIN):
    x = 0.9 + k * (w + 0.14)
    col = WHO[who]
    rect(s, x, 4.08, w, 1.32, CARD, line=col, lw=1.5)
    rect(s, x, 4.08, w, 0.07, col)
    tb(s, x, 4.2, w, 0.3, [(who, 10, col if who != "数据" else GREY, True)], align=PP_ALIGN.CENTER)
    tb(s, x + 0.04, 4.5, w - 0.08, 0.82, [(text, 12.5, WHITE, True)], align=PP_ALIGN.CENTER)
    if k < 6:
        arrow(s, x + w - 0.01, 4.5, 0.5, 14)
for k, (head, body) in enumerate([("四个视角 · 7 份反馈", "上级  ·  平级  ·  下属  ·  自评"),
                                  ("五个维度", "目标达成 · 专业能力 · 协作沟通 · 责任心 · 学习成长")]):
    x = 0.9 + k * 5.85
    rect(s, x, 5.62, 5.7, 1.2, DEEP, line=MID)
    tb(s, x + 0.2, 5.7, 5.3, 0.38, [(head, 11.5, GREEN, True)])
    tb(s, x + 0.2, 6.1, 5.3, 0.6, [(body, 15, WHITE, True)])

s.notes_slide.notes_text_frame.text = """【女2】绩效封面（约 25 秒）

【接过男2，走到台中】

谢谢。我来讲绩效：把 360 反馈变成面谈材料，不给评级。【指链条】一条链走完：先校验齐不齐，缺自评就停；脚本按身份分组求均值，极差到 1 分标“疑似分歧”；Agent 只读评语，判断是共识还是分歧；面谈提纲和看板由脚本生成；定级由主管面谈后决定。

【指下方】每个人收 7 份反馈，来自上级、平级、下属和自评四个视角，看五个维度。今天现场分三轮把这条链走完。"""

p.save(DST)
print("saved", DST, len(p.slides))
