#!/usr/bin/env python3
"""绩效模块四页完善版：双闸门、匿名化、语义分层与人在回路。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import build_opening_template_pptx as base


ROOT = Path("/opt/openclaw-hr/docs")
OUT = ROOT / "PPT准备/绩效模块技能演示-4页完善版.pptx"
ASSETS = ROOT / "ppt/assets/performance"
DASHBOARD = ASSETS / "dashboard-score-table.png"
TOTAL = 4

GREEN = RGBColor(0x26, 0xD0, 0x98)
GREEN_DARK = RGBColor(0x08, 0x50, 0x57)
PURPLE = RGBColor(0x8C, 0x7C, 0xFF)
PURPLE_DARK = RGBColor(0x20, 0x1B, 0x55)
RED = RGBColor(0xF1, 0x5A, 0x66)
RED_DARK = RGBColor(0x56, 0x17, 0x26)
CODE_BG = RGBColor(0x00, 0x11, 0x2C)


def connector(slide, x1, y1, x2, y2, color=base.CYAN, width=1.1, arrow=True):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def secondary_nav(slide) -> None:
    labels = ["总控", "招聘", "入职", "培训", "绩效", "行政"]
    base.rect(slide, Inches(0), Inches(0.72), base.W, Inches(0.45), base.NAVY, line=base.LINE)
    base.text(
        slide,
        Inches(0.38),
        Inches(0.79),
        Inches(0.9),
        Inches(0.27),
        "技能地图",
        10,
        bold=True,
        color=base.MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    for i, label in enumerate(labels):
        x = Inches(1.33 + i * 1.77)
        active = i == 4
        base.rect(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            base.CYAN if active else base.NAVY_2,
            line=base.CYAN if active else base.LINE,
            radius=True,
        )
        base.text(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            label,
            10.5,
            bold=active,
            color=base.NAVY if active else base.MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def header(slide, page: int, eyebrow: str, title: str, subtitle: str) -> None:
    base.add_navigation(slide, 2, page)
    secondary_nav(slide)
    base.text(slide, Inches(0.48), Inches(1.28), Inches(5.0), Inches(0.25), eyebrow, 10.5, bold=True, color=base.CYAN)
    base.text(slide, Inches(0.46), Inches(1.53), Inches(12.2), Inches(0.49), title, 22, bold=True, font=base.TITLE_FONT)
    base.rect(slide, Inches(0.47), Inches(2.08), Inches(0.75), Inches(0.05), base.CYAN)
    base.text(slide, Inches(1.4), Inches(1.99), Inches(11.1), Inches(0.29), subtitle, 11.2, color=base.MUTED)


def card(slide, x, y, w, h, title: str, body: str, *, accent=base.CYAN, fill=base.NAVY_2, body_size=11.5):
    base.rect(slide, x, y, w, h, fill, line=accent, radius=True, transparency=3)
    base.rect(slide, x, y, Inches(0.08), h, accent)
    base.text(slide, x + Inches(0.2), y + Inches(0.14), w - Inches(0.4), Inches(0.3), title, 13, bold=True, color=accent)
    base.text(slide, x + Inches(0.2), y + Inches(0.58), w - Inches(0.4), h - Inches(0.7), body, body_size, color=base.PALE, line_spacing=0.95)


def pill(slide, x, y, w, text_value: str, color=base.CYAN):
    base.rect(slide, x, y, w, Inches(0.34), base.NAVY_3, line=color, radius=True)
    base.text(
        slide,
        x,
        y,
        w,
        Inches(0.34),
        text_value,
        10.2,
        bold=True,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def slide_positioning(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        1,
        "05  PERFORMANCE · POSITIONING",
        "绩效 Agent：把360反馈转化为可复核的面谈材料",
        "一期聚焦360反馈聚合，不替代目标管理，也不输出S/A/B/C/D最终评级",
    )

    flow = [
        ("反馈摄入", base.CYAN),
        ("评价人脱敏", base.CYAN),
        ("数据校验", base.ORANGE),
        ("机械聚合", base.ORANGE),
        ("语义分析", PURPLE),
        ("HTML看板", GREEN),
    ]
    for i, (label, accent) in enumerate(flow):
        x = Inches(0.48 + i * 2.04)
        base.rect(slide, x, Inches(2.4), Inches(1.8), Inches(0.62), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x, Inches(2.4), Inches(1.8), Inches(0.62), label, 11.5, bold=True, color=accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(flow) - 1:
            connector(slide, x + Inches(1.8), Inches(2.71), x + Inches(2.0), Inches(2.71), accent)

    card(
        slide,
        Inches(0.48),
        Inches(3.45),
        Inches(3.9),
        Inches(1.76),
        "脚本负责",
        "校验结构与分值\n计算平均分、极差\n形成唯一数字基准",
        accent=base.CYAN,
        body_size=12,
    )
    card(
        slide,
        Inches(4.7),
        Inches(3.45),
        Inches(3.9),
        Inches(1.76),
        "Agent负责",
        "读取多方评语\n分析共识与分歧\n生成面谈问题",
        accent=PURPLE,
        fill=PURPLE_DARK,
        body_size=12,
    )
    card(
        slide,
        Inches(8.92),
        Inches(3.45),
        Inches(3.9),
        Inches(1.76),
        "人工负责",
        "直接主管初评\n隔级主管复核\n校准会确认",
        accent=base.ORANGE,
        body_size=12,
    )

    base.rect(slide, Inches(0.48), Inches(5.58), Inches(12.36), Inches(0.96), base.NAVY_3, line=RED, radius=True)
    base.text(
        slide,
        Inches(0.72),
        Inches(5.72),
        Inches(11.9),
        Inches(0.56),
        "明确边界｜聚合数据 + 综合分析 + 看板 = 面谈输入材料 ≠ 绩效评级结果",
        13,
        bold=True,
        color=RED,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约55秒。完整绩效管理还包括目标设定、过程辅导和改进跟踪，一期只聚焦最适合标准化的360反馈聚合。"
        "脚本负责校验与计算，Agent负责理解评语，人负责最终评级。系统不输出S/A/B/C/D，"
        "产物只作为主管面谈和校准复核的输入材料。",
    )


def slide_gate_aggregate(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        2,
        "05  DATA GATE · DETERMINISTIC AGGREGATION",
        "先脱敏、再校验、后聚合：不让脏数据进入评分基准",
        "评价人身份最小化｜check前置闸门｜aggregate二次校验｜数字全部由脚本计算",
    )

    steps = [
        ("01", "压缩包摄入", "DOCX → JSON"),
        ("02", "角色代号", "物理脱敏"),
        ("03", "前置校验", "退出码 0 / 1"),
        ("04", "二次校验", "拒绝绕过"),
        ("05", "机械聚合", "唯一数字基准"),
    ]
    accents = [base.CYAN, base.CYAN, base.ORANGE, base.ORANGE, GREEN]
    for i, ((num, title, sub), accent) in enumerate(zip(steps, accents)):
        x = Inches(0.48 + i * 2.5)
        base.rect(slide, x, Inches(2.4), Inches(2.22), Inches(1.02), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x + Inches(0.15), Inches(2.53), Inches(0.4), Inches(0.25), num, 10, bold=True, color=accent)
        base.text(slide, x + Inches(0.55), Inches(2.49), Inches(1.48), Inches(0.28), title, 11.8, bold=True, color=base.PALE, align=PP_ALIGN.CENTER)
        base.text(slide, x + Inches(0.18), Inches(2.9), Inches(1.86), Inches(0.25), sub, 10.2, color=base.MUTED, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            connector(slide, x + Inches(2.22), Inches(2.91), x + Inches(2.47), Inches(2.91), accent)

    card(
        slide,
        Inches(0.48),
        Inches(3.78),
        Inches(3.92),
        Inches(1.55),
        "身份最小化",
        "领导1 · 同事1 · 下属1 · 自评\n真人姓名不进入报告和看板",
        accent=base.CYAN,
        body_size=11.8,
    )
    card(
        slide,
        Inches(4.7),
        Inches(3.78),
        Inches(3.92),
        Inches(1.55),
        "完整性闸门",
        "五维完整 · 评分1—5整数\n至少有上级 + 自评｜失败 exit 1",
        accent=RED,
        fill=RED_DARK,
        body_size=11.8,
    )
    card(
        slide,
        Inches(8.92),
        Inches(3.78),
        Inches(3.92),
        Inches(1.55),
        "确定性计算",
        "平均分 · 最小值 · 最大值 · 极差\naggregated-data.json 是唯一基准",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=11.8,
    )

    base.rect(slide, Inches(0.48), Inches(5.63), Inches(12.36), Inches(0.92), base.NAVY_3, line=PURPLE, radius=True)
    base.text(slide, Inches(0.7), Inches(5.78), Inches(2.15), Inches(0.25), "法律锚点", 11.5, bold=True, color=PURPLE)
    base.text(
        slide,
        Inches(2.1),
        Inches(5.72),
        Inches(10.4),
        Inches(0.54),
        "《个人信息保护法》第六条｜只处理完成反馈分析所必需的信息，不让评价人真实身份进入聚合产物",
        11.6,
        color=base.PALE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约1分15秒。DOCX反馈先解析为JSON，评价人姓名被物理改写为角色代号。"
        "check_feedback检查字段、五维、1到5整数，并确认至少有上级和自评；失败列出问题并exit 1。"
        "aggregate再跑同一套规则，防止绕过前置检查。平均分、最小值、最大值和视角极差全部由脚本计算，"
        "aggregated-data.json是下游唯一数字基准。",
    )


def slide_semantic_layer(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        3,
        "05  SIGNAL ≠ VERDICT · SEMANTIC LAYER",
        "机械信号只负责提醒，业务结论必须结合评语",
        "极差≥1.0只标记“疑似分歧”；看板判定来自综合分析报告，不照搬机械阈值",
    )

    card(
        slide,
        Inches(0.48),
        Inches(2.4),
        Inches(3.0),
        Inches(1.42),
        "机械层｜脚本",
        "四视角 × 五维度\n计算平均分与视角极差\n极差≥1.0 → 疑似分歧",
        accent=base.CYAN,
        body_size=11.3,
    )
    card(
        slide,
        Inches(0.48),
        Inches(4.02),
        Inches(3.0),
        Inches(1.42),
        "语义层｜Agent",
        "读取各视角原始评语\n判断叙事是否真正冲突\n形成归因假设与面谈问题",
        accent=PURPLE,
        fill=PURPLE_DARK,
        body_size=11.3,
    )
    connector(slide, Inches(1.98), Inches(3.82), Inches(1.98), Inches(4.0), PURPLE)

    base.rect(slide, Inches(3.75), Inches(2.4), Inches(2.45), Inches(3.04), base.NAVY_2, line=base.ORANGE, radius=True)
    base.text(slide, Inches(3.98), Inches(2.6), Inches(1.99), Inches(0.26), "关键案例｜责任心", 12, bold=True, color=base.ORANGE, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(4.1), Inches(3.08), Inches(1.75), Inches(0.4), "极差 = 1.00", 19, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)
    pill(slide, Inches(4.12), Inches(3.62), Inches(1.7), "机械：疑似分歧", base.ORANGE)
    connector(slide, Inches(4.97), Inches(3.98), Inches(4.97), Inches(4.18), base.ORANGE)
    pill(slide, Inches(4.12), Inches(4.2), Inches(1.7), "语义：共识", GREEN)
    base.text(slide, Inches(3.98), Inches(4.68), Inches(1.99), Inches(0.52), "方向一致，只是对\n“责任范围”期待不同", 9.8, color=base.PALE, align=PP_ALIGN.CENTER)

    base.rect(slide, Inches(6.46), Inches(2.4), Inches(6.38), Inches(3.5), base.NAVY_2, line=GREEN, radius=True)
    base.text(slide, Inches(6.67), Inches(2.55), Inches(5.96), Inches(0.25), "真实看板证据｜员工1 · 模拟数据", 12, bold=True, color=GREEN)
    if DASHBOARD.exists():
        slide.shapes.add_picture(str(DASHBOARD), Inches(6.67), Inches(2.91), width=Inches(5.96), height=Inches(2.75))

    base.rect(slide, Inches(0.48), Inches(6.08), Inches(12.36), Inches(0.52), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(
        slide,
        Inches(0.68),
        Inches(6.08),
        Inches(11.96),
        Inches(0.52),
        "数字只取 aggregated-data.json｜判定与卡片只取 aggregated-report.md｜看板不重新算分",
        11.2,
        bold=True,
        color=base.CYAN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约1分20秒。极差大于等于1.0只触发疑似分歧，不是业务结论。员工1责任心极差刚好1.00，"
        "机械层触发关注；但读完评语后发现各方都认可其负责，只对责任范围是否延伸到复盘和协作期待不同，"
        "语义层将其归为共识。看板数字只取aggregated-data.json，判定和卡片只取aggregated-report.md，"
        "自身不重新算分，也不输出评级。",
    )


def slide_live_demo(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        4,
        "05  FAIL CLOSED · LIVE DEMO",
        "缺少关键视角，系统宁可停止，也不生成“伪360”",
        "现场演示：缺自评失败 → 完整样本通过 → 机械聚合 → 打开看板；本模块不手写代码",
    )

    card(
        slide,
        Inches(0.48),
        Inches(2.4),
        Inches(3.7),
        Inches(1.52),
        "01｜失败样本",
        "yuangong1-noself-demo\n仅上级 + 平级，没有 self\n预期：退出码 1，不进入聚合",
        accent=RED,
        fill=RED_DARK,
        body_size=11.3,
    )
    connector(slide, Inches(4.18), Inches(3.16), Inches(4.52), Inches(3.16), RED)
    card(
        slide,
        Inches(4.55),
        Inches(2.4),
        Inches(3.7),
        Inches(1.52),
        "02｜完整样本",
        "yuangong1\n上级 · 平级 · 下属 · 自评\n预期：校验通过，角色代号可见",
        accent=base.ORANGE,
        body_size=11.3,
    )
    connector(slide, Inches(8.25), Inches(3.16), Inches(8.59), Inches(3.16), base.ORANGE)
    card(
        slide,
        Inches(8.62),
        Inches(2.4),
        Inches(4.2),
        Inches(1.52),
        "03｜聚合与看板",
        "aggregated-data.json\n雷达图 · 对照表 · 共识/分歧\n页面不出现 S/A/B/C/D",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=11.3,
    )

    base.rect(slide, Inches(0.48), Inches(4.18), Inches(8.35), Inches(1.72), CODE_BG, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.7), Inches(4.34), Inches(7.9), Inches(0.22), "现场命令｜现行 Skill 入口", 10.5, bold=True, color=base.CYAN)
    code = (
        "$ python3 skills/aggregating-360-feedback/scripts/check_feedback.py \\\n"
        "    feedback-360/source_feedback/yuangong1-noself-demo\n"
        "$ python3 skills/aggregating-360-feedback/scripts/check_feedback.py \\\n"
        "    feedback-360/source_feedback/yuangong1\n"
        "$ python3 skills/aggregating-360-feedback/scripts/aggregate.py \\\n"
        "    feedback-360/source_feedback/yuangong1\n"
        "$ bash skills/generating-360-dashboard/scripts/build_dashboard.sh \\\n"
        "    feedback-360/source_feedback/yuangong1 --publish --force"
    )
    base.text(slide, Inches(0.7), Inches(4.64), Inches(7.9), Inches(1.14), code, 7.7, color=base.PALE, font="Courier New", line_spacing=0.78)

    base.rect(slide, Inches(9.08), Inches(4.18), Inches(3.74), Inches(1.72), base.NAVY_3, line=PURPLE, radius=True)
    base.text(slide, Inches(9.3), Inches(4.34), Inches(3.3), Inches(0.25), "最终决策仍由人完成", 12, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(9.3), Inches(4.82), Inches(3.3), Inches(0.66), "直接主管初评\n→ 隔级主管复核\n→ 校准会确认", 11.2, color=base.PALE, align=PP_ALIGN.CENTER)

    base.rect(slide, Inches(0.48), Inches(6.12), Inches(12.36), Inches(0.48), base.NAVY_3, line=base.ORANGE, radius=True)
    base.text(
        slide,
        Inches(0.68),
        Inches(6.12),
        Inches(11.96),
        Inches(0.48),
        "《个人信息保护法》第二十四条｜对个人权益有重大影响的决定，不应仅由自动化结果完成",
        10.8,
        bold=True,
        color=base.ORANGE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约3分钟。先运行yuangong1-noself-demo，脚本指出缺少self并exit 1，不产生聚合结果。"
        "再切换完整yuangong1，展示四类视角和角色代号，运行aggregate生成唯一数字基准，"
        "最后打开Canvas看板，确认有雷达图、对照表、共识和分歧，但没有S/A/B/C/D评级。"
        "收口：脚本保证完整和稳定，Agent帮助理解反馈，人保留最终决定。"
        "培训若已保留手写代码，本模块只运行现有脚本。",
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL
    slide_positioning(prs)
    slide_gate_aggregate(prs)
    slide_semantic_layer(prs)
    slide_live_demo(prs)
    prs.core_properties.title = "小慧 · 绩效模块技能演示完善版"
    prs.core_properties.subject = "360反馈双闸门、评价人匿名化、确定性聚合、语义分层与人在回路"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
