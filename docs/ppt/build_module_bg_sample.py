"""样张：封面（全貌版）+ 每模块 1 张内页，加模块背景变化。v2.2 不改。"""
import copy
from lxml import etree
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

COVER_SCRIPT = "/opt/openclaw-hr/docs/ppt/build_covers_overview.py"
DST = "/opt/openclaw-hr/docs/PPT准备/模块背景-样张.pptx"
PAIRS = [(14, 16), (19, 21), (24, 26), (29, 30)]
ns = {}
code = open(COVER_SCRIPT, encoding="utf-8").read()
code = code.replace("KEEP = [14, 19, 24, 29]", "KEEP = [14, 16, 19, 21, 24, 26, 29, 30]")
code = code.replace("p.save(DST)", "pass").replace('print("saved", DST, len(p.slides))', "")
exec(compile(code, COVER_SCRIPT, "exec"), ns)
p, S, E, A = ns["p"], ns["S"], ns["E"], ns["A"]

MODULE = {14: ("招聘", "00D8FF"), 19: ("入职", "FFA421"), 24: ("行政", "FF6B8A"), 29: ("绩效", "26D098")}
INNER = {16: 14, 21: 19, 26: 24, 30: 29}


def q(tag):
    return "{%s}%s" % ("http://schemas.openxmlformats.org/drawingml/2006/main", tag)


def send_behind_content(s, el):
    bg = next(sh for sh in s.shapes if sh.shape_id == 2)._element
    bg.addnext(el)


def glow(s, cx, cy, r, color, alpha):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    o.line.fill.background()
    o.shadow.inherit = False
    spPr = o._element.spPr
    for child in list(spPr):
        if child.tag in (q("solidFill"), q("gradFill"), q("noFill")):
            spPr.remove(child)
    grad = etree.SubElement(spPr, q("gradFill"), rotWithShape="1")
    gs_lst = etree.SubElement(grad, q("gsLst"))
    for pos, a in ((0, alpha), (100000, 0)):
        gs = etree.SubElement(gs_lst, q("gs"), pos=str(pos))
        clr = etree.SubElement(gs, q("srgbClr"), val=color)
        etree.SubElement(clr, q("alpha"), val=str(a))
    path = etree.SubElement(grad, q("path"), path="circle")
    etree.SubElement(path, q("fillToRect"), l="50000", t="50000", r="50000", b="50000")
    ln = spPr.find(q("ln"))
    if ln is not None:
        spPr.remove(ln)
        spPr.append(ln)
    send_behind_content(s, o._element)


def wash(s, color, alpha, fade_at=65000):
    """整页从左到右渐弱的模块色光。"""
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    r.line.fill.background()
    r.shadow.inherit = False
    spPr = r._element.spPr
    for child in list(spPr):
        if child.tag in (q("solidFill"), q("gradFill"), q("noFill")):
            spPr.remove(child)
    grad = etree.Element(q("gradFill"), rotWithShape="1")
    gs_lst = etree.SubElement(grad, q("gsLst"))
    for pos, a in ((0, alpha), (fade_at // 2, alpha // 3), (fade_at, 0), (100000, 0)):
        gs = etree.SubElement(gs_lst, q("gs"), pos=str(pos))
        clr = etree.SubElement(gs, q("srgbClr"), val=color)
        etree.SubElement(clr, q("alpha"), val=str(a))
    etree.SubElement(grad, q("lin"), ang="0", scaled="0")
    spPr.insert(list(spPr).index(spPr.find(q("prstGeom"))) + 1, grad)
    send_behind_content(s, r._element)


def watermark(s, text, color, alpha):
    b = s.shapes.add_textbox(Inches(6.6), Inches(0.95), Inches(6.5), Inches(2.9))
    tf = b.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT
    r = para.add_run()
    r.text = text
    r.font.size = Pt(190)
    r.font.bold = True
    r.font.name = "微软雅黑"
    rpr = r._r.get_or_add_rPr()
    fill = etree.SubElement(rpr, q("solidFill"))
    clr = etree.SubElement(fill, q("srgbClr"), val=color)
    etree.SubElement(clr, q("alpha"), val=str(alpha))
    rpr.append(rpr.makeelement(A + "ea", {"typeface": "微软雅黑"}))
    first = rpr.find(q("solidFill"))
    rpr.remove(first)
    rpr.insert(0, first)
    send_behind_content(s, b._element)


def tint_grid(s, color, alpha):
    for sh in s.shapes:
        if sh.shape_type == 9 and 3 <= sh.shape_id <= 9:
            ln = sh._element.spPr.find(q("ln"))
            if ln is None:
                continue
            for child in list(ln):
                if child.tag in (q("solidFill"), q("gradFill"), q("noFill")):
                    ln.remove(child)
            fill = etree.Element(q("solidFill"))
            clr = etree.SubElement(fill, q("srgbClr"), val=color)
            etree.SubElement(clr, q("alpha"), val=str(alpha))
            ln.insert(0, fill)


def strip(s, color, w):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.18), Inches(w), Inches(6.32))
    r.fill.solid()
    r.fill.fore_color.rgb = ns["C"](color)
    r.line.fill.background()
    r.shadow.inherit = False


for n, (name, col) in MODULE.items():
    s = S[n]
    wash(s, col, 17000 if col == "FFA421" else 26000)
    tint_grid(s, col, 30000)
    strip(s, col, 0.12)

for n, cover in INNER.items():
    s = S[n]
    col = MODULE[cover][1]
    for sh in s.shapes:
        l, t = sh.left / E, sh.top / E
        if sh.has_text_frame and abs(l - 0.48) < 0.05 and 1.2 < t < 1.4:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    r.font.color.rgb = ns["C"](col)
        elif sh.shape_type == 1 and abs(l - 0.48) < 0.05 and 2.05 < t < 2.15 and sh.height / E < 0.08 and sh.width / E < 0.8:
            sh.fill.solid()
            sh.fill.fore_color.rgb = ns["C"](col)
    wash(s, col, 10000 if col == "FFA421" else 15000, fade_at=55000)
    tint_grid(s, col, 22000)
    strip(s, col, 0.08)

lst = p.slides._sldIdLst
order = [n for pair in PAIRS for n in pair]
els = {n: el for n, el in zip(ns["KEEP"], list(lst))}
for el in list(lst):
    lst.remove(el)
for n in order:
    lst.append(els[n])

p.save(DST)
print("saved", DST, len(p.slides))
