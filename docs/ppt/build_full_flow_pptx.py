#!/usr/bin/env python3
"""生成“小慧”完整流程演示单页PPT。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import build_opening_template_pptx as base


OUT = Path("/opt/openclaw-hr/docs/小慧-完整流程演示-模板初版.pptx")
TOTAL = 1
GREEN = base.RGBColor(0x26, 0xD0, 0x98)
GREEN_DARK = base.RGBColor(0x08, 0x50, 0x57)
PURPLE = base.RGBColor(0x8C, 0x7C, 0xFF)
RED = base.RGBColor(0xF1, 0x5A, 0x66)


def arrow(slide, x, y, w, color=base.CYAN):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Inches(0.23))
    base.set_fill(shape, color)
    shape.line.fill.background()
    return shape


def task_card(slide, y, number: str, title: str, accent) -> None:
    base.rect(slide, Inches(0.66), y, Inches(2.38), Inches(0.52), base.NAVY_3, line=base.LINE, radius=True)
    base.text(slide, Inches(0.78), y + Inches(0.1), Inches(0.35), Inches(0.24), number, 10, bold=True, color=accent)
    base.text(slide, Inches(1.13), y + Inches(0.07), Inches(1.73), Inches(0.3), title, 11.5, bold=True, color=base.PALE)


def flow_lane(slide, y, module: str, output: str, mechanism: str, human_gate: str, accent) -> None:
    base.rect(slide, Inches(6.0), y, Inches(6.78), Inches(0.64), base.NAVY_2, line=base.LINE, radius=True)
    base.rect(slide, Inches(6.0), y, Inches(0.08), Inches(0.64), accent)
    base.text(slide, Inches(6.18), y + Inches(0.1), Inches(0.9), Inches(0.27), module, 12.5, bold=True, color=accent)
    base.text(slide, Inches(7.12), y + Inches(0.08), Inches(1.72), Inches(0.31), output, 11, bold=True, color=base.PALE)
    base.text(slide, Inches(8.92), y + Inches(0.08), Inches(2.06), Inches(0.31), mechanism, 9.5, color=base.MUTED)
    base.rect(slide, Inches(11.08), y + Inches(0.1), Inches(1.48), Inches(0.28), base.NAVY_3, line=accent, radius=True)
    base.text(
        slide,
        Inches(11.08),
        y + Inches(0.1),
        Inches(1.48),
        Inches(0.28),
        human_gate,
        9,
        bold=True,
        color=accent,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def build_slide(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    base.add_navigation(slide, 3, 1)
    base.add_title(
        slide,
        "04  FULL PROCESS",
        "完整流程联调：回应开场中的HR一天",
        "同一入口接收任务，Router分线处理，结果统一返回，关键决策仍由HR完成",
    )

    # 边界提示：这里展示五条独立任务，不制造自动串行的假象。
    base.rect(slide, Inches(0.48), Inches(2.02), Inches(12.36), Inches(0.38), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(
        slide,
        Inches(0.72),
        Inches(2.02),
        Inches(11.88),
        Inches(0.38),
        "五条独立请求  ·  一次只进入一个业务模块  ·  入职只做进度查询，不重复技能段审批演示",
        10.5,
        bold=True,
        color=base.CYAN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    # 左：开场中的任务队列
    base.rect(slide, Inches(0.48), Inches(2.58), Inches(2.75), Inches(3.5), base.NAVY_2, line=base.LINE, radius=True)
    base.text(slide, Inches(0.72), Inches(2.75), Inches(2.24), Inches(0.32), "HR任务队列", 15, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)
    tasks = [
        ("01", "筛选/检索候选人", base.CYAN),
        ("02", "查询员工1入职进度", GREEN),
        ("03", "准备指定培训资料", PURPLE),
        ("04", "汇总360绩效反馈", base.ORANGE),
        ("05", "行政问答或HR月报", RED),
    ]
    for i, values in enumerate(tasks):
        task_card(slide, Inches(3.2 + i * 0.55), *values)

    # 中：Router集中入口
    arrow(slide, Inches(3.4), Inches(4.07), Inches(0.5), base.CYAN)
    base.rect(slide, Inches(3.98), Inches(3.13), Inches(1.72), Inches(2.14), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(4.14), Inches(3.43), Inches(1.4), Inches(0.34), "小慧总控", 16, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(4.14), Inches(3.89), Inches(1.4), Inches(0.27), "判断领域", 11, bold=True, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(4.14), Inches(4.23), Inches(1.4), Inches(0.27), "派发专线", 11, bold=True, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(4.14), Inches(4.57), Inches(1.4), Inches(0.27), "汇总转述", 11, bold=True, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(4.14), Inches(4.94), Inches(1.4), Inches(0.2), "WebChat / CLI", 8.8, color=base.MUTED, align=PP_ALIGN.CENTER)
    arrow(slide, Inches(5.73), Inches(4.07), Inches(0.24), base.CYAN)

    # 右：五条并列业务专线，每条都同时展示产物、机制与人工节点。
    base.text(slide, Inches(6.0), Inches(2.62), Inches(1.5), Inches(0.27), "业务模块", 10, bold=True, color=base.MUTED)
    base.text(slide, Inches(7.12), Inches(2.62), Inches(1.6), Inches(0.27), "返回结果", 10, bold=True, color=base.MUTED)
    base.text(slide, Inches(8.92), Inches(2.62), Inches(2.0), Inches(0.27), "机制证据", 10, bold=True, color=base.MUTED)
    base.text(slide, Inches(11.12), Inches(2.62), Inches(1.4), Inches(0.27), "人工节点", 10, bold=True, color=base.MUTED, align=PP_ALIGN.CENTER)
    lanes = [
        ("招聘", "人才检索 / 画像", "脱敏 · 评分 · 数字回核", "HR决定录用", base.CYAN),
        ("入职", "员工1当前卡点", "只读查询回传", "审批见技能段", GREEN),
        ("培训", "点名培训资料", "素材边界 · 结构自检", "讲师审核", PURPLE),
        ("绩效", "360报告 / 看板", "校验 · 聚合 · 不定级", "主管校准", base.ORANGE),
        ("行政", "制度答复 / 月报", "身份 · 依据 · 脚本计算", "争议转人工", RED),
    ]
    for i, values in enumerate(lanes):
        flow_lane(slide, Inches(2.94 + i * 0.72), *values)

    # 四个联调验收点
    base.text(slide, Inches(0.5), Inches(6.26), Inches(1.25), Inches(0.27), "联调验收", 11, bold=True, color=base.MUTED)
    checks = ["正确分发", "机制正常执行", "产物成功返回", "人工决策未被越过"]
    for i, label in enumerate(checks):
        x = Inches(1.7 + i * 2.72)
        base.rect(slide, x, Inches(6.18), Inches(2.43), Inches(0.45), GREEN_DARK, line=GREEN, radius=True)
        base.text(
            slide,
            x,
            Inches(6.18),
            Inches(2.43),
            Inches(0.45),
            f"✓  {label}",
            11.5,
            bold=True,
            color=GREEN,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    base.text(
        slide,
        Inches(0.5),
        Inches(6.78),
        Inches(12.2),
        Inches(0.24),
        "现场只看“入口—分发—处理—返回—人工确认”。入职技能段已验证的HCM更新与Lobster审批门，本段不再重复。",
        9.5,
        color=base.MUTED,
        align=PP_ALIGN.CENTER,
    )

    base.add_notes(
        slide,
        "建议总时长6—8分钟。开场：前面的技能演示分别验证了各模块的核心功能和机制，"
        "接下来回到开场中HR同时面对的多项任务，通过同一入口连续调用五类能力。"
        "\n依次输入："
        "\n1. 从人才库中查找与员工1背景相似的候选人，并说明推荐依据。"
        "\n2. 查询员工1目前的入职进度和待处理事项。"
        "\n3. 查看入职流程培训的评估反馈表是否通过交付前自检。"
        "\n4. 生成员工1的360度绩效反馈看板，并说明主要共识和分歧。"
        "\n5. 生成2026年9月HR月报。"
        "\n一次只发一条。每项只验证Router是否正确分发、模块是否返回结果，不重复技能段代码。"
        "\n入职边界：本段只用员工1做只读查询；不修改HCM、不走审批门、不讲Lobster。"
        "员工2的HCM协同与两道人工门已在入职技能段完成。"
        "\n收口：五项业务不是Agent自动串行，而是HR发起的五条独立任务。"
        "小慧接住重复事务，但录用、审批、培训交付和绩效定级仍由人完成。随后进入结尾成果页。",
    )


def main() -> None:
    prs = Presentation(str(base.TEMPLATE))
    base.clear_slides(prs)
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL
    build_slide(prs)
    prs.core_properties.title = "小慧 · 完整流程演示"
    prs.core_properties.subject = "统一入口、Router分发、五模块结果回传与人工决策边界"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
