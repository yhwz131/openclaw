#!/usr/bin/env python3
"""入职模块七页：跨部门编排、Lobster 核心机制与 Cron 飞书证据。"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import build_opening_template_pptx as base


ROOT = Path("/opt/openclaw-hr/docs")
OUT = ROOT / "PPT准备/入职模块技能演示-7页跨部门编排版.pptx"
PREVIEW = ROOT / "PPT准备/入职模块技能演示-7页跨部门编排版-预览.png"
ASSETS = ROOT / "ppt/assets/onboarding"
DASHBOARD = ASSETS / "onboarding-dashboard.png"
DASHBOARD_CROP = ASSETS / "onboarding-dashboard-crop.png"
FEISHU_REVIEW = ASSETS / "feishu-review-day7-employee2.jpg"

TOTAL = 7
GREEN = RGBColor(0x26, 0xD0, 0x98)
GREEN_DARK = RGBColor(0x08, 0x50, 0x57)
PURPLE = RGBColor(0x8C, 0x7C, 0xFF)
RED = RGBColor(0xF1, 0x5A, 0x66)
RED_DARK = RGBColor(0x56, 0x17, 0x26)
CODE_BG = RGBColor(0x00, 0x11, 0x2C)


def connector(slide, x1, y1, x2, y2, color=base.CYAN, width=1.1, arrow=True):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def secondary_nav(slide) -> None:
    labels = ["总控", "招聘", "入职", "培训", "绩效", "行政"]
    base.rect(slide, Inches(0), Inches(0.72), base.W, Inches(0.45), base.NAVY, line=base.LINE)
    base.text(
        slide,
        Inches(0.38),
        Inches(0.79),
        Inches(0.9),
        Inches(0.27),
        "技能地图",
        10,
        bold=True,
        color=base.MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    for i, label in enumerate(labels):
        x = Inches(1.33 + i * 1.77)
        active = i == 2
        base.rect(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            base.CYAN if active else base.NAVY_2,
            line=base.CYAN if active else base.LINE,
            radius=True,
        )
        base.text(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            label,
            10.5,
            bold=active,
            color=base.NAVY if active else base.MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def header(slide, page: int, eyebrow: str, title: str, subtitle: str) -> None:
    base.add_navigation(slide, 2, page)
    secondary_nav(slide)
    base.text(slide, Inches(0.48), Inches(1.28), Inches(4.2), Inches(0.25), eyebrow, 10.5, bold=True, color=base.CYAN)
    base.text(slide, Inches(0.46), Inches(1.53), Inches(12.2), Inches(0.49), title, 22, bold=True, font=base.TITLE_FONT)
    base.rect(slide, Inches(0.47), Inches(2.08), Inches(0.75), Inches(0.05), base.CYAN)
    base.text(slide, Inches(1.4), Inches(1.99), Inches(11.1), Inches(0.29), subtitle, 11.2, color=base.MUTED)


def card(slide, x, y, w, h, title: str, body: str, *, accent=base.CYAN, fill=base.NAVY_2, body_size=11.2):
    base.rect(slide, x, y, w, h, fill, line=base.LINE, radius=True, transparency=3)
    base.rect(slide, x, y, Inches(0.08), h, accent)
    base.text(slide, x + Inches(0.18), y + Inches(0.12), w - Inches(0.36), Inches(0.28), title, 12.5, bold=True, color=accent)
    base.text(slide, x + Inches(0.18), y + Inches(0.5), w - Inches(0.36), h - Inches(0.62), body, body_size, color=base.PALE, line_spacing=0.9)


def crop_dashboard() -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.open(DASHBOARD).convert("RGB")
    w, h = image.size
    image.crop((0, 0, w, int(h * 0.62))).save(DASHBOARD_CROP, quality=95)
    return DASHBOARD_CROP


def slide_timeline(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        1,
        "03  ONBOARDING · TIMELINE",
        "入职不是清单，而是一条跨部门时间轴",
        "主命题：编排跨部门、长周期、多人负责的入职事项，不越权、不丢状态、不假完成",
    )

    phases = [
        ("入职前准备", "HCM 五部门节点", "合同 · 账号 · 设备\nBuddy · 培训", base.CYAN),
        ("报到审批", "Lobster 两道人工门", "合同门 · 账号权限门\nresumeToken 恢复", base.ORANGE),
        ("入职后跟进", "7 / 30 / 60 / 90", "到期 needs_review\nHR 回填后完成", GREEN),
    ]
    for i, (phase, owner, body, accent) in enumerate(phases):
        x = Inches(0.48 + i * 4.18)
        base.rect(slide, x, Inches(2.4), Inches(3.9), Inches(2.35), base.NAVY_2, line=accent, radius=True)
        base.rect(slide, x, Inches(2.4), Inches(3.9), Inches(0.42), base.NAVY_3)
        base.text(slide, x, Inches(2.4), Inches(3.9), Inches(0.42), phase, 14, bold=True, color=accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        base.text(slide, x + Inches(0.2), Inches(3.0), Inches(3.5), Inches(0.35), owner, 13.5, bold=True, color=base.PALE, align=PP_ALIGN.CENTER)
        base.text(slide, x + Inches(0.25), Inches(3.5), Inches(3.4), Inches(0.95), body, 12.5, color=base.MUTED, align=PP_ALIGN.CENTER)
        if i < 2:
            connector(slide, x + Inches(3.9), Inches(3.55), x + Inches(4.12), Inches(3.55), accent, 1.2)

    owners = [
        ("HCM", "业务事实", "不存审批 token"),
        ("Lobster", "审批顺序与暂停", "不替部门填写"),
        ("试用期跟进", "到期提醒与回填", "不决定录用"),
        ("入职小慧", "读取 · 催办 · 提醒", "不代批 · 不写 HCM"),
    ]
    for i, (name, a, b) in enumerate(owners):
        x = Inches(0.48 + i * 3.15)
        base.rect(slide, x, Inches(5.05), Inches(2.95), Inches(0.95), base.NAVY_3, line=base.LINE, radius=True)
        base.text(slide, x + Inches(0.15), Inches(5.18), Inches(2.65), Inches(0.25), name, 12, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(0.15), Inches(5.48), Inches(2.65), Inches(0.38), f"{a}\n{b}", 10.5, color=base.PALE)

    base.rect(slide, Inches(0.48), Inches(6.22), Inches(12.36), Inches(0.48), base.NAVY_2, line=base.CYAN, radius=True)
    base.text(
        slide,
        Inches(0.68),
        Inches(6.22),
        Inches(11.96),
        Inches(0.48),
        "本段直接进入入职 Agent｜完整流程联调再用 Router + 员工1 做一次进度查询，不重复审批机制",
        11,
        bold=True,
        color=base.CYAN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约50秒。招聘确认录用后才进入入职。入职是跨部门时间轴：HCM管准备事实，Lobster管两道人工门，"
        "7/30/60/90管入职后跟进。看板只汇总真实状态。本段不展示Router；完整流程段只做一次员工1查询。",
    )


def slide_why_hard(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        2,
        "03  WHY HARD · RISK BOUNDARY",
        "状态分散、责任分散，模型最容易说“已经办完了”",
        "法律不单独成页：嵌在对应风险旁，解释为什么要这样设计",
    )

    depts = ["合同签署", "IT 账号", "行政设备", "Buddy 匹配", "入职培训"]
    for i, name in enumerate(depts):
        x = Inches(0.48 + i * 2.5)
        base.rect(slide, x, Inches(2.4), Inches(2.28), Inches(0.48), base.NAVY_2, line=base.LINE, radius=True)
        base.text(slide, x, Inches(2.4), Inches(2.28), Inches(0.48), name, 12, bold=True, color=base.PALE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    card(
        slide,
        Inches(0.48),
        Inches(3.15),
        Inches(3.9),
        Inches(1.55),
        "三类状态源",
        "HCM｜部门是否真完成\nLobster｜审批是否可恢复\nreviews｜跟进是否回填",
        accent=base.CYAN,
        body_size=12,
    )
    card(
        slide,
        Inches(4.7),
        Inches(3.15),
        Inches(3.9),
        Inches(1.55),
        "假成功风险",
        "旧快照冒充实时进度\n把展示当成代替审批\n把审批通过当成部门完成",
        accent=RED,
        fill=RED_DARK,
        body_size=12,
    )
    card(
        slide,
        Inches(8.92),
        Inches(3.15),
        Inches(3.9),
        Inches(1.55),
        "设计原则",
        "业务事实只读 HCM\n关键决定写入 Lobster 门\n跟进完成必须 HR 回填",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=12,
    )

    rows = [
        ("《劳动合同法》第八、十条", "书面合同与告知义务", "合同节点守卫 + 合同审批门", base.ORANGE),
        ("现行《网络安全法》第二十三条", "防止未经授权访问", "账号权限门 + 材料不足不建议批", GREEN),
        ("《个人信息保护法》第六、十三条", "目的限定、最小必要", "HCM 只 GET · 不可用即停止", base.CYAN),
    ]
    for i, (law, risk, mech, accent) in enumerate(rows):
        y = Inches(4.95 + i * 0.52)
        base.rect(slide, Inches(0.48), y, Inches(12.36), Inches(0.44), base.NAVY_2, line=base.LINE, radius=True)
        base.rect(slide, Inches(0.48), y, Inches(0.08), Inches(0.44), accent)
        base.text(slide, Inches(0.72), y, Inches(3.5), Inches(0.44), law, 11, bold=True, color=accent, anchor=MSO_ANCHOR.MIDDLE)
        base.text(slide, Inches(4.3), y, Inches(3.2), Inches(0.44), risk, 11, color=base.PALE, anchor=MSO_ANCHOR.MIDDLE)
        base.text(slide, Inches(7.7), y, Inches(4.8), Inches(0.44), mech, 11, bold=True, color=base.PALE, anchor=MSO_ANCHOR.MIDDLE)

    base.add_notes(
        slide,
        "约70秒。五个责任部门状态分散。三类状态源必须拆开。法律嵌在风险旁："
        "劳动合同法对应合同门，网络安全法对应账号门，个人信息保护法对应只读与失败关闭。"
        "两道门是项目风险控制设计，不是法律原句。",
    )


def slide_waiting_departments(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        3,
        "03  LIVE · WAITING DEPARTMENTS",
        "合同未开始：拒绝启动，等待真实部门更新",
        "现场输入：请为员工2启动入职流程，岗位前端工程师，薪资30K",
    )

    steps = [
        ("01", "合同 = pending", "start 退出 6", RED),
        ("02", "记 waiting_departments", "不调用 Lobster", base.ORANGE),
        ("03", "队友更新 HCM", "合同 → in_progress", base.CYAN),
        ("04", "再次 start", "只读快照后进门", GREEN),
    ]
    for i, (num, title, sub, accent) in enumerate(steps):
        x = Inches(0.48 + i * 3.15)
        base.rect(slide, x, Inches(2.4), Inches(2.95), Inches(1.2), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x + Inches(0.15), Inches(2.55), Inches(0.5), Inches(0.25), num, 11, bold=True, color=accent)
        base.text(slide, x + Inches(0.15), Inches(2.9), Inches(2.65), Inches(0.28), title, 13, bold=True)
        base.text(slide, x + Inches(0.15), Inches(3.25), Inches(2.65), Inches(0.22), sub, 11, color=base.MUTED)
        if i < 3:
            connector(slide, x + Inches(2.95), Inches(3.0), x + Inches(3.1), Inches(3.0), accent, 1.0)

    card(
        slide,
        Inches(0.48),
        Inches(3.95),
        Inches(6.0),
        Inches(2.15),
        "系统真实反应",
        "人事系统合同未齐 → 明确失败\n记下等待的责任部门\n催办对应部门\n小慧不改人事主数据",
        accent=RED,
        fill=RED_DARK,
        body_size=13,
    )
    card(
        slide,
        Inches(6.78),
        Inches(3.95),
        Inches(6.06),
        Inches(2.15),
        "协同原则",
        "“等待部门”是合法状态，不是演示失败\n本地快照只服务当次流程\n部门更新后才能进入 Lobster\n跨部门协同 = 读真实状态，不替人执行",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=12.5,
    )
    base.add_notes(
        slide,
        "约2分钟。先用员工2启动，合同pending时退出6并记waiting_departments，不调Lobster。"
        "队友在HCM把合同改为in_progress后再次启动，只读快照再进Lobster。"
        "强调等待部门是合法状态，小慧不写HCM。",
    )


def slide_lobster_core(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        4,
        "03  CORE · LOBSTER HITL",
        "Lobster：把人工决策写成不可跳过的节点",
        "现场验证：start → contract_gate halt → resume → 从暂停点恢复",
    )

    chain = ["入职 Skill", "run-onboarding.sh", "Lobster 引擎", "hr-onboarding.lobster"]
    for i, label in enumerate(chain):
        x = Inches(0.48 + i * 3.15)
        accent = base.ORANGE if i == 2 else base.CYAN
        base.rect(slide, x, Inches(2.4), Inches(2.95), Inches(0.55), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x, Inches(2.4), Inches(2.95), Inches(0.55), label, 12.5, bold=True, color=accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 3:
            connector(slide, x + Inches(2.95), Inches(2.67), x + Inches(3.1), Inches(2.67), accent, 1.0)

    nodes = [("prep", base.CYAN), ("contract_gate", base.ORANGE), ("account_gate", base.ORANGE), ("buddy", base.CYAN), ("complete", GREEN)]
    for i, (name, accent) in enumerate(nodes):
        x = Inches(0.7 + i * 2.4)
        hot = i in (1, 2)
        base.rect(slide, x, Inches(3.2), Inches(2.1), Inches(0.42), base.NAVY_3 if hot else base.NAVY_2, line=accent, radius=True)
        base.text(slide, x, Inches(3.2), Inches(2.1), Inches(0.42), name, 11, bold=hot, color=accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 4:
            connector(slide, x + Inches(2.1), Inches(3.41), x + Inches(2.35), Inches(3.41), base.LINE, 0.9)

    base.rect(slide, Inches(0.48), Inches(3.9), Inches(6.1), Inches(2.35), CODE_BG, line=base.LINE, radius=True)
    base.text(slide, Inches(0.7), Inches(4.05), Inches(5.7), Inches(0.28), "hr-onboarding.lobster｜核心语义", 12, bold=True, color=base.CYAN)
    code = (
        "- id: contract_gate\n"
        "  run: build-approval-prompt.py\n"
        "  approval: true          # 自动 halt\n\n"
        "- id: complete\n"
        "  when: contract && account approved"
    )
    base.text(slide, Inches(0.7), Inches(4.45), Inches(5.6), Inches(1.55), code, 11.2, color=base.PALE, font="Courier New", line_spacing=0.88)

    card(
        slide,
        Inches(6.9),
        Inches(3.9),
        Inches(5.94),
        Inches(2.35),
        "现场必须看见",
        "status = needs_approval\ncurrent_step = contract_gate\n审批正文原样展示\n本次专属 resumeToken\nresume 后从暂停点继续，不重跑整链",
        accent=base.ORANGE,
        body_size=12.2,
    )
    base.add_notes(
        slide,
        "约2分钟。调用链：Skill→启动器→Lobster→工作流文件。"
        "run生成审批正文，approval true强制halt，when限制后续节点。"
        "现场看needs_approval、正文和resumeToken；批准后用token从暂停点恢复。"
        "状态不依赖模型记忆，审批节点不能跳过。",
    )


def slide_second_gate(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        5,
        "03  SECOND GATE · DUAL THRESHOLD",
        "账号权限门＋双重完成：拦住两类假成功",
        "新门生成新 token；审批完成 ≠ 部门准备完成 ≠ 跟进完成",
    )

    card(
        slide,
        Inches(0.48),
        Inches(2.4),
        Inches(4.0),
        Inches(2.0),
        "第二道｜账号权限门",
        "审批人：IT 安全负责人\n批准 → workflow complete\n拒绝 → cancelled\n材料不足：不建议批准",
        accent=base.ORANGE,
        body_size=12.5,
    )
    card(
        slide,
        Inches(4.72),
        Inches(2.4),
        Inches(4.0),
        Inches(2.0),
        "token 硬规则",
        "每道门新 resumeToken\n旧 token 不可复用\napprovalId 不能代替\n正文不可改写",
        accent=base.CYAN,
        body_size=12.5,
    )
    card(
        slide,
        Inches(8.96),
        Inches(2.4),
        Inches(3.86),
        Inches(2.0),
        "假成功 A",
        "两门通过 ≠ 五节点完成\n账号/设备/Buddy/培训\n仍可能待办",
        accent=RED,
        fill=RED_DARK,
        body_size=12.5,
    )

    base.rect(slide, Inches(0.48), Inches(4.7), Inches(3.5), Inches(1.2), base.NAVY_3, line=base.ORANGE, radius=True)
    base.text(slide, Inches(0.68), Inches(4.9), Inches(3.1), Inches(0.35), "Lobster 两门批准", 14, bold=True, color=base.ORANGE, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(0.68), Inches(5.3), Inches(3.1), Inches(0.3), "workflow = complete", 11, color=base.PALE, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(4.1), Inches(5.05), Inches(0.45), Inches(0.35), "＋", 22, bold=True, color=base.PALE, align=PP_ALIGN.CENTER)
    base.rect(slide, Inches(4.65), Inches(4.7), Inches(3.5), Inches(1.2), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(4.85), Inches(4.9), Inches(3.1), Inches(0.35), "HCM 五节点全绿", 14, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(4.85), Inches(5.3), Inches(3.1), Inches(0.3), "all nodes completed", 11, color=base.PALE, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(8.25), Inches(5.05), Inches(0.45), Inches(0.35), "→", 22, bold=True, color=base.PALE, align=PP_ALIGN.CENTER)
    base.rect(slide, Inches(8.8), Inches(4.7), Inches(4.02), Inches(1.2), GREEN_DARK, line=GREEN, radius=True)
    base.text(slide, Inches(9.0), Inches(4.9), Inches(3.62), Inches(0.35), "开通 7/30/60/90", 14, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(9.0), Inches(5.3), Inches(3.62), Inches(0.3), "sync-reviews-if-ready", 11, color=base.PALE, align=PP_ALIGN.CENTER)

    base.rect(slide, Inches(0.48), Inches(6.15), Inches(12.36), Inches(0.5), base.NAVY_2, line=GREEN, radius=True)
    base.text(
        slide,
        Inches(0.68),
        Inches(6.15),
        Inches(11.96),
        Inches(0.5),
        "现场验证：任一节点未完成 → sync-reviews-if-ready 跳过，不挂 Cron",
        12,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约1分40秒。合同门通过后进入账号权限门，新token。IT安全负责人决定。"
        "两门批准只是workflow complete；五节点全绿才开通跟进。未全绿不挂Cron。",
    )


def slide_cron_feishu(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        6,
        "03  CRON · FEISHU DELIVERY",
        "到期不是自动完成：Cron主动催办，HR回填闭环",
        "一次性定时任务在隔离会话运行，飞书只负责通知，不直接修改跟进结果",
    )

    base.text(slide, Inches(0.48), Inches(2.4), Inches(6.1), Inches(0.28), "触发链路｜五节点全绿后才注册", 12.5, bold=True, color=base.CYAN)
    steps = [
        ("01", "五节点全绿", "sync-reviews-if-ready"),
        ("02", "注册一次性 Cron", "7 / 30 / 60 / 90"),
        ("03", "到期触发 Agent", "隔离会话运行 review"),
        ("04", "飞书 announce", "推送重点与待回填"),
        ("05", "HR 确认后回填", "CLI record 才完成"),
    ]
    for i, (num, title, sub) in enumerate(steps):
        y = Inches(2.82 + i * 0.69)
        accent = GREEN if i == 4 else base.CYAN
        base.rect(slide, Inches(0.48), y, Inches(5.9), Inches(0.56), base.NAVY_2, line=accent, radius=True)
        base.text(slide, Inches(0.66), y + Inches(0.1), Inches(0.42), Inches(0.22), num, 9.8, bold=True, color=accent)
        base.text(slide, Inches(1.15), y + Inches(0.07), Inches(2.08), Inches(0.26), title, 11.5, bold=True, color=base.PALE)
        base.text(slide, Inches(3.25), y + Inches(0.07), Inches(2.85), Inches(0.26), sub, 10.3, color=base.MUTED, align=PP_ALIGN.RIGHT)
        if i < 4:
            connector(slide, Inches(3.42), y + Inches(0.56), Inches(3.42), y + Inches(0.68), base.LINE, 0.8)

    base.rect(slide, Inches(0.48), Inches(6.39), Inches(5.9), Inches(0.39), GREEN_DARK, line=GREEN, radius=True)
    base.text(
        slide,
        Inches(0.68),
        Inches(6.39),
        Inches(5.5),
        Inches(0.39),
        "边界：推送成功 ≠ 跟进完成｜尚未回填保持 due",
        10.5,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    base.rect(slide, Inches(6.68), Inches(2.4), Inches(6.16), Inches(4.38), base.NAVY_2, line=GREEN, radius=True)
    base.text(slide, Inches(6.9), Inches(2.55), Inches(5.72), Inches(0.3), "真实推送证据｜员工2 · 第7天", 12.5, bold=True, color=GREEN)
    if FEISHU_REVIEW.exists():
        image = Image.open(FEISHU_REVIEW).convert("RGB")
        w, h = image.size
        cropped = ASSETS / "feishu-review-day7-employee2-crop.jpg"
        image.crop((0, 45, w, h - 20)).save(cropped, quality=95)
        slide.shapes.add_picture(str(cropped), Inches(8.02), Inches(2.94), height=Inches(3.65))
    base.text(
        slide,
        Inches(6.91),
        Inches(6.48),
        Inches(5.72),
        Inches(0.22),
        "HCM：onboarding｜五节点：completed｜Review：due",
        9.4,
        bold=True,
        color=base.MUTED,
        align=PP_ALIGN.CENTER,
    )
    base.add_notes(
        slide,
        "约1分钟。五节点全绿后，系统注册7、30、60、90天四个一次性Cron。"
        "到期任务在隔离会话运行onboarding-review脚本，并通过announce投递飞书。"
        "截图显示员工2第7天、HCM五节点全部completed、Review为due。"
        "这证明推送链路已运行，但推送成功不等于跟进完成；只有HR真实沟通并执行CLI record后才完成。"
        "补充异常守卫：员工总体状态为cancelled时，不新开通跟进，也不输出needs_review催办。",
    )


def slide_close_loop(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        7,
        "03  FOLLOW-UP · EVIDENCE",
        "跟进只提醒，完成靠回填；所有结论都能追溯",
        "本页收束证据；自然语言查询留给完整流程联调，避免重复",
    )

    card(
        slide,
        Inches(0.48),
        Inches(2.4),
        Inches(4.2),
        Inches(2.55),
        "入职后跟进边界",
        "7 / 30 / 60 / 90 天\n到期只提醒，不自动完成\n人事确认沟通后回填\n查询与看板只读",
        accent=GREEN,
        body_size=13,
    )

    base.rect(slide, Inches(4.95), Inches(2.4), Inches(7.87), Inches(2.55), base.NAVY_2, line=base.CYAN, radius=True)
    base.text(slide, Inches(5.15), Inches(2.55), Inches(7.45), Inches(0.3), "真实证据｜看板 + 审批事件", 13, bold=True, color=base.CYAN)
    if DASHBOARD.exists():
        path = crop_dashboard()
        slide.shapes.add_picture(str(path), Inches(5.15), Inches(2.95), width=Inches(4.55), height=Inches(1.75))
    events = "合同门    暂停\n合同门    已批准\n账号门    暂停\n账号门    已批准\n入职准备  完成"
    base.text(slide, Inches(9.9), Inches(3.05), Inches(2.7), Inches(1.55), events, 10.5, color=base.PALE, font="Courier New", line_spacing=0.88)

    base.rect(slide, Inches(0.48), Inches(5.2), Inches(12.36), Inches(1.45), base.NAVY_3, line=base.ORANGE, radius=True)
    base.text(slide, Inches(0.72), Inches(5.4), Inches(3.0), Inches(0.3), "交接完整流程联调", 14, bold=True, color=base.ORANGE)
    base.text(
        slide,
        Inches(0.72),
        Inches(5.85),
        Inches(11.8),
        Inches(0.55),
        "后面从总控入口查询员工1入职进度：只验证正确分发与结果返回。\n"
        "不再修改人事主数据、不再走两道审批门、不再展开编排细节。",
        13,
        color=base.PALE,
    )
    base.add_notes(
        slide,
        "约50秒。跟进到期只提醒，人事回填才完成。展示看板和审批事件。"
        "收口并交接：完整流程段用员工1做一次查询，验证总控分发，不重复本段机制。",
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL
    slide_timeline(prs)
    slide_why_hard(prs)
    slide_waiting_departments(prs)
    slide_lobster_core(prs)
    slide_second_gate(prs)
    slide_cron_feishu(prs)
    slide_close_loop(prs)
    prs.core_properties.title = "小慧 · 入职模块技能演示跨部门编排版"
    prs.core_properties.subject = "跨部门时间轴、等待部门、Lobster人工门、双重完成与跟进闭环"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
