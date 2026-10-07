#!/usr/bin/env python3
"""基于“26-科技教育”模板生成小慧开头四页初版。"""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path("/opt/openclaw-hr/docs")
TEMPLATE = ROOT / "26-科技教育(PPT模板）.pptx"
OUT = ROOT / "小慧-开头部分-模板初版.pptx"
POLICY_DIR = ROOT / "ppt/assets/policy"
COVER_BG = ROOT / "ppt/assets/cover/xiaohui-cover-hr.png"

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x02, 0x16, 0x3B)
NAVY_2 = RGBColor(0x03, 0x2B, 0x5C)
NAVY_3 = RGBColor(0x04, 0x45, 0x7D)
CYAN = RGBColor(0x00, 0xD8, 0xFF)
BLUE = RGBColor(0x16, 0x83, 0xD8)
PALE = RGBColor(0xE9, 0xF8, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x9E, 0xC2, 0xD8)
LINE = RGBColor(0x12, 0x68, 0xA1)
ORANGE = RGBColor(0xFF, 0xA4, 0x21)
RED = RGBColor(0xE9, 0x4F, 0x4F)
FONT = "微软雅黑"
TITLE_FONT = "优设标题黑"
NAV_ITEMS = ["项目背景", "方案规划", "技能演示", "完整流程", "成果展望"]
TOTAL = 4


def set_alpha(color_element, value: int) -> None:
    """设置 OOXML 颜色透明度，value 取 0..100000。"""
    for old in color_element.findall(qn("a:alpha")):
        color_element.remove(old)
    alpha = etree.SubElement(color_element, qn("a:alpha"))
    alpha.set("val", str(value))


def set_fill(shape, color: RGBColor, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if transparency:
        solid = shape.fill._xPr.solidFill
        color_el = next(iter(solid), None)
        if color_el is not None:
            set_alpha(color_el, max(0, 100000 - transparency * 1000))


def set_run(run, size: float, bold=False, color=WHITE, name=FONT) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", name)


def text(
    slide,
    left,
    top,
    width,
    height,
    content,
    size,
    *,
    bold=False,
    color=WHITE,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    margin=0.02,
    font=FONT,
    line_spacing=1.0,
):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    for i, line in enumerate(str(content).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color, font)
    return shape


def rect(slide, left, top, width, height, fill, *, line=None, radius=False, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    set_fill(shape, fill, transparency)
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.0):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    sh.line.color.rgb = color
    sh.line.width = Pt(width)
    return sh


def add_notes(slide, content: str) -> None:
    slide.notes_slide.notes_text_frame.text = content


def clear_slides(prs: Presentation) -> None:
    ids = prs.slides._sldIdLst
    for slide_id in list(ids):
        prs.part.drop_rel(slide_id.rId)
        ids.remove(slide_id)


def new_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    return slide


def add_tech_background(slide) -> None:
    # 模仿模板中的蓝色光带和线路，不使用泛化机器人素材。
    rect(slide, Inches(0), Inches(0), W, H, NAVY)
    for y in (1.55, 3.62, 6.56):
        line(slide, Inches(0), Inches(y), W, Inches(y - 0.52), NAVY_3, 0.8)
    for x in (0.45, 3.15, 7.15, 11.7):
        line(slide, Inches(x), Inches(0.72), Inches(x + 0.75), Inches(7.18), NAVY_2, 0.7)
    for x, y in ((0.6, 1.48), (3.55, 2.92), (7.65, 1.77), (12.1, 4.8), (9.3, 6.08)):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.08), Inches(0.08))
        set_fill(dot, CYAN)
        dot.line.fill.background()


def add_navigation(slide, active: int, page: int) -> None:
    rect(slide, Inches(0), Inches(0), W, Inches(0.72), NAVY_2, line=LINE)
    text(
        slide,
        Inches(0.35),
        Inches(0.11),
        Inches(2.2),
        Inches(0.42),
        "小慧 · HR数字员工",
        13.5,
        bold=True,
        font=FONT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    start_x, item_w = 2.65, 1.75
    for i, label in enumerate(NAV_ITEMS):
        x = Inches(start_x + i * item_w)
        shape = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, x, Inches(0.12), Inches(1.67), Inches(0.39))
        set_fill(shape, CYAN if i == active else NAVY_3, 0 if i == active else 18)
        shape.line.color.rgb = CYAN if i == active else LINE
        shape.line.width = Pt(1)
        text(
            slide,
            x,
            Inches(0.12),
            Inches(1.67),
            Inches(0.39),
            label,
            12,
            bold=i == active,
            color=NAVY if i == active else MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    text(
        slide,
        Inches(11.72),
        Inches(0.12),
        Inches(1.15),
        Inches(0.38),
        f"{page:02d} / {TOTAL:02d}",
        11,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def fit_picture(slide, path: Path, left, top, box_w, box_h):
    """把官网截图等比放入框内，不拉伸。"""
    from PIL import Image

    im = Image.open(path)
    ar = im.width / float(im.height)
    box_ar = box_w / box_h
    if ar >= box_ar:
        w, h = box_w, box_w / ar
    else:
        h, w = box_h, box_h * ar
    x = left + (box_w - w) / 2
    y = top + (box_h - h) / 2
    return slide.shapes.add_picture(str(path), int(x), int(y), width=int(w), height=int(h))


def add_title(slide, eyebrow: str, title_value: str, subtitle: str = "") -> None:
    text(slide, Inches(0.48), Inches(0.93), Inches(2.4), Inches(0.28), eyebrow, 11, bold=True, color=CYAN)
    text(slide, Inches(0.46), Inches(1.17), Inches(12.1), Inches(0.53), title_value, 25, bold=True, font=TITLE_FONT)
    rect(slide, Inches(0.47), Inches(1.77), Inches(0.75), Inches(0.05), CYAN)
    if subtitle:
        text(slide, Inches(1.4), Inches(1.64), Inches(11.3), Inches(0.34), subtitle, 12.5, color=MUTED)


def add_cover_background(source: Presentation, slide) -> None:
    """封面/致谢底板：HR 盖章 + 光态分身。左侧压字，右侧少压色。"""
    if COVER_BG.exists():
        slide.shapes.add_picture(str(COVER_BG), 0, 0, width=W, height=H)
    else:
        for shape in source.slides[0].shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width >= W * 0.95 and shape.height >= H * 0.95:
                slide.shapes.add_picture(BytesIO(shape.image.blob), 0, 0, width=W, height=H)
    # 参数越大越透。全幅只薄罩一层，标题区再压一块，右侧人+分身保持可见。
    rect(slide, Inches(0), Inches(0), W, H, NAVY, transparency=78)
    rect(slide, Inches(0), Inches(0), Inches(6.4), H, NAVY, transparency=48)
    rect(slide, Inches(0), Inches(0), Inches(0.12), H, CYAN)


def cover(prs: Presentation, source: Presentation) -> None:
    slide = new_slide(prs)
    add_cover_background(source, slide)
    text(
        slide,
        Inches(0.68),
        Inches(0.46),
        Inches(11.9),
        Inches(0.36),
        "2026 云南省职业院校技能大赛  ·  人工智能赛道  ·  高职组",
        12,
        color=PALE,
    )
    rect(slide, Inches(0.7), Inches(1.33), Inches(1.0), Inches(0.08), CYAN)
    text(slide, Inches(0.68), Inches(1.58), Inches(7.4), Inches(0.95), "小 慧", 50, bold=True, font=TITLE_FONT)
    text(slide, Inches(0.7), Inches(2.52), Inches(7.7), Inches(0.55), "HR 数字员工", 26, bold=True, color=CYAN)
    text(
        slide,
        Inches(0.7),
        Inches(3.25),
        Inches(9.1),
        Inches(0.8),
        "一人决策  +  Agent辅助",
        25,
        bold=True,
        font=FONT,
    )
    text(
        slide,
        Inches(0.72),
        Inches(4.08),
        Inches(8.5),
        Inches(0.6),
        "让重复事务交给智能体，让关键决策始终掌握在人手中",
        15,
        color=PALE,
    )
    labels = ["招聘", "入职", "培训", "绩效", "行政"]
    for i, label in enumerate(labels):
        x = Inches(0.72 + i * 1.35)
        rect(slide, x, Inches(5.0), Inches(1.15), Inches(0.42), NAVY_2, line=CYAN, radius=True, transparency=12)
        text(slide, x, Inches(5.0), Inches(1.15), Inches(0.42), label, 12, color=PALE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # 首页也给评委一张完整路线图。
    rect(slide, Inches(0), Inches(6.65), W, Inches(0.85), NAVY_2, transparency=5)
    for i, label in enumerate(NAV_ITEMS):
        x = Inches(0.78 + i * 2.48)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, Inches(6.91), Inches(0.17), Inches(0.17))
        set_fill(dot, CYAN if i == 0 else BLUE)
        dot.line.fill.background()
        text(slide, x + Inches(0.27), Inches(6.77), Inches(1.75), Inches(0.42), label, 11.5, bold=i == 0, color=WHITE if i == 0 else MUTED, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(NAV_ITEMS) - 1:
            line(slide, x + Inches(1.82), Inches(6.99), x + Inches(2.32), Inches(6.99), LINE, 1.2)
    add_notes(
        slide,
        "情景演示结束后切入本页。口播：这就是我们研发HR数字员工小慧的目的。"
        "我们采用一人决策、Agent辅助的方案，让Agent处理筛简历、追入职、追培训、绩效整理、月报和行政问答等重复事务，"
        "并将结果如实汇报，辅助HR安排工作和进行决策。",
    )


def policy(prs: Presentation) -> None:
    slide = new_slide(prs)
    add_tech_background(slide)
    add_navigation(slide, 0, 2)
    add_title(slide, "01  PROJECT BACKGROUND", "政策驱动：从“人工智能+”到智能体规范应用", "政策给出发展方向，规范文件划定人机决策边界")

    cards = [
        (
            "2025.08",
            "国务院“人工智能+”行动",
            "到2027年，新一代智能终端、智能体等应用普及率超70%",
            "国发〔2025〕11号",
            POLICY_DIR / "书证-国务院-卡片.png",
        ),
        (
            "2026.03",
            "国家“十五五”规划纲要",
            "全面实施“人工智能+”行动，鼓励智能体等技术创新",
            "第十三章",
            POLICY_DIR / "书证-十五五-卡片.png",
        ),
        (
            "2026.05",
            "云南省数字产业实施方案",
            "构建重点领域智能体，鼓励人工智能OPC创新发展",
            "云政办发〔2026〕27号",
            POLICY_DIR / "书证-云南-卡片.png",
        ),
    ]
    for i, (date, heading, body, source, shot) in enumerate(cards):
        x = Inches(0.48 + i * 4.18)
        rect(slide, x, Inches(2.02), Inches(3.83), Inches(2.28), NAVY_2, line=LINE, radius=True, transparency=8)
        rect(slide, x + Inches(0.08), Inches(2.1), Inches(3.67), Inches(1.08), WHITE)
        fit_picture(slide, shot, x + Inches(0.08), Inches(2.1), Inches(3.67), Inches(1.08))
        text(slide, x + Inches(0.16), Inches(3.22), Inches(0.9), Inches(0.28), date, 11, bold=True, color=CYAN, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + Inches(1.08), Inches(3.22), Inches(2.55), Inches(0.28), heading, 13, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + Inches(0.16), Inches(3.5), Inches(3.5), Inches(0.52), body, 12.5, color=PALE)
        text(slide, x + Inches(0.16), Inches(4.02), Inches(3.5), Inches(0.22), source, 9.5, color=MUTED)
        if i < 2:
            line(slide, x + Inches(3.83), Inches(3.16), x + Inches(4.18), Inches(3.16), CYAN, 1.5)

    rect(slide, Inches(0.48), Inches(4.42), Inches(12.36), Inches(2.38), NAVY_2, line=CYAN, radius=True, transparency=2)
    rect(slide, Inches(0.48), Inches(4.42), Inches(0.12), Inches(2.38), CYAN)
    text(slide, Inches(0.78), Inches(4.5), Inches(2.3), Inches(0.32), "2026.05  应用规范", 12, bold=True, color=CYAN, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(3.0), Inches(4.48), Inches(8.9), Inches(0.36), "《智能体规范应用与创新发展实施意见》", 16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, Inches(0.68), Inches(4.88), Inches(7.95), Inches(1.08), WHITE)
    fit_picture(
        slide,
        POLICY_DIR / "书证-实施意见-卡片.png",
        Inches(0.68),
        Inches(4.88),
        Inches(7.95),
        Inches(1.08),
    )
    text(
        slide,
        Inches(0.72),
        Inches(6.02),
        Inches(7.9),
        Inches(0.52),
        "“厘清仅限用户本人决策、需由用户授权决策和智能体自主决策等各种决策方式的合理边界”\n"
        "“确保用户对智能体自主决策享有知情权和最终决策权”",
        12.5,
        bold=True,
        color=PALE,
        line_spacing=1.02,
    )
    # 右侧直接落到项目机制，避免政策页只剩文件罗列。
    rect(slide, Inches(8.84), Inches(4.88), Inches(3.72), Inches(1.62), NAVY_3, line=LINE, radius=True)
    text(slide, Inches(9.02), Inches(5.0), Inches(3.36), Inches(0.28), "小慧的人机边界", 11, bold=True, color=CYAN)
    text(slide, Inches(9.02), Inches(5.34), Inches(3.36), Inches(0.95), "Agent：理解、执行、汇报\nHR：授权、复核、最终决策", 14, bold=True)
    text(
        slide,
        Inches(0.72),
        Inches(6.52),
        Inches(11.8),
        Inches(0.2),
        "来源：中国政府网｜国家发展改革委｜云南省人民政府｜中央网信办",
        9,
        color=MUTED,
    )
    add_notes(
        slide,
        "口播控制在60秒内。前三张卡片只说明发展方向。重点指向下方规范："
        "项目据此把仅限HR决策、需HR授权和Agent可自主执行的事项分开。"
        "必须说“知情权和最终决策权”，不要说“叫停权”。"
        "书证是官网抬头截图，含国徽/站点身份和文件标题，不要念截图全文。"
        "\n来源：https://www.gov.cn/zhengce/zhengceku/202508/content_7037862.htm；"
        "https://www.gov.cn/yaowen/liebiao/202603/content_7062633.htm；"
        "https://www.yn.gov.cn/zwgk/zcwj/zxwj/202605/t20260520_326925.html；"
        "https://www.cac.gov.cn/2026-05/08/c_1779979789523320.htm",
    )


def module_card(slide, x, y, code, name, items, color=CYAN) -> None:
    rect(slide, x, y, Inches(3.88), Inches(2.13), NAVY_2, line=LINE, radius=True, transparency=5)
    rect(slide, x, y, Inches(0.1), Inches(2.13), color)
    text(slide, x + Inches(0.22), y + Inches(0.16), Inches(0.62), Inches(0.29), code, 10.5, bold=True, color=color)
    text(slide, x + Inches(0.88), y + Inches(0.1), Inches(2.74), Inches(0.42), name, 15, bold=True)
    text(slide, x + Inches(0.23), y + Inches(0.63), Inches(3.38), Inches(1.2), "\n".join(f"· {item}" for item in items), 13, color=PALE, line_spacing=0.86)


def research(prs: Presentation) -> None:
    slide = new_slide(prs)
    add_tech_background(slide)
    add_navigation(slide, 0, 3)
    add_title(slide, "01  RESEARCH FINDINGS", "HR岗位调研：6大模块，30+业务子项", "从岗位真实工作内容出发，再决定Agent如何分工")
    modules = [
        ("01", "招聘 Recruitment", ["JD撰写发布", "简历收发初筛", "人才画像比对", "面试调度纪要", "Offer制作跟进"]),
        ("02", "入职 Onboarding", ["合同与资料准备", "账号设备申请", "政策介绍", "入职培训", "试用期跟进"]),
        ("03", "培训发展", ["需求调研", "计划制定", "内容生成", "实施记录", "效果测评"]),
        ("04", "绩效管理", ["OKR/KPI设定", "周期评估", "360反馈", "面谈纪要", "改进跟踪"]),
        ("05", "薪酬福利", ["薪资核算", "调薪审批", "社保公积金", "个税申报", "离职结算"]),
        ("06", "员工关系与行政", ["会议与内部通知", "HR Helpdesk", "考勤异常", "档案管理", "数据看板"]),
    ]
    positions = [
        (Inches(0.48), Inches(2.03)),
        (Inches(4.73), Inches(2.03)),
        (Inches(8.98), Inches(2.03)),
        (Inches(0.48), Inches(4.38)),
        (Inches(4.73), Inches(4.38)),
        (Inches(8.98), Inches(4.38)),
    ]
    for (code, name, items), (x, y) in zip(modules, positions):
        module_card(slide, x, y, code, name, items, ORANGE if code == "05" else CYAN)
    text(slide, Inches(0.5), Inches(6.72), Inches(7.5), Inches(0.3), "调研结论：高频重复事务可由Agent辅助，高风险决策必须保留人工环节", 10.5, bold=True, color=CYAN)
    text(slide, Inches(8.2), Inches(6.72), Inches(4.5), Inches(0.3), "30项典型事项节选｜完整清单见调研材料", 9.5, color=MUTED, align=PP_ALIGN.RIGHT)
    add_notes(
        slide,
        "口播：在与合作企业HR的调研中，我们按岗位工作内容归纳出六个模块。"
        "本页展示30项典型事项，完整调研清单还有渠道运营、人才库维护、Buddy匹配、发展路径等内容。"
        "这一步的目的不是堆功能，而是为后面的Agent模块划分提供依据。",
    )


def agent_card(slide, x, name, body, tag) -> None:
    rect(slide, x, Inches(4.17), Inches(2.35), Inches(1.62), NAVY_2, line=LINE, radius=True, transparency=3)
    text(slide, x + Inches(0.14), Inches(4.3), Inches(2.05), Inches(0.36), name, 15, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    text(slide, x + Inches(0.16), Inches(4.76), Inches(2.02), Inches(0.52), body, 12.5, color=PALE, align=PP_ALIGN.CENTER)
    rect(slide, x + Inches(0.38), Inches(5.42), Inches(1.58), Inches(0.25), NAVY_3, radius=True)
    text(slide, x + Inches(0.38), Inches(5.42), Inches(1.58), Inches(0.25), tag, 9, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def scheme(prs: Presentation) -> None:
    slide = new_slide(prs)
    add_tech_background(slide)
    add_navigation(slide, 1, 4)
    add_title(slide, "02  SOLUTION PLANNING", "一期方案规划：1协调 + 5模块", "把调研得到的岗位事项，拆成职责清晰、边界可控的Agent模块")

    # 概念来源与项目概念分开标明，防止把 OPD 说成政策词。
    rect(slide, Inches(0.48), Inches(2.0), Inches(12.36), Inches(0.7), NAVY_2, line=LINE, radius=True)
    steps = [("政策场景", "OPC · 一人公司"), ("概念延伸", "从公司到部门"), ("项目目标", "OPD · 一人部门")]
    for i, (a, b) in enumerate(steps):
        x = Inches(0.72 + i * 4.12)
        text(slide, x, Inches(2.09), Inches(1.05), Inches(0.24), a, 10, bold=True, color=CYAN)
        text(slide, x + Inches(1.13), Inches(2.04), Inches(2.48), Inches(0.35), b, 13.5, bold=True)
        if i < 2:
            text(slide, x + Inches(3.43), Inches(2.05), Inches(0.45), Inches(0.3), "→", 18, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

    rect(slide, Inches(4.3), Inches(2.92), Inches(4.73), Inches(0.91), NAVY_3, line=CYAN, radius=True)
    text(slide, Inches(4.52), Inches(3.01), Inches(4.29), Inches(0.33), "小慧总控 Agent", 17, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    text(slide, Inches(4.52), Inches(3.4), Inches(4.29), Inches(0.25), "识别任务 · 分发模块 · 返回结果", 10.5, color=PALE, align=PP_ALIGN.CENTER)
    # 总控到模块的分发线。
    for cx in (1.66, 4.18, 6.70, 9.22, 11.74):
        line(slide, Inches(6.67), Inches(3.83), Inches(cx), Inches(4.17), LINE, 1.1)

    cards = [
        ("招聘Agent", "简历解析画像\n面试纪要", "脱敏 · 回核"),
        ("入职Agent", "流程追踪\n试用期跟进", "节点 · 催办"),
        ("培训Agent", "指定主题\n培训资料", "生成 · 审核"),
        ("绩效Agent", "360反馈\n绩效报表", "聚合 · 复核"),
        ("行政Agent", "制度问答\n月报生成", "检索 · 计算"),
    ]
    for i, values in enumerate(cards):
        agent_card(slide, Inches(0.48 + i * 2.52), *values)

    rect(slide, Inches(0.48), Inches(6.05), Inches(12.36), Inches(0.67), NAVY_2, line=ORANGE, radius=True, transparency=2)
    text(slide, Inches(0.72), Inches(6.17), Inches(2.1), Inches(0.29), "后期开发计划｜薪酬模块", 11, bold=True, color=ORANGE)
    text(
        slide,
        Inches(2.78),
        Inches(6.13),
        Inches(9.7),
        Inches(0.36),
        "涉及金税四期、社保公积金、个税申报及外部系统对接，属于强合规领域，一期暂不自动化。",
        11.5,
        color=PALE,
    )
    add_notes(
        slide,
        "口播：我们把OPC“一人公司”的概念延伸到部门范围，以实现OPD“一人部门”为目标进行规划。"
        "注意这是项目概念延伸，不是政策原文。第一期采用1协调+5模块：总控只负责分发和返回；"
        "招聘、入职、培训、绩效、行政分别承担对应业务。薪酬因强合规和外部系统对接，列入后期。"
        "本页结束后进入技能演示。",
    )


def main() -> None:
    source = Presentation(str(TEMPLATE))
    prs = Presentation(str(TEMPLATE))
    clear_slides(prs)
    prs.slide_width, prs.slide_height = W, H
    cover(prs, source)
    policy(prs)
    research(prs)
    scheme(prs)
    prs.core_properties.title = "小慧 · HR数字员工 · 开头部分初版"
    prs.core_properties.subject = "云南省职业院校技能大赛人工智能赛道展示"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
