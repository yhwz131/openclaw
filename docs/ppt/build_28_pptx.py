#!/usr/bin/env python3
"""小慧赛场控页稿：26 页（上限 28）。新内容只换槽，不往后加页。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

import build_admin_pptx as admin
import build_closing_pptx as closing
import build_control_recruitment_pptx as recruit
import build_full_flow_pptx as flow
import build_onboarding_coordination_pptx as onboard
import build_opening_template_pptx as base
import build_paradigm_pptx as paradigm
import build_performance_pptx as performance
import build_quant_pptx as quant
import build_team_slide as team
import build_training_pptx as training


OUT = Path("/opt/openclaw-hr/docs/PPT准备/小慧-控页26版.pptx")
TOTAL = 26
PAGE = [1]  # 封面不走导航，下一页从 02 起


def patch_navigation() -> None:
    real = base.add_navigation
    base.COMPOSE_TOTAL = TOTAL

    def add_navigation(slide, active, page=None):
        PAGE[0] += 1
        base.TOTAL = TOTAL
        real(slide, active, PAGE[0])

    base.add_navigation = add_navigation
    closing.END_ITEMS = ["一期成果", "方法创新", "价值规划", "致谢"]


def build() -> None:
    patch_navigation()
    source = Presentation(str(base.TEMPLATE))
    prs = Presentation(str(base.TEMPLATE))
    base.clear_slides(prs)
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL

    # 前场 5：封面 / 团队 / 政策 / 调研 / 方案
    # 企业调研只替换第 4 槽，不另开页。
    base.cover(prs, source)
    team.build(prs)
    base.policy(prs)
    base.research(prs)
    base.scheme(prs)

    # 技能：架构一页带过 + 总控一页 + 招聘四页
    recruit.architecture(prs)
    recruit.router_slide(prs)
    recruit.recruitment_map(prs)
    recruit.desensitization(prs)
    recruit.number_audit(prs)
    recruit.close_loop(prs)

    # 入职 7→4：时间轴、审批门、第二门、闭环
    onboard.slide_timeline(prs)
    onboard.slide_lobster_core(prs)
    onboard.slide_second_gate(prs)
    onboard.slide_close_loop(prs)

    # 培训 4→2：自检闸 + 现场
    training.slide_quality_gate(prs)
    training.slide_live_demo(prs)

    # 绩效 4→2：缺自评停 + 现场
    performance.slide_gate_aggregate(prs)
    performance.slide_live_demo(prs)

    # 行政 4→2：有依据才答 + 权限闸
    admin.slide_policy(prs)
    admin.slide_permission_demo(prs)

    # 联调 1
    flow.build_slide(prs)

    # 收尾 4：成果含量化 / 方法创新 / 展望 / 致谢
    quant.result_with_sample(prs)
    paradigm.innovation_slide(prs)
    closing.value_future_slide(prs)
    closing.thanks_slide(prs, source)

    base.TOTAL = TOTAL
    n = len(prs.slides)
    if n != TOTAL:
        raise SystemExit(f"expected {TOTAL} slides, got {n}; nav counter={PAGE[0]}")

    prs.core_properties.title = "小慧 · 控页26版"
    prs.core_properties.subject = "赛场控页稿：新内容换槽，不追加"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={n} last_nav={PAGE[0]}")


if __name__ == "__main__":
    build()
