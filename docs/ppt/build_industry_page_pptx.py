#!/usr/bin/env python3
"""政策页之后：三个行业问题。做法放到亮点页，本页只留问题。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

import build_opening_template_pptx as base


OUT = Path("/opt/openclaw-hr/docs/PPT准备/小慧-行业约束页.pptx")
RED_DARK = RGBColor(0x56, 0x17, 0x26)
RED_LIGHT = RGBColor(0xFF, 0x1A, 0x1A)
GLOW_FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"


def glow_text_png(text: str, *, pt: float = 26) -> BytesIO:
    """红灯外发光：透明底，只有红色光晕和字，不再垫深色底。"""
    scale = 6
    font = ImageFont.truetype(GLOW_FONT, int(pt * scale), index=2)
    dummy = ImageDraw.Draw(Image.new("RGBA", (1, 1)))
    x0, y0, x1, y1 = dummy.textbbox((0, 0), text, font=font)
    pad = int(28 * scale)
    width, height = (x1 - x0) + pad * 2, (y1 - y0) + pad * 2
    pos = (pad - x0, pad - y0)

    def layer(fill, blur=0, stack=1):
        im = Image.new("RGBA", (width, height), (0, 0, 0, 0))
        ImageDraw.Draw(im).text(pos, text, font=font, fill=fill)
        if blur:
            im = im.filter(ImageFilter.GaussianBlur(blur))
        acc = im
        for _ in range(stack - 1):
            acc = Image.alpha_composite(acc, im)
        return acc

    glow = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    for item in (
        layer((255, 0, 0, 255), int(14 * scale), 3),
        layer((255, 16, 0, 255), int(6 * scale), 2),
        layer((255, 48, 24, 255), int(2.2 * scale), 1),
        layer((255, 32, 16, 255), 0, 1),
        layer((255, 240, 230, 110), int(0.6 * scale), 1),
    ):
        glow = Image.alpha_composite(glow, item)

    buf = BytesIO()
    glow.save(buf, format="PNG")
    buf.seek(0)
    return buf


def add_glow_label(slide, left, top, height, text: str, *, pt: float) -> None:
    img = glow_text_png(text, pt=pt)
    with Image.open(img) as im:
        ratio = im.width / float(im.height)
    img.seek(0)
    slide.shapes.add_picture(img, left, top, height=height, width=int(height * ratio))


def build(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    base.TOTAL = 27
    base.add_navigation(slide, 0, 4)
    base.add_title(
        slide,
        "01  PROJECT BACKGROUND",
        "智能体进职场，卡在三道坎上",
        "说得不准，做得越界，查得不清",
    )

    rows = [
        (
            "01",
            "说得不准",
            "一本正经地编数字、编事实。\n信通院测评：幻觉输出率超过10%。",
        ),
        (
            "02",
            "做得越界",
            "没让它做的事，它也可能动手。\n《实施意见》：执行不得超出用户授权。",
        ),
        (
            "03",
            "查得不清",
            "58%的受访者说不清它能碰哪些数据。\n《实施意见》：重要行为要可验证、可追溯。",
        ),
    ]
    for i, (num, problem, explain) in enumerate(rows):
        y = Inches(2.16 + i * 1.42)
        base.rect(slide, Inches(0.48), y, Inches(12.36), Inches(1.28), RED_DARK, line=base.RED, radius=True)
        base.rect(slide, Inches(0.48), y, Inches(0.08), Inches(1.28), base.RED)
        add_glow_label(slide, Inches(0.62), y + Inches(0.24), Inches(0.96), problem, pt=28)
        base.text(slide, Inches(0.82), y + Inches(0.10), Inches(0.6), Inches(0.26), num, 13, bold=True, color=RED_LIGHT)
        base.text(
            slide,
            Inches(4.7),
            y + Inches(0.18),
            Inches(7.8),
            Inches(0.92),
            explain,
            18,
            color=base.PALE,
            line_spacing=1.18,
            anchor=MSO_ANCHOR.MIDDLE,
        )

    base.text(
        slide,
        Inches(0.5),
        Inches(6.5),
        Inches(12.3),
        Inches(0.42),
        "信通院《人工智能安全治理研究报告（2025）》　｜　《实施意见》第6、7条　｜　21世纪经济报道 2025年7月2日，受访43人",
        12,
        color=base.MUTED,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约35秒。本页只讲问题，不讲我们怎么做。"
        "先念标题：智能体进职场，卡在三道坎上——说得不准，做得越界，查得不清。"
        "第一，说得不准：一本正经地编数字、编事实，信通院测评幻觉输出率超过10%。"
        "第二，做得越界：没让它做的事，它也可能动手，《实施意见》写明执行不得超出用户授权。"
        "第三，查得不清：21世纪经济报道访问43人，58%说不清它能碰哪些数据；《实施意见》要求重要行为可验证、可追溯。"
        "下一页再讲：让模型去写，让脚本去核，让人去定。",
    )


def main() -> None:
    prs = Presentation(str(base.TEMPLATE))
    base.clear_slides(prs)
    prs.slide_width, prs.slide_height = base.W, base.H
    build(prs)
    prs.core_properties.title = "智能体进职场，卡在三道坎上"
    prs.core_properties.subject = "三个行业问题"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
