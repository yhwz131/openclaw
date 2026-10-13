"""从 v2.2 生成 v2.3：去重复、压文字、版式差异化。v2.2 不改。"""
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v2.2.pptx"
DST = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v2.3.pptx"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
E = 914400

NAVY, CARD, MID, DEEP = "02163B", "032B5C", "04457D", "00112C"
CYAN, GREEN, ORANGE, PINK, PURPLE = "00D8FF", "26D098", "FFA421", "FF6B8A", "8C7CFF"
WHITE, LIGHT, GREY, ALERT = "FFFFFF", "E9F8FF", "9EC2D8", "FF9DB0"
C = RGBColor.from_string

p = Presentation(SRC)
S = {i + 1: s for i, s in enumerate(p.slides)}   # 按 v2.2 页码取


def by_id(s, sid):
    return next(x for x in s.shapes if x.shape_id == sid)


def remove(s, ids):
    for sh in list(s.shapes):
        if sh.shape_id in ids:
            if sh.shape_type == 13:
                s.part.drop_rel(sh._element.find(".//" + A + "blip").get(R + "embed"))
            sh._element.getparent().remove(sh._element)


def set_text(sh, text):
    tf = sh.text_frame
    paras = tf.paragraphs
    runs = paras[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""
    for para in paras[1:]:
        para._p.getparent().remove(para._p)


def rect(s, l, t, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = C(fill)
    if line:
        r.line.color.rgb = C(line)
        r.line.width = Pt(1)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def tb(s, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, font="微软雅黑"):
    b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, (text, size, color, bold) in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run()
        r.text = text
        f = r.font
        f.size, f.color.rgb, f.bold, f.name = Pt(size), C(color), bold, font
        rpr = r._r.get_or_add_rPr()
        rpr.append(rpr.makeelement(A + "ea", {"typeface": font}))
    return b


def chip(s, l, t, w, h, text, fill, color=NAVY, size=12):
    rect(s, l, t, w, h, fill)
    tb(s, l, t, w, h, [(text, size, color, True)], align=PP_ALIGN.CENTER)


# ① 第 8 页：表格保留，每格压成关键词
P8 = {
    42: "意图分派 · LoRA 对照", 44: "只给派发工具", 46: "说不清就问人",
    50: "解析 · 评分 · 检索 · 转写", 52: "删敏感字段 · 数字回核", 54: "约面 · 录用",
    58: "读进度 · 报卡点", 60: "Lobster 审批门 · 到点催办", 62: "合同 · 账号审批",
    66: "判断该答 · 原文作答", 68: "未登记拦截 · 身份闸门", 70: "红线 · 争议 · 审批",
    74: "ADDIE 七步 · 八类物料", 76: "读写依赖 · 缺输入就停", 78: "讲师审稿",
    82: "读评语 · 判共识分歧", 84: "校验 · 聚合 · 缺自评停", 86: "定级 · 面谈",
}
for sid, text in P8.items():
    sh = by_id(S[8], sid)
    set_text(sh, text)
    for r in sh.text_frame.paragraphs[0].runs:
        r.font.size = Pt(14)

# ② 第 10 页：每条专线加核心技术标签，并入“一期 18 项”
s10 = S[10]
for x, text, color in [(2.45, "数字回核", CYAN), (4.54, "审批门真停", ORANGE), (6.63, "ADDIE 编排", PURPLE),
                       (8.72, "分组极差", GREEN), (10.81, "身份闸门", PINK)]:
    chip(s10, x + 0.23, 3.8, 1.5, 0.27, text, color, size=10.5)
tb(s10, 1.12, 5.55, 1.25, 0.6, [("一期覆盖", 9.5, GREY, False), ("18 项", 17, CYAN, True)])
nt = s10.notes_slide.notes_text_frame
nt.text = nt.text.replace(
    "【指图】系统是 1 个总控加 5 条业务专线。",
    "调研时我们和企业人事岗位逐项核对了 30 多项事项，一期做 18 项：重复、可回溯的交给 Agent，录用、审批、定级这类决定留给人。\n\n【指图】系统是 1 个总控加 5 条业务专线。【指箭头上的标签】每条线各有一项核心技术。",
)

# ③ 第 12 页：去掉派发结构图，左侧只留工具清单，放大成能用 / 禁用两块
s12 = S[12]
remove(s12, set(range(44, 72)))
rect(s12, 0.48, 2.4, 5.9, 3.7, CARD)
tb(s12, 0.68, 2.48, 5.5, 0.32, [("总控的工具清单（openclaw.yaml › router.tools）", 11, CYAN, False)])
for y, head, color, items in [(2.92, "✓ 能用", GREEN, "派发任务 · 等待结果 · 读文件 · 看图"),
                              (4.5, "✗ 禁用", "E94F4F", "执行命令 · 写改文件 · 浏览器 · 定时任务")]:
    rect(s12, 0.68, y, 5.5, 1.4, DEEP, line=color)
    tb(s12, 0.86, y + 0.1, 5.2, 0.45, [(head, 18, color, True)])
    tb(s12, 0.86, y + 0.62, 5.2, 0.6, [(items, 16, WHITE, True)])

# ④ 四个模块封面改成转场页
COVERS = {
    14: ("03  RECRUITMENT AGENT", "招聘", CYAN, "第一道坎 · 说得不准",
         "简历变画像，帮 HR 快速初筛",
         [("模型", "画像与评语"), ("脚本", "删敏感字段、核对数字"), ("人", "约面与录用")]),
    19: ("03  ONBOARDING AGENT", "入职", ORANGE, "第二道坎 · 做得越界",
         "跨部门时间轴：部门办事，Agent 跟进度、报卡点",
         [("Lobster", "工作流编排与审批门"), ("脚本", "抽取 HCM 业务数据"), ("人", "合同与账号审批")]),
    24: ("03  ADMIN AGENT", "行政", PINK, "第三道坎 · 查得不清",
         "员工问制度、HR 查数据：制度有出处，数据有权限",
         [("模型", "按原文作答"), ("脚本", "按身份放行"), ("人", "个案与争议")]),
    29: ("03  PERFORMANCE AGENT", "绩效", GREEN, "把决定留给人",
         "把 360 反馈变成可复核的面谈材料，不给评级",
         [("模型", "读评语、写面谈提纲"), ("脚本", "依赖检查、分数计算"), ("人", "最终评级")]),
}
LAYER = [CYAN, GREEN, ORANGE]
for n, (label, name, color, hurdle, pos, triad) in COVERS.items():
    s = S[n]
    remove(s, set(range(40, 58)))
    rect(s, 0, 1.18, 4.9, 6.32, color)
    tb(s, 0.6, 1.9, 4.0, 0.36, [(label, 13, NAVY, True)])
    tb(s, 0.55, 2.3, 4.2, 1.4, [(name, 72, NAVY, True)])
    tb(s, 0.62, 3.7, 4.0, 0.6, [("Agent", 30, NAVY, True)])
    rect(s, 0.68, 4.45, 0.9, 0.07, NAVY)
    tb(s, 0.6, 4.7, 4.1, 0.5, [(hurdle, 18, NAVY, True)])
    tb(s, 5.5, 2.2, 7.4, 1.3, [(pos, 26, WHITE, True)], anchor=MSO_ANCHOR.TOP)
    for k, (who, what) in enumerate(triad):
        y = 3.9 + k * 0.72
        chip(s, 5.55, y, 1.25, 0.46, who, LAYER[k], size=14)
        tb(s, 6.95, y, 5.8, 0.46, [(what, 18, LIGHT, False)])

# ⑤ 三张风险页：法、规各压成一行，空出位置放两个风险实例
RISK = {
    15: (set(range(71, 91)),
         "《就业促进法》不得就业歧视；《个人信息保护法》自动化决策须公平 → 评分不能看身份",
         "交给 HR 的材料必须真实：每个数字、每段经历，都要在本人简历里查得到",
         ["年龄、性别、院校进了评分，就无法排除它们影响打分", "评分写“58 万笔/秒”，原简历是“12.8 万 TPS”"]),
    20: (set(range(80, 100)),
         "《劳动合同法》《实施意见》：关键动作要有人授权，执行不得超出授权",
         "入职 SOP：合同由 HR 主管审，账号权限由 IT 核；进度以 HCM 记录为准",
         ["HR 主管还没看合同，Agent 就推进签约", "IT 还没核权限，Agent 就开通系统访问"]),
    25: (set(range(80, 100)),
         "《个人信息保护法》第六、五十一条：最小必要，合理确定操作权限",
         "制度以公司原文为准；争议、法律定性、审批，判断权在人",
         ["员工查同事的工资，Agent 照查照答", "凭常识答一句“按国家规定”，出处说不清"]),
}
for n, (ids, law, rule, risks) in RISK.items():
    s = S[n]
    remove(s, ids)
    for y, tag, color, text in [(3.3, "法", "F5C542", law), (3.86, "规", CYAN, rule)]:
        rect(s, 0.48, y, 12.37, 0.48, CARD)
        chip(s, 0.6, y + 0.07, 0.5, 0.34, tag, color, size=14)
        tb(s, 1.22, y, 11.5, 0.48, [(text, 13, LIGHT, False)])
    for k, text in enumerate(risks):
        x = 0.48 + k * 6.28
        rect(s, x, 4.5, 6.08, 0.9, "3A1530", line=PINK)
        tb(s, x + 0.15, 4.52, 1.2, 0.3, [("风险实例", 10.5, PINK, True)])
        tb(s, x + 0.15, 4.8, 5.8, 0.55, [(text, 15.5, WHITE, True)])

# ⑥ 页内重复：底部“删什么 / 谁来删 / 怎么证明”三格与正文重复，删掉
TRIOS = {16: range(74, 86), 17: range(87, 99), 21: range(87, 99), 26: range(90, 102), 27: range(81, 93)}
for n, ids in TRIOS.items():
    remove(S[n], set(ids))

# ⑦ 模块主色：标题下短线与英文小标随模块换色
MODULE_COLOR = {}
for n in (15, 16, 17, 18):
    MODULE_COLOR[n] = CYAN
for n in (20, 21, 22, 23):
    MODULE_COLOR[n] = ORANGE
for n in (25, 26, 27, 28):
    MODULE_COLOR[n] = PINK
for n in (30, 31, 32):
    MODULE_COLOR[n] = GREEN
for n, color in MODULE_COLOR.items():
    s = S[n]
    for sh in s.shapes:
        top, left = sh.top / E, sh.left / E
        if sh.has_text_frame and abs(left - 0.48) < 0.05 and 1.2 < top < 1.4 and re.match(r"^\d\d\s", sh.text_frame.text):
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    r.font.color.rgb = C(color)
        if not sh.has_text_frame and abs(left - 0.48) < 0.05 and 2.05 < top < 2.15 and sh.width / E < 0.8 and sh.height / E < 0.08:
            sh.fill.solid()
            sh.fill.fore_color.rgb = C(color)
    rect(s, 0, 1.18, 0.12, 6.32, color)

for n in (15, 20, 25):
    for para in S[n].notes_slide.notes_text_frame.paragraphs:
        for r in para.runs:
            r.text = r.text.replace("【指左栏 · 法】", "【指“法”一行】").replace("【指右栏 · 规】", "【指“规”一行】")

# ⑧ 删除第 9 页
lst = p.slides._sldIdLst
el9 = list(lst)[8]
p.part.drop_rel(el9.get(R + "id"))
lst.remove(el9)

# 页码：只改右上角页码框
total = len(p.slides)
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.top is not None and sh.top / E < 0.6 and sh.left / E > 11 \
                and re.fullmatch(r"\d{2} / \d{2}", sh.text_frame.text.strip()):
            set_text(sh, f"{i:02d} / {total:02d}")

p.save(DST)
print("saved", DST, total)
