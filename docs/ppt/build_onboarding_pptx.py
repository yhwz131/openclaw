#!/usr/bin/env python3
"""按“小慧”统一科技风生成入职模块四页技能演示 PPT。"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import build_opening_template_pptx as base


ROOT = Path("/opt/openclaw-hr/docs")
OUT = ROOT / "PPT准备/入职模块技能演示-4页完善版.pptx"
PREVIEW = ROOT / "PPT准备/入职模块技能演示-4页完善版-预览.png"
ASSETS = ROOT / "ppt/assets/onboarding"
DASHBOARD = ASSETS / "onboarding-dashboard.png"
DASHBOARD_CROP = ASSETS / "onboarding-dashboard-crop.png"

TOTAL = 4
GREEN = RGBColor(0x26, 0xD0, 0x98)
GREEN_DARK = RGBColor(0x08, 0x50, 0x57)
PURPLE = RGBColor(0x8C, 0x7C, 0xFF)
RED = RGBColor(0xF1, 0x5A, 0x66)
RED_DARK = RGBColor(0x56, 0x17, 0x26)
CODE_BG = RGBColor(0x00, 0x11, 0x2C)


def connector(slide, x1, y1, x2, y2, color=base.CYAN, width=1.1, arrow=True):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def secondary_nav(slide) -> None:
    labels = ["总控", "招聘", "入职", "培训", "绩效", "行政"]
    base.rect(slide, Inches(0), Inches(0.72), base.W, Inches(0.45), base.NAVY, line=base.LINE)
    base.text(
        slide,
        Inches(0.38),
        Inches(0.79),
        Inches(0.9),
        Inches(0.27),
        "技能地图",
        10,
        bold=True,
        color=base.MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    for i, label in enumerate(labels):
        x = Inches(1.33 + i * 1.77)
        active = i == 2
        base.rect(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            base.CYAN if active else base.NAVY_2,
            line=base.CYAN if active else base.LINE,
            radius=True,
        )
        base.text(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            label,
            10.5,
            bold=active,
            color=base.NAVY if active else base.MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def header(slide, page: int, eyebrow: str, title: str, subtitle: str) -> None:
    base.add_navigation(slide, 2, page)
    secondary_nav(slide)
    base.text(slide, Inches(0.48), Inches(1.28), Inches(3.5), Inches(0.25), eyebrow, 10.5, bold=True, color=base.CYAN)
    base.text(slide, Inches(0.46), Inches(1.53), Inches(12.2), Inches(0.49), title, 23, bold=True, font=base.TITLE_FONT)
    base.rect(slide, Inches(0.47), Inches(2.08), Inches(0.75), Inches(0.05), base.CYAN)
    base.text(slide, Inches(1.4), Inches(1.99), Inches(11.1), Inches(0.29), subtitle, 11.5, color=base.MUTED)


def card(slide, x, y, w, h, title: str, body: str, *, accent=base.CYAN, fill=base.NAVY_2, body_size=11.5):
    base.rect(slide, x, y, w, h, fill, line=base.LINE, radius=True, transparency=3)
    base.rect(slide, x, y, Inches(0.08), h, accent)
    base.text(slide, x + Inches(0.2), y + Inches(0.14), w - Inches(0.4), Inches(0.3), title, 13, bold=True, color=accent)
    base.text(slide, x + Inches(0.2), y + Inches(0.55), w - Inches(0.4), h - Inches(0.68), body, body_size, color=base.PALE, line_spacing=0.92)


def source_card(slide, x, title: str, source: str, items: str, accent) -> None:
    card(slide, x, Inches(2.43), Inches(3.82), Inches(2.26), title, items, accent=accent, body_size=12)
    base.rect(slide, x + Inches(0.2), Inches(4.29), Inches(3.42), Inches(0.25), base.NAVY_3, line=accent, radius=True)
    base.text(
        slide,
        x + Inches(0.2),
        Inches(4.29),
        Inches(3.42),
        Inches(0.25),
        f"真相源｜{source}",
        9.5,
        bold=True,
        color=accent,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def overview(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        1,
        "03  ONBOARDING AGENT · MECHANISM",
        "入职小慧：三类状态源，两道人工审批门",
        "不是替部门办事，而是让真实状态按正确顺序流动",
    )
    source_card(slide, Inches(0.48), "入职中｜部门准备", "HCM · 只读 GET", "合同 / 账号 / 设备\nBuddy / 培训\n五节点分别追踪", base.CYAN)
    source_card(slide, Inches(4.75), "入职中｜审批流程", "Lobster · workflow.json", "合同审批门\n账号权限审批门\nresumeToken 恢复", base.ORANGE)
    source_card(slide, Inches(9.02), "入职后｜持续跟进", "reviews.json", "第 7 / 30 / 60 / 90 天\n到期只提醒\nHR 确认后完成", GREEN)

    base.rect(slide, Inches(0.48), Inches(5.0), Inches(12.36), Inches(1.2), base.NAVY_3, line=base.CYAN, radius=True)
    parts = [("HCM 五节点全部完成", base.CYAN), ("＋", base.PALE), ("两道审批门均批准", base.ORANGE), ("→", base.PALE), ("开通 7/30/60/90", GREEN)]
    positions = [(0.76, 2.6), (3.56, 0.5), (4.22, 2.7), (7.28, 0.5), (8.0, 3.9)]
    for (label, color), (x, w) in zip(parts, positions):
        base.text(slide, Inches(x), Inches(5.22), Inches(w), Inches(0.42), label, 16 if label not in ("＋", "→") else 20, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    base.text(
        slide,
        Inches(0.76),
        Inches(5.75),
        Inches(11.8),
        Inches(0.25),
        "Agent 读取、推进与提醒｜脚本校验状态｜HR 与责任人审批并确认",
        11,
        bold=True,
        color=base.MUTED,
        align=PP_ALIGN.CENTER,
    )
    base.rect(slide, Inches(0.48), Inches(6.38), Inches(12.36), Inches(0.39), base.NAVY_2, line=base.LINE, radius=True)
    base.text(slide, Inches(0.68), Inches(6.38), Inches(11.96), Inches(0.39), "边界：HCM 只读｜小慧不替部门填写｜小慧不替负责人批准", 11.5, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    base.add_notes(
        slide,
        "约50秒。招聘模块完成后，是否录用仍由HR决定。确定录用后才进入入职模块。"
        "入职有三类状态源：HCM保存员工与五个准备节点，Lobster保存两道审批门，本地跟进表保存7、30、60、90天状态。"
        "小慧只读取、推进和提醒，不写HCM、不代替审批。五节点全完成且两门均批准，才允许开通后续跟进。",
    )


def start_and_halt(prs: Presentation, page: int = 2) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        page,
        "03  LIVE DEMO · START & HALT",
        "启动不是生成结果：校验通过后，流程主动暂停",
        "现场输入：请为员工2启动入职流程，岗位是前端工程师，薪资30K",
    )

    checks = [
        ("01", "HCM 已建档？", "否 → 提示先建档", base.CYAN),
        ("02", "姓名 / 拼音匹配？", "否 → 报错，不写状态", base.CYAN),
        ("03", "合同已开始？", "否 → waiting_departments", base.CYAN),
        ("OK", "只读五节点快照", "进入 contract_gate", GREEN),
    ]
    for i, (num, title, sub, accent) in enumerate(checks):
        x = Inches(0.48 + i * 3.08)
        base.rect(slide, x, Inches(2.42), Inches(2.78), Inches(1.16), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x + Inches(0.16), Inches(2.57), Inches(0.43), Inches(0.25), num, 10.5, bold=True, color=accent)
        base.text(slide, x + Inches(0.58), Inches(2.51), Inches(1.98), Inches(0.32), title, 13, bold=True)
        base.text(slide, x + Inches(0.18), Inches(3.02), Inches(2.42), Inches(0.25), sub, 9.8, color=base.MUTED)
        if i < 3:
            connector(slide, x + Inches(2.78), Inches(3.0), x + Inches(3.02), Inches(3.0), accent, 1.0)

    base.rect(slide, Inches(0.48), Inches(3.88), Inches(5.62), Inches(2.4), CODE_BG, line=base.LINE, radius=True)
    base.text(slide, Inches(0.72), Inches(4.05), Inches(5.14), Inches(0.3), "现场必须看到｜当前审批暂停结果", 12, bold=True, color=base.CYAN)
    terminal = (
        "status        needs_approval\n"
        "current_step  contract_gate\n"
        "approver      HR 主管\n"
        "prompt        审批正文原样展示\n"
        "resumeToken   本次专属恢复钥匙"
    )
    base.text(slide, Inches(0.72), Inches(4.52), Inches(5.05), Inches(1.42), terminal, 12.3, color=base.PALE, font="Courier New", line_spacing=0.9)

    base.rect(slide, Inches(6.38), Inches(3.88), Inches(6.46), Inches(2.4), base.NAVY_2, line=GREEN, radius=True)
    base.text(slide, Inches(6.62), Inches(4.05), Inches(5.98), Inches(0.3), "真实运行记录节选｜workflow.json", 12, bold=True, color=GREEN)
    history = (
        "15:10:36  contract_gate  halted\n"
        "15:12:04  contract_gate  approved\n"
        "15:12:04  account_gate   halted\n"
        "15:16:39  account_gate   approved\n"
        "15:16:39  complete       completed"
    )
    base.text(slide, Inches(6.62), Inches(4.52), Inches(5.92), Inches(1.42), history, 12.3, color=base.PALE, font="Courier New", line_spacing=0.9)

    base.rect(slide, Inches(0.48), Inches(6.48), Inches(12.36), Inches(0.32), RED_DARK, line=RED, radius=True)
    base.text(slide, Inches(0.68), Inches(6.48), Inches(11.96), Inches(0.32), "HCM 不可用 / 查不到员工 / 合同未开始 → 明确停止，不拿旧数据冒充成功", 10.5, bold=True, color=RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    base.add_notes(
        slide,
        "约2分钟。现场输入员工2的入职任务。先观察HCM建档、姓名拼音、合同节点三道校验。"
        "校验通过后只读五节点并进入contract_gate。重点让评委看见needs_approval、审批正文和resumeToken。"
        "右侧是仓库当前workflow.json的真实历史记录，证明两道门均发生过暂停和人工恢复。"
        "正式演示前必须重置为可演示起点，不要临场使用不确定的--force。",
    )


def gates_and_completion(prs: Presentation, page: int = 3) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        page,
        "03  STATE MACHINE · HUMAN IN THE LOOP",
        "两道人工门＋双重完成条件，拦住“假成功”",
        "AI 可以推进流程，但合同责任和系统访问权限必须由人决定",
    )

    card(
        slide,
        Inches(0.48),
        Inches(2.42),
        Inches(3.62),
        Inches(1.75),
        "第一道｜合同审批门",
        "审批人：HR 主管\n批准 → 使用当次 token 恢复\n拒绝 → cancelled，流程终止",
        accent=base.ORANGE,
        body_size=11.7,
    )
    connector(slide, Inches(4.12), Inches(3.29), Inches(4.6), Inches(3.29), base.ORANGE, 1.5)
    card(
        slide,
        Inches(4.64),
        Inches(2.42),
        Inches(3.62),
        Inches(1.75),
        "第二道｜账号权限门",
        "审批人：IT 安全负责人\n批准 → 审批流程 complete\n拒绝 → cancelled，流程终止",
        accent=base.ORANGE,
        body_size=11.7,
    )
    connector(slide, Inches(8.28), Inches(3.29), Inches(8.76), Inches(3.29), GREEN, 1.5)
    card(
        slide,
        Inches(8.8),
        Inches(2.42),
        Inches(4.04),
        Inches(1.75),
        "每道门的硬规则",
        "审批正文不可改写\n新门生成新 token，旧 token 不复用\napprovalId 不能代替 resumeToken",
        accent=GREEN,
        body_size=11.2,
    )

    base.text(slide, Inches(0.48), Inches(4.42), Inches(12.0), Inches(0.3), "两类假成功｜审批完成与业务完成必须分别判断", 12.5, bold=True, color=base.CYAN)
    card(
        slide,
        Inches(0.48),
        Inches(4.82),
        Inches(5.92),
        Inches(1.46),
        "假成功 A｜两道门通过＝全部准备完成？",
        "不成立。账号、设备、Buddy、培训仍可能待办。\n正确：HCM 五节点全部 completed 后才能开通跟进。",
        accent=RED,
        fill=RED_DARK,
        body_size=11.2,
    )
    card(
        slide,
        Inches(6.68),
        Inches(4.82),
        Inches(6.16),
        Inches(1.46),
        "假成功 B｜提醒发出＝跟进已经完成？",
        "不成立。到期只产生 needs_review。\n正确：HR 确认实际沟通后，才记录 mark-done。",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=11.2,
    )
    base.rect(slide, Inches(0.48), Inches(6.48), Inches(12.36), Inches(0.32), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.68), Inches(6.48), Inches(11.96), Inches(0.32), "审批流程 complete ＋ HCM 五节点 completed → 开通 7 / 30 / 60 / 90 天跟进", 10.8, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    base.add_notes(
        slide,
        "约1分40秒。合同门由HR主管决定，账号权限门由IT安全负责人决定。"
        "每道门都生成新的resumeToken，任何一道拒绝都立即终止。"
        "两道门通过不等于部门工作完成；提醒到期也不等于跟进完成。"
        "审批流程complete且HCM五节点completed，才开通7、30、60、90天跟进，结果仍由HR确认。",
    )


def crop_dashboard() -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.open(DASHBOARD).convert("RGB")
    w, h = image.size
    image.crop((0, 0, w, int(h * 0.64))).save(DASHBOARD_CROP, quality=95)
    return DASHBOARD_CROP


def query_and_dashboard(prs: Presentation, page: int = 4) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        page,
        "03  QUERY & DASHBOARD · EVIDENCE",
        "一次查询定位卡点，一张看板管理全局",
        "三条线分别展示，所有结论都能回到真实来源",
    )

    base.rect(slide, Inches(0.48), Inches(2.42), Inches(4.15), Inches(3.78), base.NAVY_2, line=base.LINE, radius=True)
    base.text(slide, Inches(0.72), Inches(2.6), Inches(3.67), Inches(0.3), "现场查询｜员工2到哪一步？", 13.5, bold=True, color=base.CYAN)
    rows = [
        ("HCM", "五个准备节点", base.CYAN),
        ("workflow.json", "当前审批门与归档状态", base.ORANGE),
        ("reviews.json", "7/30/60/90 跟进状态", GREEN),
    ]
    for i, (source, result, accent) in enumerate(rows):
        y = Inches(3.12 + i * 0.72)
        base.rect(slide, Inches(0.72), y, Inches(3.67), Inches(0.57), base.NAVY_3, line=accent, radius=True)
        base.text(slide, Inches(0.9), y + Inches(0.08), Inches(1.18), Inches(0.23), source, 10, bold=True, color=accent)
        base.text(slide, Inches(2.05), y + Inches(0.08), Inches(2.1), Inches(0.28), result, 10.5, color=base.PALE, align=PP_ALIGN.RIGHT)
    base.text(
        slide,
        Inches(0.72),
        Inches(5.45),
        Inches(3.67),
        Inches(0.52),
        "输出：当前卡点＋下一责任部门\n只查询，不批准，不挂 Cron",
        11.5,
        bold=True,
        color=base.PALE,
        align=PP_ALIGN.CENTER,
    )

    base.rect(slide, Inches(4.9), Inches(2.42), Inches(7.94), Inches(3.78), base.NAVY_2, line=GREEN, radius=True)
    base.text(slide, Inches(5.14), Inches(2.59), Inches(7.46), Inches(0.3), "项目真实产物｜入职进度看板", 13.5, bold=True, color=GREEN)
    image_path = crop_dashboard()
    slide.shapes.add_picture(str(image_path), Inches(5.12), Inches(3.02), width=Inches(7.5), height=Inches(2.95))

    base.rect(slide, Inches(0.48), Inches(6.37), Inches(12.36), Inches(0.43), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(
        slide,
        Inches(0.68),
        Inches(6.37),
        Inches(11.96),
        Inches(0.43),
        "五节点决定准备进度｜审批门独立显示｜跟进状态单独读取｜停留过久只催办",
        11,
        bold=True,
        color=base.CYAN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约1分30秒。现场查询员工2到哪一步以及下一责任人。"
        "小慧同时读取HCM、workflow.json和reviews.json，但不会把三条线混成一个百分比。"
        "右侧是项目当前生成的真实入职看板。五节点决定准备进度，两道门独立显示，跟进状态单独读取。"
        "收口：状态来自真实系统，流程由状态机推进，关键决定由人完成，所有结论都可查询和追溯。"
        "看板地址以比赛现场设备实际地址为准。",
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL
    overview(prs)
    start_and_halt(prs)
    gates_and_completion(prs)
    query_and_dashboard(prs)
    prs.core_properties.title = "小慧 · 入职模块技能演示完善版"
    prs.core_properties.subject = "三类状态源、两道人工审批门、双重完成条件与可追溯看板"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
