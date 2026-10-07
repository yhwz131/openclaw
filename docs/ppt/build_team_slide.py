#!/usr/bin/env python3
"""生成“小慧”团队技术方向分工单页PPT。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

import build_opening_template_pptx as base


OUT = Path("/opt/openclaw-hr/docs/PPT准备/团队介绍页-方向分工版.pptx")
GREEN = base.RGBColor(0x26, 0xD0, 0x98)
PURPLE = base.RGBColor(0x8C, 0x7C, 0xFF)


def member_card(
    slide,
    x,
    number: str,
    role: str,
    modules: str,
    action: str,
    accent,
) -> None:
    y = Inches(2.72)
    base.rect(
        slide,
        x,
        y,
        Inches(2.91),
        Inches(3.18),
        base.NAVY_2,
        line=base.LINE,
        radius=True,
        transparency=3,
    )
    base.rect(slide, x, y, Inches(2.91), Inches(0.08), accent)

    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        x + Inches(1.05),
        y + Inches(0.28),
        Inches(0.8),
        Inches(0.8),
    )
    base.set_fill(badge, base.NAVY_3)
    badge.line.color.rgb = accent
    base.text(
        slide,
        x + Inches(1.05),
        y + Inches(0.28),
        Inches(0.8),
        Inches(0.8),
        number,
        18,
        bold=True,
        color=accent,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.text(
        slide,
        x + Inches(0.18),
        y + Inches(1.22),
        Inches(2.55),
        Inches(0.54),
        role,
        14.2,
        bold=True,
        color=base.PALE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.rect(
        slide,
        x + Inches(0.28),
        y + Inches(1.86),
        Inches(2.35),
        Inches(0.4),
        base.NAVY_3,
        line=accent,
        radius=True,
    )
    base.text(
        slide,
        x + Inches(0.28),
        y + Inches(1.86),
        Inches(2.35),
        Inches(0.4),
        modules,
        10.5,
        bold=True,
        color=accent,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.text(
        slide,
        x + Inches(0.28),
        y + Inches(2.43),
        Inches(2.35),
        Inches(0.52),
        action,
        10.5,
        color=base.MUTED,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        line_spacing=0.9,
    )


def build(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    if not getattr(base, "COMPOSE_TOTAL", None):
        base.TOTAL = 36
    base.add_navigation(slide, 0, 2)
    base.add_title(
        slide,
        "01  TEAM COLLABORATION",
        "团队协作：同一岗位族，四类技术方向",
        "围绕AI应用开发，分别承担系统集成、流程编排、生成评测与知识数据工作",
    )

    base.rect(
        slide,
        Inches(0.48),
        Inches(2.05),
        Inches(12.36),
        Inches(0.43),
        base.NAVY_3,
        line=base.CYAN,
        radius=True,
        transparency=5,
    )
    base.text(
        slide,
        Inches(0.72),
        Inches(2.1),
        Inches(11.9),
        Inches(0.3),
        "项目统筹  →  模块开发  →  交叉验证  →  完整联调",
        11.5,
        bold=True,
        color=base.CYAN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    cards = [
        (
            "01",
            "AI应用开发工程师",
            "系统集成方向",
            "总控 · 招聘｜Router调度 · 完整联调",
            base.CYAN,
        ),
        (
            "02",
            "AI应用开发工程师",
            "流程编排方向",
            "入职｜工作流 · 审批门 · Cron跟进",
            GREEN,
        ),
        (
            "03",
            "AI应用开发工程师",
            "生成评测方向",
            "培训 · 绩效｜ADDIE · 360反馈",
            PURPLE,
        ),
        (
            "04",
            "AI应用开发工程师",
            "知识数据方向",
            "行政｜IMA检索 · 权限 · 月报",
            base.ORANGE,
        ),
    ]
    for i, values in enumerate(cards):
        member_card(slide, Inches(0.48 + i * 3.12), *values)

    base.rect(
        slide,
        Inches(0.48),
        Inches(6.16),
        Inches(12.36),
        Inches(0.62),
        base.NAVY_2,
        line=base.LINE,
        radius=True,
    )
    base.text(
        slide,
        Inches(0.7),
        Inches(6.24),
        Inches(3.0),
        Inches(0.35),
        "同一目标",
        12.5,
        bold=True,
        color=base.CYAN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.text(
        slide,
        Inches(3.7),
        Inches(6.24),
        Inches(3.0),
        Inches(0.35),
        "分线负责",
        12.5,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.text(
        slide,
        Inches(6.7),
        Inches(6.24),
        Inches(3.0),
        Inches(0.35),
        "同步操作",
        12.5,
        bold=True,
        color=PURPLE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.text(
        slide,
        Inches(9.7),
        Inches(6.24),
        Inches(2.9),
        Inches(0.35),
        "共同验证",
        12.5,
        bold=True,
        color=base.ORANGE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    base.add_notes(
        slide,
        "本页控制在30—40秒。团队属于同一AI应用开发岗位族，按技术方向承担不同职责。"
        "建议四人依次说：我负责系统集成方向，承担总控、招聘与完整联调；"
        "我负责流程编排方向，承担入职工作流、审批门和跟进任务；"
        "我负责生成评测方向，承担培训资料生成与绩效反馈处理；"
        "我负责知识数据方向，承担政策检索、权限控制和月报复核。"
        "最后共同口号：同一目标，分线负责，同步操作，共同验证。"
        "定稿时应继续核对四人的现场操作与本页方向是否一致。",
    )


def main() -> None:
    prs = Presentation(str(base.TEMPLATE))
    base.clear_slides(prs)
    prs.slide_width, prs.slide_height = base.W, base.H
    build(prs)
    prs.core_properties.title = "小慧 · 团队介绍页方向分工版"
    prs.core_properties.subject = "AI应用开发技术方向与现场协作"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
