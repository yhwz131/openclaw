"""从 v2.3 生成 v3.0：按“背景 → 设计 → 工作流程 → 核心技能 → 成果 → 价值创新”重排。v2.2 / v2.3 不改。"""
import copy
import re
import subprocess
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

sys.path.insert(0, "/opt/openclaw-hr/docs/ppt")
subprocess.run([sys.executable, "/opt/openclaw-hr/docs/ppt/build_v23.py"], check=True)

SRC = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v2.3.pptx"
DST = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v3.0.pptx"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
E = 914400

NAVY, CARD, MID, DEEP = "02163B", "032B5C", "04457D", "00112C"
CYAN, GREEN, ORANGE, PINK, PURPLE = "00D8FF", "26D098", "FFA421", "FF6B8A", "8C7CFF"
WHITE, LIGHT, GREY, RED, WINE = "FFFFFF", "E9F8FF", "9EC2D8", "E94F4F", "3A1530"
from pptx.dml.color import RGBColor
C = RGBColor.from_string

p = Presentation(SRC)
V = {i + 1: s for i, s in enumerate(p.slides)}   # 按 v2.3 页码取


def by_id(s, sid):
    return next(x for x in s.shapes if x.shape_id == sid)


def remove(s, ids):
    for sh in list(s.shapes):
        if sh.shape_id in ids:
            if sh.shape_type == 13:
                s.part.drop_rel(sh._element.find(".//" + A + "blip").get(R + "embed"))
            sh._element.getparent().remove(sh._element)


def set_text(sh, text):
    paras = sh.text_frame.paragraphs
    runs = paras[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""
    for para in paras[1:]:
        para._p.getparent().remove(para._p)


def rect(s, l, t, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = C(fill)
    if line:
        r.line.color.rgb = C(line)
        r.line.width = Pt(1.25)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def tb(s, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, font="微软雅黑", spacing=None):
    b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, (text, size, color, bold) in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.line_spacing = spacing
        r = para.add_run()
        r.text = text
        f = r.font
        f.size, f.color.rgb, f.bold, f.name = Pt(size), C(color), bold, font
        rpr = r._r.get_or_add_rPr()
        rpr.append(rpr.makeelement(A + "ea", {"typeface": font}))
    return b


def chip(s, l, t, w, h, text, fill, color=NAVY, size=12):
    rect(s, l, t, w, h, fill)
    tb(s, l, t, w, h, [(text, size, color, True)], align=PP_ALIGN.CENTER)


def find(s, test):
    return next((sh for sh in s.shapes if test(sh)), None)


def header(s, label=None, title=None, sub=None):
    """按位置找英文小标、标题、副标题（两种版头高度都适用）。"""
    for sh in s.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        l, t, w = sh.left / E, sh.top / E, sh.width / E
        if label and abs(l - 0.48) < 0.05 and 0.85 < t < 1.4 and abs(w - 6.0) < 0.1:
            set_text(sh, label)
        elif title and abs(l - 0.46) < 0.05 and 1.0 < t < 1.6 and abs(w - 12.2) < 0.1:
            set_text(sh, title)
        elif sub and abs(l - 1.4) < 0.05 and 1.6 < t < 2.1 and abs(w - 11.4) < 0.1:
            set_text(sh, sub)


def clone(src, max_id):
    """新建一页，复制源页的版头（shape_id ≤ max_id，不含图片）。"""
    s = p.slides.add_slide(src.slide_layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    tree = s.shapes._spTree
    for sh in src.shapes:
        if sh.shape_id <= max_id and sh.shape_type != 13 and sh.shape_id not in (10, 12) and sh.top / E < 2.3:
            tree.append(copy.deepcopy(sh._element))
    return s


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# ---------- 项目设计：1+5 并进总控配置页（v2.3 第 11 页） ----------
s = V[11]
remove(s, set(range(26, 40)) | set(range(89, 97)))
for sh in s.shapes:
    sid = sh.shape_id
    if sid in (40, 41):
        sh.top -= Inches(0.37)
    elif sid == 42:
        sh.top = Inches(1.79)
    elif sid == 43:
        sh.top = Inches(1.68)
    elif 72 <= sid <= 88:
        sh.top -= Inches(0.3)
header(s, label="02  TECHNICAL ARCHITECTURE · WHY 1 + 5",
       title="为什么拆成 1 + 5：每条线的边界写在配置里",
       sub="权限、规则、故障必须隔开；配置里没给的工具，模型想用也用不了")
set_text(by_id(s, 88), "一个模型管得越多，越难证明它没越界；拆开之后，每条线的边界都能单独验证")
rect(s, 0.48, 2.1, 5.9, 3.54, CARD)
tb(s, 0.68, 2.18, 5.5, 0.34, [("一个全能 Agent  vs  1 个总控 + 5 条专线", 12, CYAN, True)])
for k, (tag, one, split) in enumerate([
        ("权限", "权限是所有模块之和", "每条线只拿自己的工具"),
        ("规则", "规则挤在一个上下文，越聊越忘", "每条线只带自己那份规则"),
        ("故障", "一处卡住，全部停", "一条线出错，只停这一条")]):
    y = 2.62 + k * 0.98
    rect(s, 0.68, y, 5.5, 0.86, DEEP)
    chip(s, 0.8, y + 0.2, 0.75, 0.46, tag, CYAN, size=13)
    tb(s, 1.7, y + 0.05, 4.4, 0.38, [("全能：" + one, 11.5, GREY, False)])
    tb(s, 1.7, y + 0.42, 4.4, 0.4, [("拆开：" + split, 14, WHITE, True)])
notes(s, """【男1】为什么拆成 1 + 5（约 60 秒）

如果做成一个全能 Agent，【指左侧三行】它的权限是所有模块之和，任何一句话都可能碰到不该碰的数据；脱敏、审批门、红线全挤在一个上下文，对话一长，前面的规则就被挤掉；一处卡住，所有业务一起停。拆开以后，每条线只拿自己的工具、只带自己的规则，出错只停这一条。

【指配置节选】边界靠什么守住？不靠提示词，靠配置。以总控为例：能用的只有派发、等待、读文件、看图；执行命令、写文件、浏览器、定时任务全部禁用。派发名单只有总控有，五条专线都是空的，所以专线之间不能互相派活。最多 8 个子任务，1800 秒超时即停。

【指底栏】一个模型管得越多，越难证明它没越界；拆开之后，每条线的边界都能单独验证。

接下来看一条数据在系统里怎么走。""")

nt = V[9].notes_slide.notes_text_frame
nt.text = nt.text.replace("【→ 接 02 总控】", "【→ 接 1 + 5】")

# ---------- 工作流程：开发流程（原 OPD 页） ----------
s = V[35]
header(s, label="03  DEVELOPMENT FLOW", title="开发流程：从部门 SOP 到可运行系统",
       sub="先盘能力、再拆业务、再定架构、按阶段开发验收；方法可迁移到其他部门")
notes(s, """【男1】开发流程（约 40 秒）

小慧不是从工具反推出来的功能集合，是按部门 SOP 一步步做出来的。【指四步】第一步能力盘点，列清手里的能力：生成、检索、工作流、脚本校验、系统接口，不绑定平台；第二步业务分析，拆开事项、对上技术、标出人做到哪一步；第三步系统架构，决定怎么拆 Agent、哪些交给脚本；第四步按阶段开发，每个阶段设检查点，通过了才进下一阶段。

【指右侧】HR 是第一个实例；换成财务、采购、客服，同样四步就能复用。""")

# ---------- 工作流程：数据流程（新） ----------
s_data = clone(V[8], 29)
header(s_data, label="03  DATA FLOW", title="数据流程：一条数据从进来到留痕，每一步谁来做",
       sub="颜色就是分工：模型去写，脚本去核，人去定")
STEPS = [
    ("① 采集", "入口", MID, ["眼镜语音", "WebChat · CLI", "飞书 · 邮箱", "HCM 接口", "360 表格"]),
    ("② 理解分派", "模型", CYAN, ["总控识别意图", "点名专线派发", "说不清就问人", "LoRA 小模型对照"]),
    ("③ 预处理", "脚本", GREEN, ["文档转 Markdown", "删 8 类敏感字段", "读数据前认身份", "缺输入就停"]),
    ("④ 生成分析", "模型", CYAN, ["八维评分", "向量检索", "面试转写分人", "评语语义 · 物料"]),
    ("⑤ 核验", "脚本", GREEN, ["数字回原文", "结构校验", "分组算极差", "看退出码"]),
    ("⑥ 决定", "人", ORANGE, ["审批门停下等人", "录用、合同、账号", "定级、争议", "人批才往下走"]),
    ("⑦ 输出留痕", "系统", MID, ["Canvas 看板", "飞书回执", "记录落盘", "事后可回溯"]),
]
for k, (step, who, color, items) in enumerate(STEPS):
    x = 0.48 + k * 1.78
    dark = color in (CYAN, GREEN, ORANGE)
    rect(s_data, x, 2.2, 1.64, 3.2, CARD, line=color if dark else None)
    rect(s_data, x, 2.2, 1.64, 0.5, color)
    tb(s_data, x, 2.2, 1.64, 0.5, [(step, 14, NAVY if dark else WHITE, True)], align=PP_ALIGN.CENTER)
    tb(s_data, x, 2.78, 1.64, 0.3, [(who, 12, color if dark else GREY, True)], align=PP_ALIGN.CENTER)
    tb(s_data, x + 0.02, 3.18, 1.6, 2.15, [(it, 12.5, LIGHT, False) for it in items],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, spacing=1.45)
    if k < 6:
        tb(s_data, x + 1.6, 3.6, 0.22, 0.4, [("›", 18, GREY, True)], align=PP_ALIGN.CENTER)
for k, (who, color) in enumerate([("模型去写", CYAN), ("脚本去核", GREEN), ("人去定", ORANGE)]):
    chip(s_data, 0.48 + k * 1.55, 5.62, 1.4, 0.38, who, color, size=12)
tb(s_data, 5.2, 5.62, 7.6, 0.38, [("模型的产出都要经过脚本这一关，关键决定停下等人", 12.5, GREY, False)])
rect(s_data, 0.48, 6.22, 12.37, 0.62, CARD, line=MID)
tb(s_data, 0.68, 6.22, 12.0, 0.62, [("例 · 一份简历：邮箱收件 → 转 Markdown → 删敏感字段 → 八维评分 → 数字回原简历 → HR 决定约面 → 画像入人才库",
                                   13.5, WHITE, True)])
notes(s_data, """【男1】数据流程（约 50 秒）

【指七步】一条数据在小慧里走七步：采集、理解分派、预处理、生成分析、核验、决定、留痕。颜色就是分工：青色是模型，绿色是脚本，橙色是人。

模型做两件事：判断该交给谁，以及写，比如评分、检索、转写、读评语。它写的每一样东西，都要经过绿色的脚本这一关：进模型之前，先删敏感字段、先认身份；出模型之后，数字回原文核对，分数用脚本算，过没过看退出码。到了录用、合同、账号、定级，流程停下，等人决定。最后所有结果落在看板、回执和记录里，事后能查。

【指底栏】拿一份简历举例：邮箱收件，转成 Markdown，删掉敏感字段，八维评分，评分里的数字回原简历核对，HR 决定约不约面，画像进人才库。

下面按核心技能，一条线一条线看。""")

# ---------- 核心技能：五张技能页（左边主视觉各不相同，右边四格同一结构） ----------
ROW_TAGS = [("痛点", RED, WHITE), ("核心技术", CYAN, NAVY), ("我们的改良", GREEN, NAVY), ("和常见做法比", ORANGE, NAVY)]


def skill_page(live, color, label, title, sub, rows, hero, note):
    s = clone(live, 43)
    header(s, label=label, title=title, sub=sub)
    for sh in s.shapes:
        if sh.shape_id == 42:
            sh.fill.solid()
            sh.fill.fore_color.rgb = C(color)
        if sh.shape_id == 40:
            for r in sh.text_frame.paragraphs[0].runs:
                r.font.color.rgb = C(color)
    rect(s, 0, 1.18, 0.12, 6.32, color)
    rect(s, 0.48, 2.5, 5.75, 4.4, CARD)
    hero(s)
    for k, ((tag, tcolor, tfont), text) in enumerate(zip(ROW_TAGS, rows)):
        y = 2.5 + k * 1.14
        rect(s, 6.45, y, 6.4, 1.0, CARD)
        rect(s, 6.45, y, 0.07, 1.0, tcolor)
        chip(s, 6.65, y + 0.3, 1.3, 0.4, tag, tcolor, color=tfont, size=11.5)
        tb(s, 8.08, y + 0.04, 4.7, 0.92, [(text, 13.5, LIGHT, False)])
    notes(s, note)
    return s


def hero_router(s):
    tb(s, 0.68, 2.6, 5.4, 0.36, [("84 道分派题 · 派错模块的次数", 13, CYAN, True)])
    base = 5.75
    for k, (n, name, col) in enumerate([(20, "微调前小模型", RED), (2, "LoRA 微调后", CYAN), (1, "云端主链路", GREY)]):
        x = 0.95 + k * 1.75
        h = max(0.12, n * 0.11)
        rect(s, x, base - h, 1.2, h, col)
        tb(s, x - 0.2, base - h - 0.62, 1.6, 0.58, [(str(n), 30, col, True)], align=PP_ALIGN.CENTER)
        tb(s, x - 0.25, base + 0.06, 1.7, 0.34, [(name, 11.5, LIGHT, True)], align=PP_ALIGN.CENTER)
    rect(s, 0.75, base, 5.2, 0.02, GREY)
    tb(s, 0.68, 6.3, 5.4, 0.5, [("单卡 76 秒 · 只训 1.2% 参数 · 适配器 74 MB · 本地约 0.1 秒/条", 11, GREY, False)],
       align=PP_ALIGN.CENTER)


def hero_recruit(s):
    tb(s, 0.68, 2.6, 5.4, 0.36, [("回核拦下的一处真实错误", 13, CYAN, True)])
    for x, head, val, col, fill in [(0.75, "评分里写的", "58 万笔/秒", "FF9DB0", WINE), (3.55, "原简历上写的", "12.8 万 TPS", GREEN, DEEP)]:
        rect(s, x, 3.05, 2.4, 1.35, fill, line=col)
        tb(s, x, 3.12, 2.4, 0.32, [(head, 11, GREY, False)], align=PP_ALIGN.CENTER)
        tb(s, x, 3.48, 2.4, 0.7, [(val, 22, col, True)], align=PP_ALIGN.CENTER)
    tb(s, 3.1, 3.4, 0.5, 0.6, [("≠", 26, WHITE, True)], align=PP_ALIGN.CENTER)
    for x, big, col, small in [(0.75, "11 / 23", CYAN, "个数字对不上原简历"), (3.55, "0", GREEN, "评分输入剩余敏感字段")]:
        rect(s, x, 4.65, 2.4, 1.55, DEEP)
        tb(s, x, 4.72, 2.4, 0.8, [(big, 34, col, True)], align=PP_ALIGN.CENTER)
        tb(s, x, 5.55, 2.4, 0.5, [(small, 11.5, LIGHT, False)], align=PP_ALIGN.CENTER)
    tb(s, 0.68, 6.35, 5.4, 0.4, [("8 类敏感字段评分前物理删除，残留即 exit 1", 11, GREY, False)], align=PP_ALIGN.CENTER)


def hero_onboard(s):
    tb(s, 0.68, 2.6, 5.4, 0.36, [("入职时间轴上的两道门", 13, ORANGE, True)])
    steps = [("报到前 · 五个部门准备", False, "HCM 进度只读"),
             ("停 ｜ 合同门 · HR 主管批", True, "存档 · 进程退出"),
             ("停 ｜ 账号门 · IT 安全批", True, "带恢复钥匙才继续"),
             ("报到后 · 第 7/30/60/90 天跟进", False, "到点提醒 HR")]
    for k, (text, gate, side) in enumerate(steps):
        y = 3.05 + k * 0.78
        if gate:
            rect(s, 0.75, y, 3.5, 0.55, ORANGE)
            tb(s, 0.75, y, 3.5, 0.55, [(text, 14, NAVY, True)], align=PP_ALIGN.CENTER)
        else:
            rect(s, 0.75, y, 3.5, 0.55, DEEP, line=MID)
            tb(s, 0.75, y, 3.5, 0.55, [(text, 13, WHITE, True)], align=PP_ALIGN.CENTER)
        tb(s, 4.35, y, 1.8, 0.55, [(side, 10.5, ORANGE if gate else GREY, gate)])
        if k < 3:
            tb(s, 2.3, y + 0.53, 0.4, 0.26, [("↓", 12, GREY, True)], align=PP_ALIGN.CENTER)
    tb(s, 0.68, 6.2, 5.4, 0.5, [("人不批，一步不动；超时只催办，不放行", 14, ORANGE, True)], align=PP_ALIGN.CENTER)


def hero_admin(s):
    tb(s, 0.68, 2.6, 5.4, 0.36, [("同一句话：“请查询员工1的合同薪资”", 13, PINK, True)])
    for k, (who, ok, text) in enumerate([("HR", True, "过闸读 HCM，返回结果"),
                                         ("员工", False, "闸门拒绝 · exit 3 · 没读数据"),
                                         ("未登记", False, "只收到配对码，进不了 Agent")]):
        y = 3.1 + k * 0.95
        rect(s, 0.75, y, 5.2, 0.78, DEEP, line=GREEN if ok else RED)
        chip(s, 0.9, y + 0.17, 1.05, 0.44, who, PINK, size=13)
        tb(s, 2.05, y, 0.5, 0.78, [("✓" if ok else "✗", 20, GREEN if ok else RED, True)], align=PP_ALIGN.CENTER)
        tb(s, 2.55, y, 3.35, 0.78, [(text, 13.5, WHITE, True)])
    tb(s, 0.68, 6.05, 5.4, 0.7, [("发送人由飞书网关给出，不看消息内容；", 11, GREY, False),
                                 ("模型手里没有这条数据，也就谈不上说漏嘴", 11, GREY, False)], align=PP_ALIGN.CENTER)


def hero_perf(s):
    tb(s, 0.68, 2.6, 5.4, 0.36, [("员工1 · 责任心：一个分数，三方各管一段", 13, GREEN, True)])
    blocks = [("脚本", GREEN, "分组求均值，算极差", "差距 1.00 ≥ 1 → 疑似分歧"),
              ("Agent", CYAN, "读评语：都说他负责，没有负面", "改判：共识"),
              ("人", ORANGE, "主管面谈后定级", "S / A / B / C 不由系统给")]
    for k, (who, col, how, out) in enumerate(blocks):
        y = 3.05 + k * 1.07
        rect(s, 0.75, y, 5.2, 0.9, DEEP, line=col)
        chip(s, 0.9, y + 0.23, 1.0, 0.44, who, col, size=13)
        tb(s, 2.05, y + 0.06, 3.8, 0.36, [(how, 11.5, GREY, False)])
        tb(s, 2.05, y + 0.42, 3.8, 0.42, [(out, 15, col, True)])
        if k < 2:
            tb(s, 3.1, y + 0.86, 0.4, 0.22, [("↓", 11, GREY, True)], align=PP_ALIGN.CENTER)
    tb(s, 0.68, 6.35, 5.4, 0.4, [("极差只提醒哪里该看，结论要读了评语才能下", 11, GREY, False)], align=PP_ALIGN.CENTER)


s_ctl = skill_page(
    V[12], CYAN, "04  CONTROL AGENT · INTENT ROUTING",
    "意图路由：一句话进对专线，总控自己不动手",
    "六类入口汇到一个总控；派错一次，任务就进了错的专线",
    ["眼镜、WebChat、命令行都能下任务，一句话可能涉及几个模块，派错就越界",
     "总控只有派发工具；派发必须点名专线，专线之间不互派，说不清就问人",
     "Qwen2.5-1.5B 做 LoRA 微调当本地对照：派错 20 次降到 2 次；主链路仍用云端",
     "常见是一个全能 Agent 包办；我们判断与执行分开，本地模型可作断网兜底"],
    hero_router,
    """【男1】核心技能一 · 意图路由（约 50 秒）

【指痛点】眼镜、WebChat、命令行都能下任务，一句话可能涉及好几个模块。派错一次，任务就进了错的专线。

【指核心技术】总控只有派发工具，自己不执行、不写文件；派发必须点名专线，专线之间不能互相派活；说不清的时候问人。

【指左侧柱图】分派这件事，能不能交给本地小模型？我们用 1.5B 的开源模型做 LoRA 微调，和主链路的云端模型考同一套 84 道题：微调前派错 20 次，微调后 2 次，云端 1 次。单卡训练 76 秒，只训 1.2% 的参数，本地每条约 0.1 秒。差距很小，但派错一次代价太大，所以主链路仍用云端，本地留作断网兜底。

下面现场看眼镜派单。""")

s_rec = skill_page(
    V[17], CYAN, "04  RECRUITMENT AGENT · NUMBER BACKTRACE",
    "防幻觉回核：评分里的每个数字，都回原简历对一遍",
    "评分前删敏感字段，评分后数字回核；模型可以写，但不能编",
    ["模型会把岗位要求里的数字安到候选人头上；年龄、性别、院校进了评分，偏见排除不了",
     "评分前递归删除 8 类敏感字段并自检；评分后脚本抽出全部数字，回原简历比对",
     "混合回核：脚本先比，对不上的才交模型判断，最后由脚本改回原值",
     "常见是提示“请勿编造”或让模型自查；我们只认原简历，过没过看退出码"],
    hero_recruit,
    """【男1】核心技能二 · 防幻觉回核（约 55 秒）

【指痛点】招聘这条线，“说得不准”有两种：一是看了不该看的，年龄、性别、院校进了评分，就排除不了偏见，这是《就业促进法》不允许的；二是写了简历没有的。

【指左侧】开发中真出过：模型把岗位要求里的 58 万笔每秒，写进了候选人的评分，而他的简历上是 12.8 万 TPS。一份评分里抽出 23 个数字，11 个对不上原简历。

【指核心技术】所以我们设了两道闸：评分之前，脚本递归删掉 8 类敏感字段，删完自检，残留一个就退出码 1；评分之后，脚本把评分里的数字全抽出来，回原简历比对。【指改良】对不上的才交给模型判断是不是换了说法，最后由脚本改回原值。

和常见做法比，我们不靠提示词叮嘱“请勿编造”，只认原简历，过没过看退出码。现场看。""")

s_onb = skill_page(
    V[22], ORANGE, "04  ONBOARDING AGENT · APPROVAL GATE",
    "确定性审批门：流程停在门上，人不批，一步不动",
    "合同归 HR 主管，账号归 IT 安全；Agent 只推进，不代批",
    ["只在提示词里写“请等人批准”，停没停只能听模型自己说，这是假停",
     "Lobster 工作流把审批写死：到门即存档、交出恢复钥匙、进程退出",
     "恢复必须带钥匙、明确批或不批；超时只催办不放行；HCM 五节点全绿才注册定时跟进",
     "常见是靠提示词让模型自觉暂停；我们由流程引擎真停，记录可查"],
    hero_onboard,
    """【女1】核心技能三 · 确定性审批门（约 50 秒）

【指痛点】入职跨五个部门，合同涉及劳动关系，账号涉及系统安全。只在提示词里写“请等人批准”，停没停只能听它自己说，这是假停。

【指左侧时间轴】我们要的是真停。入职 Agent 只读进度、推进审批、到点提醒。走到合同门、账号门，Lobster 工作流引擎把进度存档，交出审批问题和一把恢复钥匙，然后进程退出。流程不是在“等”，是已经停了。

【指改良】恢复必须带钥匙、明确批或不批，批一次只走一步；超时只催办，不放行；HCM 五个部门节点全绿，才注册第 7、30、60、90 天的定时跟进。

人不批，一步不动。现场看。""")

s_adm = skill_page(
    V[27], PINK, "04  ADMIN AGENT · IDENTITY GATE",
    "按身份分权 + 原文有据：同一句话，三种身份三种结果",
    "读数据之前先认身份；答制度只用公司原文，并标出处",
    ["员工查同事工资，Agent 照查照答；答制度凭常识，说不清出自公司哪份文件",
     "飞书配对认人，按发送人分配工具；闸门脚本在读数据前判，非 HR 一律 exit 3",
     "14 份制度进知识库，照原文答、标出处；知识库无命中就失败退出，转人工",
     "常见是在提示词里写“不要泄露”；我们读之前就拒绝，工具层面就没给"],
    hero_admin,
    """【男2】核心技能四 · 按身份分权、原文有据（约 50 秒）

58% 的受访者说不清，智能体能碰哪些数据。

【指痛点】行政的风险有两个：员工查同事的工资，Agent 照查照答；员工问制度，Agent 凭常识答一句“按国家规定”，说不清出自公司哪份文件。

【指左侧】同一句“请查询员工1的合同薪资”：HR 过闸，读到结果；已登记的员工，闸门脚本在读数据之前就拒绝，退出码 3，查询脚本根本不运行；没登记的账号，只收到一个配对码，消息都到不了 Agent。

【指改良】答制度这边，14 份制度进知识库，照原文答、标出处，没有命中就失败退出、转人工。

和常见做法比，我们不靠提示词说“不要泄露”，而是模型手里根本没有这条数据。现场看。""")

s_perf = skill_page(
    V[31], GREEN, "04  PERFORMANCE AGENT · HYBRID JUDGEMENT",
    "统计 + 语义混合判定：脚本标疑似，Agent 读评语，人来定级",
    "极差只说明分数不一致，是不是真分歧，要读了评语才知道",
    ["360 只看平均分会掩盖分歧；只看分差，又会把期望不同误判成分歧",
     "脚本按评价人身份分组求均值，极差 ≥ 1 标“疑似分歧”；缺自评即停，不出总表",
     "Agent 逐条读评语复核，可把疑似改判为共识或分歧，结论写进面谈材料",
     "常见是模型直接给结论甚至打等级；我们不定级，S/A/B/C 由主管面谈后定"],
    hero_perf,
    """【女2】核心技能五 · 统计加语义的混合判定（约 50 秒）

【指痛点】360 反馈只看平均分，会把分歧抹平；只看分差，又会把“期望不同”误判成分歧。

【指左侧三段】分工是：脚本管数字，Agent 读评语，人定评级。脚本按上级、平级、下属、自评分组求平均，最高组减最低组，到 1 分就标“疑似分歧”。拿员工1的责任心举例：差距正好 1 分，脚本标了疑似；Agent 读评语，所有人都说他负责，没有一条负面，上级打 4 分只是希望他负责的范围再大一些，所以改判为共识。

评级不由系统给：《个人信息保护法》要求，对个人权益有重大影响的决定，不能只由自动化方式作出。S、A、B、C 由主管面谈后定。现场看。""")

for live, pairs in ((V[12], [("和上一页的配置对得上", "和前面的配置对得上"), ("【→ 接 03 招聘】", "【→ 接 招聘 · 防幻觉回核】")]),
                    (V[17], [("就是风险页上那个数字", "就是上一页那个数字")])):
    nt = live.notes_slide.notes_text_frame
    for a, b in pairs:
        nt.text = nt.text.replace(a, b)

# ---------- 价值创新：创新点（新） ----------
s_inn = clone(V[8], 29)
header(s_inn, label="06  INNOVATION", title="创新点：不是训练更大的模型，而是把模型放到对的位置",
       sub="方法、机制、技术三个层面，都能回到前面的页面和现场")
COLS = [
    ("01", "范式创新", "OPD 开发范式", PURPLE,
     ["部门 SOP → 能力盘点 → 业务分析 → 架构 → 分阶段开发", "HR 是第一个实例，换部门照四步复用"]),
    ("02", "机制创新", "闸门写死，不靠提示词", GREEN,
     ["法进代码：评分前删敏感字段，读数据前认身份", "规进流程：审批门真停，无依据不答", "8 道闸门实测通过，过没过看退出码"]),
    ("03", "技术改良", "混合判定", CYAN,
     ["脚本比对 + 模型判疑：数字回核", "统计极差 + 语义复核：360 分歧", "本地 LoRA 分派：派错 20 → 2 次"]),
]
for k, (num, kind, name, col, lines) in enumerate(COLS):
    x = 0.48 + k * 4.18
    rect(s_inn, x, 2.2, 3.98, 3.75, CARD)
    rect(s_inn, x, 2.2, 3.98, 0.08, col)
    tb(s_inn, x + 0.25, 2.5, 1.2, 0.8, [(num, 36, col, True)])
    tb(s_inn, x + 1.3, 2.5, 2.6, 0.3, [(kind, 12, GREY, True)])
    tb(s_inn, x + 1.3, 2.82, 2.6, 0.45, [(name, 16, WHITE, True)])
    rect(s_inn, x + 0.25, 3.5, 3.48, 0.02, MID)
    tb(s_inn, x + 0.25, 3.7, 3.5, 2.1, [(ln, 14.5, LIGHT, False) for ln in lines],
       anchor=MSO_ANCHOR.TOP, spacing=1.5)
rect(s_inn, 0.48, 6.2, 12.37, 0.6, CARD, line=MID)
tb(s_inn, 0.48, 6.2, 12.37, 0.6, [("让模型去写  ·  让脚本去核  ·  让人去定", 16, CYAN, True)], align=PP_ALIGN.CENTER)
notes(s_inn, """【男1】创新点（约 40 秒）

我们的难点不在训练一个更大的模型，而在把会出错的模型放到对的位置上。创新落在三个层面。

【指 01】方法上，是基于部门 SOP 的 OPD 开发范式，HR 是第一个实例，换部门照四步复用。

【指 02】机制上，闸门写死，不靠提示词：法进代码，评分前删敏感字段、读数据前认身份；规进流程，审批门真停、无依据不答。八道闸门全部实测通过，过没过看退出码。

【指 03】技术上，是几处混合判定：脚本比对加模型判疑的数字回核，统计极差加语义复核的 360 分歧判断，以及本地 LoRA 分派对照。""")

# ---------- 封面底栏：四段改六段 ----------
s = V[3]
remove(s, {25, 28, 31})
dot, txt = by_id(s, 32), by_id(s, 33)
for _ in range(2):
    for src in (dot, txt):
        el = copy.deepcopy(src._element)
        el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set(
            "id", str(max(sh.shape_id for sh in s.shapes) + 1))
        txt._element.addnext(el)
dots = [sh for sh in s.shapes if sh.top / E > 6.8 and sh.shape_type == 1 and "OVAL" in str(sh.auto_shape_type)]
txts = [sh for sh in s.shapes if sh.top / E > 6.7 and sh.has_text_frame and sh.text_frame.text.strip()]
for k, (d, t) in enumerate(zip(dots, txts)):
    d.left = Inches(0.75 + k * 2.1)
    t.left = Inches(1.0 + k * 2.1)
    t.width = Inches(1.6)
    set_text(t, ["项目背景", "项目设计", "工作流程", "核心技能", "主要成果", "价值创新"][k])

# ---------- 重排、删页 ----------
ORDER = [V[1], V[2], V[3], V[4], V[5], V[6], V[7],          # 开场 + 项目背景
         V[8], V[9], V[11],                                  # 项目设计
         V[35], s_data,                                      # 工作流程
         s_ctl, V[12], s_rec, V[17], s_onb, V[22], s_adm, V[27], s_perf, V[31],   # 核心技能
         V[32], V[33],                                       # 主要成果
         s_inn, V[36],                                       # 价值创新
         V[37]]
SECTION = {}
for i in range(1, 28):
    SECTION[i] = (None if i in (2, 3, 27) else
                  0 if i <= 7 else 1 if i <= 10 else 2 if i <= 12 else 3 if i <= 22 else 4 if i <= 24 else 5)

lst = p.slides._sldIdLst
els = {id(s): el for s, el in zip(p.slides, list(lst))}
keep = [els[id(s)] for s in ORDER]
for el in list(lst):
    if el not in keep:
        p.part.drop_rel(el.get(R + "id"))
    lst.remove(el)
for el in keep:
    lst.append(el)

# ---------- 六栏导航 + 小标编号 + 页码 ----------
TABS = ["项目背景", "项目设计", "工作流程", "核心技能", "主要成果", "价值创新"]
total = len(p.slides)
OLD_TABS = ("项目背景", "方案规划", "技能演示", "成果展望")


def nav_parts(s):
    near = [sh for sh in s.shapes if sh.top is not None and abs(sh.top / E - 0.12) < 0.02 and 2.9 < sh.left / E < 11.7]
    shapes = sorted([sh for sh in near if sh.shape_type == 1 and "PARALLELOGRAM" in str(sh.auto_shape_type)],
                    key=lambda x: x.left)
    texts = sorted([sh for sh in near if sh.shape_type == 17 and sh.text_frame.text.strip() in OLD_TABS],
                   key=lambda x: x.left)
    return shapes, texts


for i, s in enumerate(p.slides, 1):
    shapes, texts = nav_parts(s)
    sec = SECTION[i]
    if len(shapes) == 4 and len(texts) == 4:
        next_id = max(sh.shape_id for sh in s.shapes) + 1
        for _ in range(2):
            for src in (shapes[-1], texts[-1]):
                el = copy.deepcopy(src._element)
                el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("id", str(next_id))
                next_id += 1
                src._element.addnext(el)
        allp = [sh for sh in s.shapes if sh.top is not None and abs(sh.top / E - 0.12) < 0.02 and 2.9 < sh.left / E < 11.7]
        shapes = [sh for sh in allp if sh.shape_type == 1 and "PARALLELOGRAM" in str(sh.auto_shape_type)]
        texts = [sh for sh in allp if sh.shape_type == 17 and sh.text_frame.text.strip() in OLD_TABS]
        for k in range(6):
            for sh in (shapes[k], texts[k]):
                sh.left = Inches(3.0 + k * 1.45)
                sh.width = Inches(1.37)
            act = (k == sec) or (sec is None and k == 0)
            shapes[k].fill.solid()
            shapes[k].fill.fore_color.rgb = C(CYAN if act else MID)
            shapes[k].line.fill.background()
            set_text(texts[k], TABS[k])
            r = texts[k].text_frame.paragraphs[0].runs[0]
            r.font.color.rgb = C(NAVY if act else GREY)
            r.font.bold = act
    if sec is not None:
        for sh in s.shapes:
            if sh.has_text_frame and abs(sh.left / E - 0.48) < 0.05 and 0.85 < sh.top / E < 1.4:
                run = sh.text_frame.paragraphs[0].runs
                if run and re.match(r"^\d\d", run[0].text):
                    run[0].text = f"{sec + 1:02d}" + run[0].text[2:]
    for sh in s.shapes:
        if sh.has_text_frame and sh.top is not None and sh.top / E < 0.6 and sh.left / E > 11 \
                and re.fullmatch(r"\d{2} / \d{2}", sh.text_frame.text.strip()):
            set_text(sh, f"{i:02d} / {total:02d}")

p.save(DST)
print("saved", DST, total)
