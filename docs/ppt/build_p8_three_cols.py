"""从 v2.2 第 8 页导出单页：设计原则改为三栏（模型 / 脚本 / 人）。v2.2 不改。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v2.2.pptx"
DST = "/opt/openclaw-hr/docs/PPT准备/第8页-设计原则-三栏版.pptx"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
CARD, MID = "032B5C", "04457D"
CYAN, GREEN, ORANGE = "00D8FF", "26D098", "FFA421"
WHITE, LIGHT, GREY = "FFFFFF", "E9F8FF", "9EC2D8"
C = RGBColor.from_string

p = Presentation(SRC)
lst = p.slides._sldIdLst
for k, el in enumerate(list(lst)):
    if k != 7:
        p.part.drop_rel(el.get(R + "id"))
        lst.remove(el)
s = p.slides[0]

for sh in list(s.shapes):
    if 30 <= sh.shape_id <= 86:
        sh._element.getparent().remove(sh._element)


def set_text(sh, text):
    paras = sh.text_frame.paragraphs
    runs = paras[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""
    for para in paras[1:]:
        para._p.getparent().remove(para._p)


def rect(l, t, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = C(fill)
    if line:
        r.line.color.rgb = C(line)
        r.line.width = Pt(1)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def tb(l, t, w, h, paras, anchor=MSO_ANCHOR.TOP, spacing=None, font="微软雅黑"):
    b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, (text, size, color, bold) in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        if spacing:
            para.line_spacing = spacing
        r = para.add_run()
        r.text = text
        f = r.font
        f.size, f.color.rgb, f.bold, f.name = Pt(size), C(color), bold, font
        rpr = r._r.get_or_add_rPr()
        rpr.append(rpr.makeelement(A + "ea", {"typeface": font}))
    return b


set_text(next(sh for sh in s.shapes if sh.shape_id == 29),
         "用确定的脚本，框住不确定的生成；把会影响人的决定，交还给人")

COLS = [
    ("01", "模型：能写，也会错", CYAN,
     ["读得懂简历，写得出评语，", "答得了制度。", "", "但产出不确定，会把岗位里的", "数字安到候选人头上。"],
     ["评分写“58 万笔/秒”", "原简历是“12.8 万 TPS”"]),
    ("02", "脚本：不认说法，只认原件", GREEN,
     ["数字对原简历，回答对制度原文，", "查询先对权限；对不上就拦下。", "", "提示词会被忘、被说动；", "脚本过没过，看退出码。"],
     ["评分前敏感字段残留 0", "非 HR 查薪资，退出码 3"]),
    ("03", "人：关键一步，亲自拍板", ORANGE,
     ["录用、合同、账号、定级，", "关系到一个人的去留和权益。", "", "小慧递材料、给建议，", "签字的始终是人。"],
     ["审批门停下，等人来批", "绩效只给材料，不出评级"]),
]
for k, (num, head, col, body, proof) in enumerate(COLS):
    x = 0.48 + k * 4.2
    rect(x, 2.2, 3.97, 3.86, CARD, line=MID)
    rect(x, 2.2, 0.07, 3.86, col)
    tb(x + 0.3, 2.36, 1.0, 0.36, [(num, 15, col, True)], anchor=MSO_ANCHOR.MIDDLE)
    tb(x + 0.3, 2.74, 3.55, 0.5, [(head, 19, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
    tb(x + 0.3, 3.3, 3.55, 1.75, [(ln, 14, LIGHT, False) for ln in body], spacing=1.2)
    rect(x + 0.3, 5.0, 3.4, 0.015, MID)
    tb(x + 0.3, 5.07, 1.2, 0.26, [("现场可见", 10.5, col, True)], anchor=MSO_ANCHOR.MIDDLE)
    tb(x + 0.3, 5.35, 3.55, 0.62, [(ln, 12.5, WHITE, True) for ln in proof], anchor=MSO_ANCHOR.MIDDLE)

s.notes_slide.notes_text_frame.text = """【男1】让模型去写，让脚本去核，让人去定（约 60 秒）

我们的答案是一句话：让模型去写，让脚本去核，让人去定。

【指 01】模型读得懂、写得快：解析简历、写评语、按原文答制度。但它的产出是不确定的，会错。开发中它就把岗位要求里的 58 万笔每秒，写进了候选人的评分，原简历上是 12.8 万 TPS。写得越快，越需要有人把关。

【指 02】把关的是脚本。脚本不认说法，只认原件：数字对原简历，回答对制度原文，查询先对权限，对不上就拦下，模型改不了。为什么不在提示词里多叮嘱几句？因为提示词可能被忘、被说动，照没照做只能听模型自己说；脚本是写死的检查，过没过，看退出码。

【指 03】录用、合同、账号、定级，关系到一个人的去留和权益。小慧递材料、给建议，签字的始终是人。

【指底栏】写得快，核得准，定得稳。展开说，就是三句话：把法写进代码，把规写进流程，把决定留给人。每栏下面这一行，后面现场都会看到。"""

p.save(DST)
print("saved", DST)
