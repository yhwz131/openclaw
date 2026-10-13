"""样张：一个模块 5 页，光随内容推进向右（A 光斑右移 / B 铺得越来越远）。v2.2 不改。"""
from lxml import etree
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

COVER_SCRIPT = "/opt/openclaw-hr/docs/ppt/build_covers_overview.py"
import sys
MODULES = {"招聘": ([14, 15, 16, 17, 18], "00D8FF"), "行政": ([24, 25, 26, 27, 28], "FF6B8A")}
NAME = sys.argv[1] if len(sys.argv) > 1 else "行政"
PAGES, COLOR = MODULES[NAME]
OUT = {"A": f"/opt/openclaw-hr/docs/PPT准备/光推进样张-{NAME}-A光斑右移.pptx",
       "B": f"/opt/openclaw-hr/docs/PPT准备/光推进样张-{NAME}-B铺得越来越远.pptx",
       "B2": f"/opt/openclaw-hr/docs/PPT准备/光推进样张-{NAME}-B2亮度不减铺满.pptx"}


def q(tag):
    return "{http://schemas.openxmlformats.org/drawingml/2006/main}%s" % tag


def load():
    ns = {}
    code = open(COVER_SCRIPT, encoding="utf-8").read()
    code = code.replace("KEEP = [14, 19, 24, 29]", "KEEP = %r" % PAGES)
    code = code.replace("p.save(DST)", "pass").replace('print("saved", DST, len(p.slides))', "")
    code = _guard_modules(code)
    exec(compile(code, COVER_SCRIPT, "exec"), ns)
    return ns


def _guard_modules(code):
    """只保留当前模块封面的改写，其余模块那几段跳过。"""
    out, skip = [], False
    for line in code.splitlines():
        if line.startswith("# ---------- "):
            skip = NAME not in line
        elif line.startswith("p.save") or line.startswith("pass"):
            skip = False
        if not skip:
            out.append(line)
    return "\n".join(out)


def light(s, stops):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    r.line.fill.background()
    r.shadow.inherit = False
    spPr = r._element.spPr
    for child in list(spPr):
        if child.tag in (q("solidFill"), q("gradFill"), q("noFill")):
            spPr.remove(child)
    grad = etree.Element(q("gradFill"), rotWithShape="1")
    lst = etree.SubElement(grad, q("gsLst"))
    for pos, a in stops:
        gs = etree.SubElement(lst, q("gs"), pos=str(int(pos * 1000)))
        clr = etree.SubElement(gs, q("srgbClr"), val=COLOR)
        etree.SubElement(clr, q("alpha"), val=str(int(a)))
    etree.SubElement(grad, q("lin"), ang="0", scaled="0")
    spPr.insert(list(spPr).index(spPr.find(q("prstGeom"))) + 1, grad)
    bg = next(sh for sh in s.shapes if sh.shape_id == 2)._element
    bg.addnext(r._element)


def stops_reach(k, alpha):
    reach = [40, 55, 70, 85, 100][k]
    return [(0, alpha), (reach / 2, alpha / 3), (reach, 0)] + ([(100, 0)] if reach < 100 else [])


def stops_flood(k, alpha):
    full = [0, 20, 40, 65, 100][k]
    reach = [40, 55, 75, 95, 100][k]
    if full >= 100:
        return [(0, alpha), (100, alpha)]
    pts = [(0, alpha)] + ([(full, alpha)] if full > 0 else [])
    pts += [((full + reach) / 2, alpha / 3), (reach, 0)]
    return pts + ([(100, 0)] if reach < 100 else [])


def stops_move(k, alpha):
    c = [0, 25, 50, 75, 100][k]
    half = 30
    pts = [(c, alpha)]
    if c - half > 0:
        pts = [(0, 0), (c - half, 0)] + pts
    elif c > 0:
        pts = [(0, alpha * (1 - c / half))] + pts
    if c + half < 100:
        pts = pts + [(c + half, 0), (100, 0)]
    elif c < 100:
        pts = pts + [(100, alpha * (1 - (100 - c) / half))]
    return pts


def extras(ns, s, cover):
    C, E = ns["C"], ns["E"]
    for sh in s.shapes:
        if sh.shape_type == 9 and 3 <= sh.shape_id <= 9:
            ln = sh._element.spPr.find(q("ln"))
            if ln is None:
                continue
            for child in list(ln):
                if child.tag in (q("solidFill"), q("gradFill"), q("noFill")):
                    ln.remove(child)
            fill = etree.Element(q("solidFill"))
            clr = etree.SubElement(fill, q("srgbClr"), val=COLOR)
            etree.SubElement(clr, q("alpha"), val="30000" if cover else "22000")
            ln.insert(0, fill)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.18), Inches(0.12 if cover else 0.08), Inches(6.32))
    r.fill.solid()
    r.fill.fore_color.rgb = C(COLOR)
    r.line.fill.background()
    r.shadow.inherit = False


MODES = sys.argv[2].split(",") if len(sys.argv) > 2 else ["A", "B", "B2"]
FNS = {"A": stops_move, "B": stops_reach, "B2": stops_flood}
for mode in MODES:
    fn = FNS[mode]
    ns = load()
    p, S = ns["p"], ns["S"]
    for k, n in enumerate(PAGES):
        s = S[n]
        cover = k == 0
        light(s, fn(k, 24000 if cover or mode == "B2" else 17000))
        extras(ns, s, cover)
    p.save(OUT[mode])
    print("saved", OUT[mode], len(p.slides))
