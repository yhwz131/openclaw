#!/usr/bin/env python3
"""基于“26-科技教育”模板生成小慧结尾四页。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

import build_opening_template_pptx as base


ROOT = Path("/opt/openclaw-hr/docs")
ASSETS = ROOT / "ppt/assets/outcomes"
CACHE = ASSETS / "closing-crops"
OUT = ROOT / "小慧-结尾部分-模板初版.pptx"
PREVIEW = ROOT / "小慧-结尾部分-模板初版-预览.png"
TOTAL = 4

GREEN = base.RGBColor(0x26, 0xD0, 0x98)
GREEN_DARK = base.RGBColor(0x08, 0x50, 0x57)
PURPLE = base.RGBColor(0x8C, 0x7C, 0xFF)
RED = base.RGBColor(0xF1, 0x5A, 0x66)
END_ITEMS = ["一期成果", "项目创新", "价值规划", "致谢"]


def closing_subnav(slide, active: int) -> None:
    base.rect(slide, Inches(0), Inches(0.72), base.W, Inches(0.45), base.NAVY, line=base.LINE)
    base.text(
        slide,
        Inches(0.4),
        Inches(0.79),
        Inches(1.15),
        Inches(0.27),
        "项目收束",
        10,
        bold=True,
        color=base.MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    for i, label in enumerate(END_ITEMS):
        x = Inches(2.05 + i * 2.45)
        base.rect(
            slide,
            x,
            Inches(0.78),
            Inches(2.15),
            Inches(0.29),
            base.CYAN if i == active else base.NAVY_2,
            line=base.CYAN if i == active else base.LINE,
            radius=True,
        )
        base.text(
            slide,
            x,
            Inches(0.78),
            Inches(2.15),
            Inches(0.29),
            label,
            10.5,
            bold=i == active,
            color=base.NAVY if i == active else base.MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def closing_header(slide, page: int, active: int, eyebrow: str, title: str, subtitle: str) -> None:
    base.add_navigation(slide, 4, page)
    closing_subnav(slide, active)
    base.text(slide, Inches(0.48), Inches(1.28), Inches(3.4), Inches(0.25), eyebrow, 10.5, bold=True, color=base.CYAN)
    base.text(slide, Inches(0.46), Inches(1.53), Inches(12.2), Inches(0.49), title, 23, bold=True, font=base.TITLE_FONT)
    base.rect(slide, Inches(0.47), Inches(2.08), Inches(0.75), Inches(0.05), base.CYAN)
    base.text(slide, Inches(1.4), Inches(1.99), Inches(11.1), Inches(0.29), subtitle, 11.5, color=base.MUTED)


def crop_image(src: Path, name: str, ratio: float = 16 / 9, focus_y: float = 0.0) -> Path:
    """裁成横向成果缩略图；只裁视觉，不改原产物。"""
    CACHE.mkdir(parents=True, exist_ok=True)
    dest = CACHE / name
    image = Image.open(src).convert("RGB")
    w, h = image.size
    target_h = int(w / ratio)
    if target_h <= h:
        top = int(max(0, min(h - target_h, (h - target_h) * focus_y)))
        image = image.crop((0, top, w, top + target_h))
    else:
        target_w = int(h * ratio)
        left = max(0, (w - target_w) // 2)
        image = image.crop((left, 0, left + target_w, h))
    image.save(dest, quality=94)
    return dest


def image_tile(slide, path: Path, x, y, title: str, sub: str, accent=base.CYAN) -> None:
    base.rect(slide, x, y, Inches(3.75), Inches(2.08), base.NAVY_2, line=base.LINE, radius=True)
    slide.shapes.add_picture(str(path), x + Inches(0.08), y + Inches(0.08), width=Inches(3.59), height=Inches(1.45))
    base.rect(slide, x + Inches(0.08), y + Inches(1.53), Inches(3.59), Inches(0.46), base.NAVY_3)
    base.text(slide, x + Inches(0.2), y + Inches(1.59), Inches(1.65), Inches(0.27), title, 11.5, bold=True, color=accent)
    base.text(slide, x + Inches(1.82), y + Inches(1.59), Inches(1.7), Inches(0.27), sub, 9.5, color=base.MUTED, align=PP_ALIGN.RIGHT)


def result_slide(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    closing_header(
        slide,
        1,
        0,
        "05  PHASE I DELIVERABLES",
        "从设计到落地：一期建设成果",
        "完整架构已经转化为可运行、可检查、可交付的业务流程与产物",
    )

    # 左侧总览
    base.rect(slide, Inches(0.48), Inches(2.4), Inches(4.18), Inches(3.9), base.NAVY_2, line=base.CYAN, radius=True)
    base.text(slide, Inches(0.74), Inches(2.62), Inches(3.64), Inches(0.34), "小慧 · HR数字员工", 18, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(0.74), Inches(3.04), Inches(3.64), Inches(0.38), "1个协调Agent + 5个业务Agent", 14, bold=True, align=PP_ALIGN.CENTER)
    modules = [
        ("总控", "识别·分发·返回"),
        ("招聘", "画像·人才库·纪要"),
        ("入职", "追踪·审批·跟进"),
        ("培训", "按需生成资料"),
        ("绩效", "360聚合·报告看板"),
        ("行政", "问答·权限·月报"),
    ]
    for i, (name, body) in enumerate(modules):
        col, row = i % 2, i // 2
        x = Inches(0.73 + col * 1.9)
        y = Inches(3.66 + row * 0.66)
        base.rect(slide, x, y, Inches(1.72), Inches(0.5), base.NAVY_3, line=base.LINE, radius=True)
        base.text(slide, x + Inches(0.1), y + Inches(0.06), Inches(0.5), Inches(0.24), name, 10.5, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(0.58), y + Inches(0.06), Inches(1.02), Inches(0.28), body, 8.5, color=base.PALE, align=PP_ALIGN.RIGHT)
    base.rect(slide, Inches(0.73), Inches(5.85), Inches(3.66), Inches(0.29), base.NAVY_3, line=GREEN, radius=True)
    base.text(slide, Inches(0.73), Inches(5.85), Inches(3.66), Inches(0.29), "可运行  ·  可检查  ·  可交付", 10.5, bold=True, color=GREEN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # 四张仓库真实产物截图。
    images = [
        (crop_image(ASSETS / "recruitment-profile.png", "recruitment-profile-wide.png", focus_y=0.0), "招聘画像", "八维评分与画像", base.CYAN),
        (crop_image(ASSETS / "onboarding-dashboard.png", "onboarding-dashboard-wide.png", focus_y=0.1), "入职看板", "HCM节点与跟进", GREEN),
        (crop_image(ASSETS / "performance-dashboard.png", "performance-dashboard-wide.png", focus_y=0.0), "绩效看板", "360反馈聚合", PURPLE),
        (crop_image(ASSETS / "monthly-report.png", "monthly-report-wide.png", focus_y=0.0), "行政月报", "报告期 2026-09", base.ORANGE),
    ]
    positions = [
        (Inches(4.94), Inches(2.4)),
        (Inches(8.91), Inches(2.4)),
        (Inches(4.94), Inches(4.66)),
        (Inches(8.91), Inches(4.66)),
    ]
    for (path, title, sub, accent), (x, y) in zip(images, positions):
        image_tile(slide, path, x, y, title, sub, accent)

    base.text(
        slide,
        Inches(0.5),
        Inches(6.55),
        Inches(4.16),
        Inches(0.22),
        "实际连接：WebChat / CLI / 飞书 / HCM / 本地人才库",
        9,
        color=base.MUTED,
        align=PP_ALIGN.CENTER,
    )
    base.add_notes(
        slide,
        "以上是小慧从接收任务、分发任务，到各业务模块完成处理并返回结果的完整流程。"
        "经过一期建设，我们完成了1协调+5模块的整体架构，并将招聘、入职、培训、绩效和行政五类HR业务"
        "形成可运行、可检查、可交付的业务流程。右侧展示的是仓库当前真实生成的招聘画像、入职看板、"
        "360反馈看板和HR月报，不把教学数据说成企业上线数据。",
    )


def innovation_card(slide, x, y, number: str, title: str, subtitle: str, body: str, accent) -> None:
    base.rect(slide, x, y, Inches(5.9), Inches(1.82), base.NAVY_2, line=base.LINE, radius=True)
    base.rect(slide, x, y, Inches(0.1), Inches(1.82), accent)
    base.text(slide, x + Inches(0.22), y + Inches(0.16), Inches(0.52), Inches(0.33), number, 15, bold=True, color=accent)
    base.text(slide, x + Inches(0.78), y + Inches(0.12), Inches(2.2), Inches(0.35), title, 16, bold=True)
    base.text(slide, x + Inches(2.95), y + Inches(0.16), Inches(2.65), Inches(0.28), subtitle, 9.5, color=base.MUTED, align=PP_ALIGN.RIGHT)
    base.text(slide, x + Inches(0.24), y + Inches(0.64), Inches(5.36), Inches(0.96), body, 11.3, color=base.PALE, line_spacing=0.93)


def innovation_slide(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    closing_header(
        slide,
        2,
        1,
        "05  INNOVATION DESIGN",
        "项目创新：不是新名词，而是落到流程里的改造",
        "四类创新都能回到前面的系统结构、现场操作与结果证据",
    )
    innovation_card(
        slide,
        Inches(0.48),
        Inches(2.4),
        "01",
        "场景创新",
        "从政策场景到HR业务",
        "将OPC的人机协作思路应用到HR部门，形成OPD项目方案；从6个模块、30多个岗位子项中规划一期范围。",
        base.CYAN,
    )
    innovation_card(
        slide,
        Inches(6.43),
        Inches(2.4),
        "02",
        "架构创新",
        "集中调度，业务隔离",
        "Router统一接收任务，五个业务Agent按HR边界分工；不同入口、数据和工具交给对应模块，结果统一返回HR。",
        PURPLE,
    )
    innovation_card(
        slide,
        Inches(0.48),
        Inches(4.43),
        "03",
        "机制创新",
        "非确定产出，可确定把关",
        "招聘脱敏与数字回核；入职暂停审批；培训按需生成与自检；绩效缺数据不聚合；行政校身份、无依据不回答。",
        base.ORANGE,
    )
    innovation_card(
        slide,
        Inches(6.43),
        Inches(4.43),
        "04",
        "决策边界创新",
        "把“人在回路”落实到节点",
        "招聘不决定录用；入职不代审批；培训资料由讲师审核；绩效不定最终等级；行政个案与争议转人工。",
        GREEN,
    )
    base.rect(slide, Inches(0.48), Inches(6.5), Inches(11.85), Inches(0.3), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(
        slide,
        Inches(0.68),
        Inches(6.5),
        Inches(11.45),
        Inches(0.3),
        "以上机制均已在前面的技能演示中通过实际操作验证，本页只总结，不重复讲解。",
        10.5,
        bold=True,
        color=base.CYAN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "创新不逐项重讲。口播：项目创新集中在四个层面。第一是场景，从OPC的人机协作思路延伸到HR部门；"
        "第二是Router集中调度和业务隔离；第三是将合规、校验、暂停和回核写进每个模块；"
        "第四是把人和Agent的决策权限落实到具体节点。这些机制都已经在前面的技能演示中得到验证。"
        "不要说行业首创、国内领先。",
    )


def value_group(slide, x, y, title: str, items: list[str], accent) -> None:
    base.rect(slide, x, y, Inches(5.68), Inches(1.65), base.NAVY_2, line=base.LINE, radius=True)
    base.text(slide, x + Inches(0.22), y + Inches(0.15), Inches(5.2), Inches(0.3), title, 13.5, bold=True, color=accent)
    for i, item in enumerate(items):
        yy = y + Inches(0.56 + i * 0.34)
        dot = slide.shapes.add_shape(base.MSO_SHAPE.OVAL, x + Inches(0.23), yy + Inches(0.06), Inches(0.09), Inches(0.09))
        base.set_fill(dot, accent)
        dot.line.fill.background()
        base.text(slide, x + Inches(0.42), yy, Inches(5.0), Inches(0.24), item, 10.8, color=base.PALE)


def roadmap_step(slide, x, y, label: str, title: str, body: str, accent) -> None:
    base.rect(slide, x, y, Inches(5.68), Inches(0.96), base.NAVY_2, line=base.LINE, radius=True)
    base.rect(slide, x, y, Inches(0.1), Inches(0.96), accent)
    base.text(slide, x + Inches(0.22), y + Inches(0.13), Inches(0.72), Inches(0.25), label, 10.5, bold=True, color=accent)
    base.text(slide, x + Inches(0.93), y + Inches(0.09), Inches(1.58), Inches(0.29), title, 12.5, bold=True)
    base.text(slide, x + Inches(2.42), y + Inches(0.11), Inches(2.98), Inches(0.54), body, 9.8, color=base.PALE)


def value_future_slide(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    closing_header(
        slide,
        3,
        2,
        "05  VALUE & ROADMAP",
        "从可用到可落地：应用价值与未来规划",
        "不虚构效率百分比，用已覆盖的工作和明确的下一步说明价值",
    )

    base.text(slide, Inches(0.48), Inches(2.39), Inches(5.68), Inches(0.3), "应用价值", 15, bold=True, color=base.CYAN)
    value_group(
        slide,
        Inches(0.48),
        Inches(2.77),
        "业务处理",
        ["减轻简历整理与初筛负担", "让入职待办和跟进过程可见", "辅助准备培训、绩效材料及行政月报"],
        base.CYAN,
    )
    value_group(
        slide,
        Inches(0.48),
        Inches(4.62),
        "管理价值",
        ["统一调度不同HR业务", "关键结果可检查、可回核、可追溯", "让HR投入更多时间进行沟通与判断"],
        GREEN,
    )

    base.text(slide, Inches(6.67), Inches(2.39), Inches(5.68), Inches(0.3), "未来规划", 15, bold=True, color=base.ORANGE)
    roadmap_step(
        slide,
        Inches(6.67),
        Inches(2.77),
        "近期",
        "补齐业务",
        "继续覆盖Offer、员工关系等未完成子项；根据合作企业反馈完善现有模块。",
        base.CYAN,
    )
    roadmap_step(
        slide,
        Inches(6.67),
        Inches(3.9),
        "中期",
        "深化连接",
        "加强企业邮箱、飞书、HCM连接；在合规评估后建设薪酬模块。",
        PURPLE,
    )
    roadmap_step(
        slide,
        Inches(6.67),
        Inches(5.03),
        "后续",
        "完善治理",
        "使用脱敏业务数据扩充测试；完善权限、校验、回核及操作留痕。",
        base.ORANGE,
    )
    base.rect(slide, Inches(0.48), Inches(6.47), Inches(11.87), Inches(0.33), base.NAVY_3, line=base.ORANGE, radius=True)
    base.text(
        slide,
        Inches(0.68),
        Inches(6.47),
        Inches(11.47),
        Inches(0.33),
        "规划原则：先补业务证据与系统连接，再扩功能；强合规模块坚持人工审批，不盲目自动化。",
        10.5,
        bold=True,
        color=base.ORANGE,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "口播分左右两部分。应用价值不讲虚构的降本百分比，只讲已经覆盖的工作：简历初筛、入职可视、"
        "培训和绩效材料、行政问答与月报；管理上强调统一调度、可检查和可追溯。"
        "下一步先补Offer和员工关系等子项，深化邮箱、飞书、HCM连接，再在合规评估后建设薪酬模块，"
        "持续完善权限、校验和回核机制。",
    )


def add_cover_background(source: Presentation, slide) -> None:
    base.add_cover_background(source, slide)


def thanks_slide(prs: Presentation, source: Presentation) -> None:
    slide = base.new_slide(prs)
    add_cover_background(source, slide)
    base.text(
        slide,
        Inches(0.72),
        Inches(0.5),
        Inches(11.8),
        Inches(0.36),
        "小慧 · HR数字员工",
        13,
        bold=True,
        color=base.PALE,
    )
    base.rect(slide, Inches(0.72), Inches(2.14), Inches(1.15), Inches(0.08), base.CYAN)
    base.text(slide, Inches(0.7), Inches(2.42), Inches(9.5), Inches(0.75), "感谢各位评委", 40, bold=True, font=base.TITLE_FONT)
    base.text(slide, Inches(0.72), Inches(3.42), Inches(8.5), Inches(0.52), "小慧项目团队", 18, bold=True, color=base.CYAN)
    base.add_notes(
        slide,
        "主控总结：从OPC到OPD，小慧不是替代HR，而是与HR分工协作，让重复工作处理得更高效，"
        "让关键业务过程更加可信，让最终决策始终由人掌握。我们的项目展示完毕，感谢各位评委。"
        "随后全体成员上台鞠躬致谢。评委不能与选手交流，因此不要说“请提问”。",
    )


def main() -> None:
    source = Presentation(str(base.TEMPLATE))
    prs = Presentation(str(base.TEMPLATE))
    base.clear_slides(prs)
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL
    result_slide(prs)
    innovation_slide(prs)
    value_future_slide(prs)
    thanks_slide(prs, source)
    prs.core_properties.title = "小慧 · 结尾部分PPT初版"
    prs.core_properties.subject = "一期成果、项目创新、应用价值与未来规划"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
