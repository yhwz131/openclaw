#!/usr/bin/env python3
"""行政模块四页无图初版：政策有出处、数据有权限、月报可复核。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import build_opening_template_pptx as base


ROOT = Path("/opt/openclaw-hr/docs")
OUT = ROOT / "PPT准备/行政模块技能演示-4页无图初版.pptx"
TOTAL = 4

GREEN = RGBColor(0x26, 0xD0, 0x98)
GREEN_DARK = RGBColor(0x08, 0x50, 0x57)
PURPLE = RGBColor(0x8C, 0x7C, 0xFF)
RED = RGBColor(0xF1, 0x5A, 0x66)
RED_DARK = RGBColor(0x56, 0x17, 0x26)


def connector(slide, x1, y1, x2, y2, color=base.CYAN, width=1.1, arrow=True):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def secondary_nav(slide) -> None:
    labels = ["总控", "招聘", "入职", "培训", "绩效", "行政"]
    base.rect(slide, Inches(0), Inches(0.72), base.W, Inches(0.45), base.NAVY, line=base.LINE)
    base.text(
        slide,
        Inches(0.38),
        Inches(0.79),
        Inches(0.9),
        Inches(0.27),
        "技能地图",
        10,
        bold=True,
        color=base.MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    for i, label in enumerate(labels):
        x = Inches(1.33 + i * 1.77)
        active = i == 5
        base.rect(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            base.CYAN if active else base.NAVY_2,
            line=base.CYAN if active else base.LINE,
            radius=True,
        )
        base.text(
            slide,
            x,
            Inches(0.78),
            Inches(1.58),
            Inches(0.29),
            label,
            10.5,
            bold=active,
            color=base.NAVY if active else base.MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def header(slide, page: int, eyebrow: str, title: str, subtitle: str) -> None:
    base.add_navigation(slide, 2, page)
    secondary_nav(slide)
    base.text(slide, Inches(0.48), Inches(1.28), Inches(4.8), Inches(0.25), eyebrow, 10.5, bold=True, color=base.CYAN)
    base.text(slide, Inches(0.46), Inches(1.53), Inches(12.2), Inches(0.49), title, 22, bold=True, font=base.TITLE_FONT)
    base.rect(slide, Inches(0.47), Inches(2.08), Inches(0.75), Inches(0.05), base.CYAN)
    base.text(slide, Inches(1.4), Inches(1.99), Inches(11.1), Inches(0.29), subtitle, 11.2, color=base.MUTED)


def card(slide, x, y, w, h, title: str, body: str, *, accent=base.CYAN, fill=base.NAVY_2, body_size=11.5):
    base.rect(slide, x, y, w, h, fill, line=accent, radius=True, transparency=3)
    base.rect(slide, x, y, Inches(0.08), h, accent)
    base.text(slide, x + Inches(0.2), y + Inches(0.14), w - Inches(0.4), Inches(0.3), title, 13, bold=True, color=accent)
    base.text(slide, x + Inches(0.2), y + Inches(0.58), w - Inches(0.4), h - Inches(0.7), body, body_size, color=base.PALE, line_spacing=0.95)


def pill(slide, x, y, w, text_value: str, color=base.CYAN):
    base.rect(slide, x, y, w, Inches(0.34), base.NAVY_3, line=color, radius=True)
    base.text(
        slide,
        x,
        y,
        w,
        Inches(0.34),
        text_value,
        10.2,
        bold=True,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def slide_positioning(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        1,
        "06  ADMIN · CONTROLLED CHAINS",
        "行政 Agent：制度有出处、数据有权限、月报可复核",
        "不做“什么都能回答”的行政助理，把高频事务拆成三条受控业务链",
    )

    items = [
        ("01", "政策答疑", "腾讯 IMA 检索原文\n关键数字与出处一致\n无依据、个案争议转人工", base.CYAN),
        ("02", "敏感查询", "飞书身份进入脚本闸门\n通过后才读取 HCM\n未知身份默认拒绝", base.ORANGE),
        ("03", "HR 月报", "脱敏业务台账统一导入\npandas 计算指标与图表\n18 项数字保留来源与算式", GREEN),
    ]
    for i, (num, title, body, accent) in enumerate(items):
        x = Inches(0.48 + i * 4.18)
        base.rect(slide, x, Inches(2.42), Inches(3.9), Inches(2.95), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x + Inches(0.2), Inches(2.64), Inches(0.72), Inches(0.48), num, 23, bold=True, color=accent)
        base.text(slide, x + Inches(1.0), Inches(2.7), Inches(2.55), Inches(0.35), title, 15, bold=True, color=base.PALE)
        base.rect(slide, x + Inches(0.22), Inches(3.22), Inches(3.46), Inches(0.04), accent)
        base.text(slide, x + Inches(0.28), Inches(3.52), Inches(3.34), Inches(1.42), body, 12.2, color=base.PALE, line_spacing=1.0)

    base.rect(slide, Inches(0.48), Inches(5.72), Inches(12.36), Inches(0.86), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(
        slide,
        Inches(0.72),
        Inches(5.86),
        Inches(11.9),
        Inches(0.52),
        "共同边界｜小慧负责检索、整理和计算；争议判断、破例审批与最终决策仍由人工完成",
        13,
        bold=True,
        color=base.CYAN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约45秒。行政风险集中在三件事：政策能否找到依据、敏感数据该不该让这个人看到、"
        "月报数字能否复核。因此拆成政策答疑、敏感查询、HR月报三条受控业务链。"
        "小慧只负责检索、整理和计算，争议、审批和最终决定仍由人完成。",
    )


def slide_policy(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        2,
        "06  POLICY · GROUNDED ANSWER",
        "有依据才回答：检索成功不等于可以回答",
        "先判断“该不该答”，再判断“能不能找到”；任何一步失败都不靠模型常识补齐",
    )

    steps = [
        ("01", "飞书提问", "自然语言"),
        ("02", "腾讯 IMA", "检索制度"),
        ("03", "原文片段", "返回 hits"),
        ("04", "数字核对", "逐字一致"),
        ("05", "引用答复", "注明出处"),
    ]
    for i, (num, title, sub) in enumerate(steps):
        x = Inches(0.48 + i * 2.5)
        base.rect(slide, x, Inches(2.4), Inches(2.22), Inches(1.02), base.NAVY_2, line=base.CYAN, radius=True)
        base.text(slide, x + Inches(0.15), Inches(2.53), Inches(0.45), Inches(0.25), num, 10, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(0.58), Inches(2.5), Inches(1.45), Inches(0.28), title, 12.2, bold=True, color=base.PALE)
        base.text(slide, x + Inches(0.18), Inches(2.9), Inches(1.86), Inches(0.25), sub, 10.5, color=base.MUTED, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            connector(slide, x + Inches(2.22), Inches(2.91), x + Inches(2.47), Inches(2.91), base.CYAN)

    card(
        slide,
        Inches(0.48),
        Inches(3.78),
        Inches(3.92),
        Inches(1.58),
        "能不能找到",
        "未命中 / 凭证异常\n明确停止，不用模型常识补答",
        accent=RED,
        fill=RED_DARK,
        body_size=12,
    )
    card(
        slide,
        Inches(4.7),
        Inches(3.78),
        Inches(3.92),
        Inches(1.58),
        "该不该回答",
        "工资差异 / 劳动争议\n法律定性 / 破例审批 → 转人工",
        accent=base.ORANGE,
        body_size=12,
    )
    card(
        slide,
        Inches(8.92),
        Inches(3.78),
        Inches(3.92),
        Inches(1.58),
        "输出纪律",
        "关键数字与原文一致\n结尾注明政策出处与处理渠道",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=12,
    )

    base.rect(slide, Inches(0.48), Inches(5.65), Inches(12.36), Inches(0.92), base.NAVY_3, line=base.ORANGE, radius=True)
    base.text(slide, Inches(0.7), Inches(5.8), Inches(2.25), Inches(0.25), "法律锚点", 11.5, bold=True, color=base.ORANGE)
    base.text(
        slide,
        Inches(2.15),
        Inches(5.74),
        Inches(10.35),
        Inches(0.55),
        "《劳动合同法》第四条｜直接涉及劳动者切身利益的规章制度，应依法制定，并公示或告知劳动者",
        11.7,
        color=base.PALE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约1分钟。员工用自然语言提问，但小慧必须先调用腾讯IMA检索制度原文，"
        "关键数字与原文一致并注明出处。检索成功不等于应该回答；工资差异、劳动争议、"
        "法律定性和破例审批涉及个人事实或管理决策，一律转人工。完整流程段如展示行政，"
        "只快速做一次普通政策问答，不重复本页机制。",
    )


def slide_monthly_report(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        3,
        "06  REPORT · DETERMINISTIC NUMBERS",
        "9月月报：数字由脚本计算，观察期不补数",
        "报告期 2026-09｜生成于 2026-09-21｜与成果页同一张演示月报",
    )

    source_labels = ["招聘过程", "岗位编制", "入职记录", "培训记录", "招聘费用", "离职登记"]
    for i, label in enumerate(source_labels):
        pill(slide, Inches(0.48 + i * 1.51), Inches(2.42), Inches(1.34), label, base.CYAN)
    connector(slide, Inches(9.52), Inches(2.59), Inches(9.96), Inches(2.59), base.CYAN)
    base.rect(slide, Inches(10.02), Inches(2.32), Inches(2.82), Inches(0.54), base.NAVY_2, line=GREEN, radius=True)
    base.text(
        slide,
        Inches(10.02),
        Inches(2.32),
        Inches(2.82),
        Inches(0.54),
        "统一数据契约 → pandas",
        11.2,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    kpis = [
        ("9 人", "本月入职", base.CYAN),
        ("94.3%", "岗位整体到位率", base.CYAN),
        ("40 天", "Time to Hire", base.CYAN),
        ("8000 元", "人均招聘成本", GREEN),
        ("91.7%", "培训完成率", GREEN),
        ("统计中", "45天内离职率", base.ORANGE),
    ]
    for i, (value, label, accent) in enumerate(kpis):
        row, col = divmod(i, 3)
        x = Inches(0.48 + col * 4.18)
        y = Inches(3.2 + row * 1.18)
        base.rect(slide, x, y, Inches(3.9), Inches(0.98), base.NAVY_2, line=accent, radius=True)
        base.text(slide, x + Inches(0.22), y + Inches(0.13), Inches(1.45), Inches(0.36), value, 18, bold=True, color=accent)
        base.text(slide, x + Inches(1.7), y + Inches(0.22), Inches(1.96), Inches(0.28), label, 11.2, bold=True, color=base.PALE, align=PP_ALIGN.RIGHT)
        if label == "45天内离职率":
            base.text(slide, x + Inches(0.22), y + Inches(0.61), Inches(3.42), Inches(0.2), "观察期未结束，不补数", 9.4, color=base.MUTED, align=PP_ALIGN.RIGHT)

    base.rect(slide, Inches(0.48), Inches(5.72), Inches(12.36), Inches(0.9), base.NAVY_3, line=GREEN, radius=True)
    base.text(slide, Inches(0.72), Inches(5.84), Inches(2.3), Inches(0.25), "数据口径", 11.5, bold=True, color=GREEN)
    base.text(
        slide,
        Inches(2.1),
        Inches(5.78),
        Inches(10.45),
        Inches(0.6),
        "脱敏业务台账验证计算与回核机制｜落地可接企业导出表或 API，数据契约不变｜18项数字保留来源和算式",
        11.4,
        color=base.PALE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    base.add_notes(
        slide,
        "约1分30秒。九月演示月报读取招聘、岗位、入职、培训、费用和离职六类脱敏业务台账，"
        "由脚本完成计算、图表和来源回核，模型不碰数字。与成果页同一张：报告期2026-09，生成于2026-09-21。"
        "不为凑完整而补观察期未结束的数。实际部署可接企业导出表或API，数据契约不变。",
    )


def slide_permission_demo(prs: Presentation) -> None:
    slide = base.new_slide(prs)
    base.add_tech_background(slide)
    header(
        slide,
        4,
        "06  PERMISSION · LIVE DEMO",
        "权限分级：同一请求，不同身份，不同结果",
        "不提前展示答案；下面切换两个真实飞书身份，验证敏感数据在读取前被拦截",
    )

    chain = [
        ("01", "飞书 sender_id"),
        ("02", "身份识别"),
        ("03", "脚本权限闸门"),
        ("04", "是否读取 HCM"),
    ]
    for i, (num, label) in enumerate(chain):
        x = Inches(0.68 + i * 3.12)
        base.rect(slide, x, Inches(2.38), Inches(2.64), Inches(0.72), base.NAVY_2, line=base.CYAN, radius=True)
        base.text(slide, x + Inches(0.16), Inches(2.57), Inches(0.38), Inches(0.22), num, 9.5, bold=True, color=base.CYAN)
        base.text(slide, x + Inches(0.55), Inches(2.49), Inches(1.88), Inches(0.31), label, 11.8, bold=True, color=base.PALE, align=PP_ALIGN.CENTER)
        if i < len(chain) - 1:
            connector(slide, x + Inches(2.64), Inches(2.74), x + Inches(3.05), Inches(2.74), base.CYAN)

    request = "请查询员工1的合同薪资"
    card(
        slide,
        Inches(0.48),
        Inches(3.45),
        Inches(4.0),
        Inches(1.74),
        "普通员工（演示）",
        f"相同请求：{request}\n\n结果：现场验证",
        accent=base.ORANGE,
        body_size=11.8,
    )
    base.rect(slide, Inches(4.75), Inches(3.68), Inches(3.83), Inches(1.28), base.NAVY_3, line=base.CYAN, radius=True)
    base.text(slide, Inches(4.98), Inches(3.84), Inches(3.35), Inches(0.28), "同一条飞书消息", 13, bold=True, color=base.CYAN, align=PP_ALIGN.CENTER)
    base.text(slide, Inches(4.98), Inches(4.28), Inches(3.35), Inches(0.28), "权限检查发生在读取数据之前", 10.8, color=base.PALE, align=PP_ALIGN.CENTER)
    card(
        slide,
        Inches(8.86),
        Inches(3.45),
        Inches(4.0),
        Inches(1.74),
        "HR 专员（演示）",
        f"相同请求：{request}\n\n结果：现场验证",
        accent=GREEN,
        fill=GREEN_DARK,
        body_size=11.8,
    )

    principles = [
        "未知身份默认拒绝",
        "拒绝返回非零退出码",
        "不依赖模型临场判断",
    ]
    for i, item in enumerate(principles):
        pill(slide, Inches(0.48 + i * 3.0), Inches(5.52), Inches(2.72), item, base.ORANGE if i == 0 else base.CYAN)

    base.rect(slide, Inches(9.68), Inches(5.38), Inches(3.16), Inches(1.07), base.NAVY_3, line=PURPLE, radius=True)
    base.text(slide, Inches(9.86), Inches(5.5), Inches(2.78), Inches(0.22), "法律锚点", 10.5, bold=True, color=PURPLE)
    base.text(slide, Inches(9.86), Inches(5.8), Inches(2.78), Inches(0.42), "《个人信息保护法》\n第六条 · 第五十一条", 10.4, color=base.PALE, align=PP_ALIGN.CENTER)

    base.add_notes(
        slide,
        "约4分钟。切换飞书现场验证。普通员工和HR专员发送完全相同的“请查询员工1的合同薪资”。"
        "系统从消息上下文取得sender_id，在读取HCM前进入脚本权限闸门。权限不通过时以非零退出码停止，"
        "模型拿不到姓名和薪资；HR身份通过后才读取HCM，并只返回脚本结果。"
        "收口：分级不是换一种说法，而是不同身份拥有不同的数据入口。"
        "对应个人信息保护法第六条最小必要原则和第五十一条权限管理要求。",
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = base.W, base.H
    base.TOTAL = TOTAL
    slide_positioning(prs)
    slide_policy(prs)
    slide_monthly_report(prs)
    slide_permission_demo(prs)
    prs.core_properties.title = "小慧 · 行政模块技能演示无图初版"
    prs.core_properties.subject = "政策有出处、敏感数据有权限、月报数字可复核"
    prs.core_properties.author = "小慧项目团队"
    prs.save(str(OUT))
    print(f"wrote {OUT} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
