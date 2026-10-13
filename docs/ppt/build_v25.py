"""v2.5：在 v2.4（39 页，含手动插入的一期产物页）上按“一页一观点”改项目背景到总控一段，统一编号。v2.4 不改。"""
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v2.4.pptx"
DST = "/opt/openclaw-hr/docs/PPT准备/小慧-图文版v2.5.pptx"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
E = 914400
CARD, MID = "032B5C", "04457D"
CYAN, GREEN, ORANGE = "00D8FF", "26D098", "FFA421"
WHITE, LIGHT, GREY = "FFFFFF", "E9F8FF", "9EC2D8"
C = RGBColor.from_string

p = Presentation(SRC)
old = list(p.slides)
POLICY, SURVEY, RULES, PITS, PRINCIPLE = old[3], old[4], old[5], old[6], old[7]
ARCH, CONTROL, RECRUIT = old[9], old[11], old[13]

lst = p.slides._sldIdLst
el = lst[6]
lst.remove(el)
lst.insert(3, el)


def shape(s, sid):
    return next(sh for sh in s.shapes if sh.shape_id == sid)


def drop(s, ids):
    for sh in list(s.shapes):
        if sh.shape_id in ids:
            sh._element.getparent().remove(sh._element)


def set_text(sh, text):
    paras = sh.text_frame.paragraphs
    runs = paras[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""
    for para in paras[1:]:
        para._p.getparent().remove(para._p)


def restyle(sh, size=None, bold=None, color=None):
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            if size:
                r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
            if color:
                r.font.color.rgb = C(color)


def move(sh, l=None, t=None, w=None, h=None):
    if l is not None:
        sh.left = Inches(l)
    if t is not None:
        sh.top = Inches(t)
    if w is not None:
        sh.width = Inches(w)
    if h is not None:
        sh.height = Inches(h)


def rect(s, l, t, w, h, fill, line=None, lw=1.0):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = C(fill)
    if line:
        r.line.color.rgb = C(line)
        r.line.width = Pt(lw)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def tb(s, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, font="微软雅黑"):
    b = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, (text, size, color, bold) in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.line_spacing = spacing
        r = para.add_run()
        r.text = text
        f = r.font
        f.size, f.color.rgb, f.bold, f.name = Pt(size), C(color), bold, font
        rpr = r._r.get_or_add_rPr()
        rpr.append(rpr.makeelement(A + "ea", {"typeface": font}))
    return b


def notes(s):
    return s.notes_slide.notes_text_frame


def replace_notes(s, pairs):
    tf = notes(s)
    text = tf.text
    for a, b in pairs:
        assert a in text, (a[:20], text[:40])
        text = text.replace(a, b)
    tf.text = text


# ---------- 封面：过渡句改接三道坎 ----------
replace_notes(old[2], [("为什么现在做，又为什么非停下来不可？先看国家怎么说。",
                        "为什么非停下来不可？因为智能体进职场，现在卡在三道坎上。")])

# ---------- 三道坎：提到第 4 页 ----------
for sh in PITS.shapes:
    if sh.has_text_frame and "三部门《智能体规范应用与创新发展实施意见》" in sh.text_frame.text:
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                r.text = r.text.replace("三部门《智能体规范应用与创新发展实施意见》第 6、7 条",
                                        "三部门《智能体规范应用与创新发展实施意见》（2026.05）第 6、7 条")
                r.text = r.text.replace("三部门《智能体规范应用与创新发展实施意见》第6、7条",
                                        "三部门《智能体规范应用与创新发展实施意见》（2026.05）第 6、7 条")
replace_notes(PITS, [(
    "【指第二、三张卡的引文】这几道坎不是我们自己想出来的担心：国家的《实施意见》点了名，企业的 AI 使用规定也写了“输出要核、决定在人”。"
    "刚才那些法和规，就是冲着这三道坎去的。\n\n线划好了，坎也摆在这儿了。怎么跨过去？",
    "【指第二、三张卡的引文】这几道坎不是我们自己想出来的担心，卡片上这两条要求出自同一份国家文件。"
    "国家一边在推智能体进千行百业，一边已经把线划了出来。")])

# ---------- 政策：只留《实施意见》这一句 ----------
drop(POLICY, {57, 58, 59})
quote = shape(POLICY, 56)
move(quote, l=5.7, w=6.9)
restyle(quote, size=20)
replace_notes(POLICY, [("【指右侧】小慧就按这条线分工", "小慧就按这条线分工")])

# ---------- 企业制度：每张卡只留规矩名和小慧的做法 ----------
CLAUSE = {67: 65, 74: 72, 81: 79, 88: 86, 95: 93, 102: 100}
DOING = {65: 69, 72: 76, 79: 83, 86: 90, 93: 97, 100: 104}
clauses = []
for body, head in CLAUSE.items():
    clauses.append("%s：%s" % (shape(RULES, head).text_frame.text.strip(), shape(RULES, body).text_frame.text.strip()))
drop(RULES, set(CLAUSE) | {68, 75, 82, 89, 96, 103})
for head, doing in DOING.items():
    h = shape(RULES, head)
    d = shape(RULES, doing)
    set_text(d, d.text_frame.text.replace("→ 小慧：", "").strip())
    move(d, t=h.top / E + 0.62, h=1.2)
    restyle(d, size=15, bold=True, color=LIGHT)
    tb(RULES, d.left / E, h.top / E + 0.44, 1.2, 0.22, [("小慧怎么做", 10, GREY, True)])
replace_notes(RULES, [
    ("我们先看条款，再看小慧怎么做。", "【指六张卡】每一条规矩，下面都对上了小慧的一个做法。"),
    ("国家和企业为什么都要把这些写得这么死？因为智能体真进了职场，会卡在三道坎上。",
     "国家划了线，企业立了规，三道坎也摆在这儿了。怎么跨过去？"),
])
notes(RULES).text = notes(RULES).text + "\n\n【条款原文，不念】\n" + "\n".join(clauses)

# ---------- 原则：换成三栏版 ----------
assert {87, 88} <= {sh.shape_id for sh in PRINCIPLE.shapes}
drop(PRINCIPLE, set(range(30, 87)))
set_text(shape(PRINCIPLE, 29), "用确定的脚本，框住不确定的生成；把会影响人的决定，交还给人")
COLS = [
    ("01", "模型：能写，也会错", CYAN,
     ["读得懂简历，写得出评语，", "答得了制度。", "", "但产出不确定，会把岗位里的", "数字安到候选人头上。"],
     ["评分写“58 万笔/秒”", "原简历是“12.8 万 TPS”"]),
    ("02", "脚本：不认说法，只认原件", GREEN,
     ["数字对原简历，回答对制度原文，", "查询先对权限；对不上就拦下。", "", "提示词会被忘、被说动；", "脚本过没过，看退出码。"],
     ["评分前敏感字段残留 0", "非 HR 查薪资，退出码 3"]),
    ("03", "人：关键一步，亲自拍板", ORANGE,
     ["录用、合同、账号、定级，", "关系到一个人的去留和权益。", "", "小慧递材料、给建议，", "签字的始终是人。"],
     ["审批门停下，等人来批", "绩效只给材料，不出评级"]),
]
for k, (num, head, col, body, proof) in enumerate(COLS):
    x = 0.48 + k * 4.2
    rect(PRINCIPLE, x, 2.2, 3.97, 3.86, CARD, line=MID)
    rect(PRINCIPLE, x, 2.2, 0.07, 3.86, col)
    tb(PRINCIPLE, x + 0.3, 2.36, 1.0, 0.36, [(num, 15, col, True)], anchor=MSO_ANCHOR.MIDDLE)
    tb(PRINCIPLE, x + 0.3, 2.74, 3.55, 0.5, [(head, 19, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
    tb(PRINCIPLE, x + 0.3, 3.3, 3.55, 1.75, [(ln, 14, LIGHT, False) for ln in body], spacing=1.2)
    rect(PRINCIPLE, x + 0.3, 5.0, 3.4, 0.015, MID)
    tb(PRINCIPLE, x + 0.3, 5.07, 1.2, 0.26, [("现场可见", 10.5, col, True)], anchor=MSO_ANCHOR.MIDDLE)
    tb(PRINCIPLE, x + 0.3, 5.35, 3.55, 0.62, [(ln, 12.5, WHITE, True) for ln in proof], anchor=MSO_ANCHOR.MIDDLE)
notes(PRINCIPLE).text = """【男1】让模型去写，让脚本去核，让人去定（约 60 秒）

我们的答案是一句话：让模型去写，让脚本去核，让人去定。

【指 01】模型读得懂、写得快：解析简历、写评语、按原文答制度。但它的产出是不确定的，会错。开发中它就把岗位要求里的 58 万笔每秒，写进了候选人的评分，原简历上是 12.8 万 TPS。写得越快，越需要有人把关。

【指 02】把关的是脚本。脚本不认说法，只认原件：数字对原简历，回答对制度原文，查询先对权限，对不上就拦下，模型改不了。为什么不在提示词里多叮嘱几句？因为提示词可能被忘、被说动，照没照做只能听模型自己说；脚本是写死的检查，过没过，看退出码。

【指 03】录用、合同、账号、定级，关系到一个人的去留和权益。小慧递材料、给建议，签字的始终是人。

【指底栏】写得快，核得准，定得稳。展开说，就是三句话：把法写进代码，把规写进流程，把决定留给人。每栏下面这一行，后面现场都会看到。"""

# ---------- 架构：专线卡只留一句话 ----------
drop(ARCH, set(range(104, 169)))
LINE = [(2.45, ["筛简历出画像", "人才库 · 面试纪要"]),
        (4.54, ["报到前跟进", "两道审批 · 试用期"]),
        (6.63, ["点名哪份", "就生成哪份"]),
        (8.72, ["360 反馈", "变成面谈材料"]),
        (10.81, ["答制度 · 查数据", "出月报"])]
for x, lines in LINE:
    tb(ARCH, x + 0.08, 4.62, 1.8, 1.4, [(ln, 14, LIGHT, True) for ln in lines],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.3)

# ---------- 总控：派发图 + 四条边界，配置挪到备注 ----------
config = shape(CONTROL, 74).text_frame.text
drop(CONTROL, {68, 69, 70, 71, 72, 73, 74})
for sh in CONTROL.shapes:
    if 46 <= sh.shape_id <= 67:
        sh.top = sh.top + Inches(0.55)
CARDS = [(75, 76, 77), (78, 79, 80), (81, 82, 83), (84, 85, 86)]
for i, (box, label, desc) in enumerate(CARDS):
    t = 2.40 + i * 0.95
    move(shape(CONTROL, box), l=6.63, t=t, w=6.22, h=0.83)
    lb, ds = shape(CONTROL, label), shape(CONTROL, desc)
    move(lb, l=6.85, t=t, w=1.6, h=0.83)
    move(ds, l=8.55, t=t, w=4.2, h=0.83)
    for x, size in ((lb, 17), (ds, 14)):
        x.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        restyle(x, size=size)
set_text(shape(CONTROL, 88), "后面每条专线都是同一套写法：配置划边界，脚本做核对  ·  分派另训本地小模型：派错 20 次 → 2 次")
notes(CONTROL).text = """【男1】总控只派不做：每个 Agent 的边界都写在配置里（约 45 秒）

【指左图】总控的边界写在配置里：能用的只有派发、等待、读文件、看图；执行命令、写文件、浏览器、定时任务全部禁用。它就算想自己出月报，也没有工具，只能派给行政。

【指右侧四条】六个 Agent 都是同一套写法：工具按岗位给，总控只派不做；派发名单只有总控有，专线之间不能互相派活；每条线有独立的工作区、规则和会话；一次最多 8 个子任务，超过 1800 秒就停。

【指底栏】分派判断我们还另训了一个本地小模型做对照：微调后派错从 20 次降到 2 次，详细结果在成果部分。

配置把总控能做的事限死了，只剩三件：判断、派发、转述。我们现场验证一下。

【配置原文，不念；评委问到时打开 openclaw.yaml】
""" + config

# ---------- 招聘封面：亮点数字前置 ----------
for k, (big, cap) in enumerate([("删 31 项，剩 0", "评分前敏感字段物理删除"),
                                ("58 万 → 12.8 万", "写错的数字，回原简历拦下")]):
    x = 6.95 + k * 2.9
    rect(RECRUIT, x, 1.72, 2.75, 1.1, CARD, line=CYAN)
    tb(RECRUIT, x + 0.1, 1.78, 2.55, 0.56, [(big, 18, CYAN, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tb(RECRUIT, x + 0.1, 2.34, 2.55, 0.4, [(cap, 11, LIGHT, False)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
replace_notes(RECRUIT, [("招聘是第一道坎：说得不准。",
                         "【指右上两个数】评分前删 31 项敏感字段，剩 0；评分里写错的 58 万，回原简历改回 12.8 万。后面现场就看这两个数怎么来的。\n\n招聘是第一道坎：说得不准。")])

# ---------- 编号与页码 ----------
SEC = {}
for n in range(4, 9):
    SEC[n] = "01-%d" % (n - 3)
for n in range(9, 12):
    SEC[n] = "02-%d" % (n - 8)
for rng, tag in ((range(12, 14), "03-1"), (range(14, 19), "03-2"), (range(19, 24), "03-3"),
                 (range(24, 29), "03-4"), (range(29, 33), "03-5")):
    for n in rng:
        SEC[n] = tag
for n in range(33, 39):
    SEC[n] = "04-%d" % (n - 32)

total = len(p.slides)
for n, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame or sh.top is None:
            continue
        t = sh.text_frame.text.strip()
        top, left = sh.top / E, sh.left / E
        if top < 0.6 and left > 11.5 and re.fullmatch(r"(\d+\s*/\s*\d+|—)", t):
            set_text(sh, "%02d / %d" % (n, total))
        elif 0.85 < top < 1.5 and re.match(r"^0\d\s{2}[A-Z]", t):
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if re.match(r"^0\d", r.text):
                        r.text = re.sub(r"^0\d\s{2}", "" if n == 1 else SEC[n] + "  ", r.text)
                        break

p.save(DST)
print("saved", DST, total)
